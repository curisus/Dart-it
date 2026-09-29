"""End-to-end latency harness: time until a flow's data is loaded into an .xlsx.

Subcommands (run from the repo root with ``uv run python``):
  resolve  untimed discovery of each company's FY audit filing and separate
           (별도) attachment -> output/perf/flows_resolved.json
  run      timed flows for one target in one process (= one session)
  ab       ``run --reps 1`` subprocesses between two project trees, arm order
           counterbalanced (ABBA) per company x rep; each child is started with
           a sanitized environment and ``--expect-code-root <its tree>``.
           ``--session-scope block`` instead runs every company in one child
           per rep and arm (ABBA over reps), so in-process caches (the company
           directory) carry over from one company to the next as in a server

Targets: ``local`` (in-process stdio MCP server object), ``remote`` (deployed
HTTPS endpoint) and ``remote-asgi`` (the same ASGI app of the code under test,
served in-process). Flow definitions live in ``flows.json`` next to this file.

Timing: ``server_ms`` is the sum of the tool-call times, ``loader_ms`` the
reference loader's workbook write (remote flows) and ``total_ms`` their sum.
Harness bookkeeping (hashing, record writing, verification) happens outside
the timed spans. Records are written under ``output/perf/<run_id>/``.

The OpenDART API key is read by the fitness runner from ``.env`` and only
ever travels in process environment or an Authorization header; it is never
written, printed or placed on a command line.
"""

from __future__ import annotations

import argparse
import asyncio
import concurrent.futures
import copy
import importlib
import importlib.util
import json
import logging
import os
import platform
import re
import secrets
import shutil
import subprocess
import sys
import threading
import time
from collections.abc import Callable, Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import UTC, datetime
from functools import cache
from pathlib import Path
from types import ModuleType
from typing import TYPE_CHECKING, Final, Protocol, cast

import httpx2
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from stage_probes import StageTimers, UpstreamMonitor, redact_url
from xlsx_invariance import flow_identity, xlsx_sha256_of_values

if TYPE_CHECKING:
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet
    from starlette.applications import Starlette

HARNESS_FILE: Final = Path(__file__).resolve()
HARNESS_ROOT: Final = HARNESS_FILE.parents[2]
FLOWS_PATH: Final = HARNESS_FILE.with_name("flows.json")
DEFAULT_OUT: Final = HARNESS_ROOT / "output" / "perf"
RESOLVED_NAME: Final = "flows_resolved.json"
ASGI_BASE_URL: Final = "http://testserver"
ASGI_URL: Final = ASGI_BASE_URL + "/api/mcp"
OPENDART_HOME: Final = "https://opendart.fss.or.kr/"
CURSOR_ENV_NAME: Final = "DART_MCP_CURSOR_SECRET"
TARGETS: Final = ("local", "remote", "remote-asgi")
REMOTE_TARGETS: Final = frozenset({"remote", "remote-asgi"})
IN_PROCESS_TARGETS: Final = frozenset({"local", "remote-asgi"})
MAX_PAGES: Final = 500
RTT_SAMPLES: Final = 5
EXIT_INVALID_FLOW: Final = 1
EXIT_USAGE: Final = 2
EXIT_BUDGET: Final = 3
_NS_PER_MS: Final = 1_000_000
MESSAGE_LIMIT: Final = 300
# A message cut before it reaches the harness (the fitness runner truncates
# HTTP error bodies) can end in the first characters of the key; a trailing
# key prefix at least this long is masked as well.
_MIN_TRAILING_KEY_PREFIX: Final = 4
# Variables that would let an ``ab`` child import code from outside its own
# project tree (or share one venv between both arms).
AB_CHILD_ENV_DROP: Final = frozenset(
    {"VIRTUAL_ENV", "PYTHONPATH", "PYTHONHOME", "UV_PROJECT_ENVIRONMENT"}
)
_RUN_ID_PATTERN: Final = re.compile(r"^[A-Za-z0-9._-]+$")

type Envelope = dict[str, object]
type JsonRow = dict[str, object]


class HarnessError(RuntimeError):
    """A configuration or input problem that stops the command."""


# ---------------------------------------------------------------- fitness reuse


@cache
def fitness_runner() -> ModuleType:
    """Import scripts/fitness/runner.py unchanged (evidence tooling, never edited)."""
    fitness_dir = str(HARNESS_ROOT / "scripts" / "fitness")
    if fitness_dir not in sys.path:
        sys.path.insert(0, fitness_dir)
    return importlib.import_module("runner")


def canonical(value: object) -> str:
    return cast("str", fitness_runner().canonical(value))


def sha256_text(text: str) -> str:
    return cast("str", fitness_runner().sha256_text(text))


def read_api_key() -> str:
    return cast("str", fitness_runner().read_env_key(HARNESS_ROOT))


def prod_url() -> str:
    return cast("str", fitness_runner().PROD_URL)


# ---------------------------------------------------------------- JSON helpers


def as_dict(value: object) -> dict[str, object] | None:
    return cast("dict[str, object]", value) if isinstance(value, dict) else None


def as_list(value: object) -> list[object] | None:
    return cast("list[object]", value) if isinstance(value, list) else None


def require_dict(value: object, where: str) -> dict[str, object]:
    result = as_dict(value)
    if result is None:
        msg = f"{where}: expected a JSON object"
        raise HarnessError(msg)
    return result


def require_str(container: dict[str, object], key: str, where: str) -> str:
    value = container.get(key)
    if not isinstance(value, str) or not value:
        msg = f"{where}: '{key}' must be a non-empty string"
        raise HarnessError(msg)
    return value


def optional_str(container: dict[str, object], key: str) -> str | None:
    value = container.get(key)
    return value if isinstance(value, str) and value else None


def str_tuple(value: object, where: str) -> tuple[str, ...]:
    items = as_list(value) if value is not None else []
    if items is None or any(not isinstance(item, str) for item in items):
        msg = f"{where}: expected a list of strings"
        raise HarnessError(msg)
    return tuple(cast("list[str]", items))


def utc_now() -> str:
    return datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%S.%fZ")


def append_jsonl(path: Path, record: dict[str, object]) -> None:
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False))
        handle.write("\n")


def ms(nanoseconds: int) -> float:
    return round(nanoseconds / _NS_PER_MS, 3)


# ---------------------------------------------------------------- flow catalog

VERIFY_KINDS: Final = frozenset(
    {"corp_code", "rcept_no", "attachment_id", "ids", "output_file", "page"}
)


@dataclass(frozen=True, slots=True)
class StepSpec:
    step: str
    tool: str
    args: dict[str, object]
    verify: tuple[str, ...]
    follow_cursor: bool
    expect_warnings: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class FlowSpec:
    flow_id: str
    surface: str
    representative: bool
    output_dir: str | None
    loader: bool
    companies: tuple[str, ...] | None
    steps: tuple[StepSpec, ...]

    @property
    def reuses(self) -> str | None:
        if self.output_dir is not None and self.output_dir.startswith("reuse:"):
            return self.output_dir.removeprefix("reuse:")
        return None

    def applies_to(self, corp_code: str) -> bool:
        return self.companies is None or corp_code in self.companies


@dataclass(frozen=True, slots=True)
class CompanySpec:
    corp_code: str
    name: str
    expected_rcept_no: str | None
    expected_attachment_id: str | None


@dataclass(frozen=True, slots=True)
class FlowCatalog:
    business_year: int
    reprt_code: str
    report_kind: str
    companies: tuple[CompanySpec, ...]
    flows: tuple[FlowSpec, ...]

    def flow(self, flow_id: str) -> FlowSpec | None:
        return next((flow for flow in self.flows if flow.flow_id == flow_id), None)


def _parse_step(raw: object, where: str) -> StepSpec:
    data = require_dict(raw, where)
    verify = str_tuple(data.get("verify"), f"{where}.verify")
    unknown = sorted(set(verify) - VERIFY_KINDS)
    if unknown:
        msg = f"{where}: unknown verify kinds {unknown}"
        raise HarnessError(msg)
    return StepSpec(
        step=require_str(data, "step", where),
        tool=require_str(data, "tool", where),
        args=require_dict(data.get("args", {}), f"{where}.args"),
        verify=verify,
        follow_cursor=data.get("follow_cursor") is True,
        expect_warnings=str_tuple(
            data.get("expect_warnings"), f"{where}.expect_warnings"
        ),
    )


def _parse_flow(raw: object, index: int) -> FlowSpec:
    where = f"flows[{index}]"
    data = require_dict(raw, where)
    surface = require_str(data, "surface", where)
    if surface not in {"local", "remote"}:
        msg = f"{where}: surface must be 'local' or 'remote'"
        raise HarnessError(msg)
    steps = as_list(data.get("steps"))
    if not steps:
        msg = f"{where}: steps must be a non-empty list"
        raise HarnessError(msg)
    companies = data.get("companies")
    return FlowSpec(
        flow_id=require_str(data, "id", where),
        surface=surface,
        representative=data.get("representative") is True,
        output_dir=optional_str(data, "output_dir"),
        loader=data.get("loader") is True,
        companies=None
        if companies is None
        else str_tuple(companies, f"{where}.companies"),
        steps=tuple(
            _parse_step(step, f"{where}.steps[{n}]") for n, step in enumerate(steps)
        ),
    )


def _parse_company(raw: object, index: int) -> CompanySpec:
    where = f"companies[{index}]"
    data = require_dict(raw, where)
    expected = as_dict(data.get("expected")) or {}
    return CompanySpec(
        corp_code=require_str(data, "corp_code", where),
        name=require_str(data, "name", where),
        expected_rcept_no=optional_str(expected, "rcept_no"),
        expected_attachment_id=optional_str(expected, "attachment_id"),
    )


def load_catalog(path: Path = FLOWS_PATH) -> FlowCatalog:
    data = require_dict(json.loads(path.read_text(encoding="utf-8")), str(path))
    year = data.get("business_year")
    if not isinstance(year, int):
        msg = "flows.json: business_year must be an integer"
        raise HarnessError(msg)
    companies = as_list(data.get("companies")) or []
    flows = as_list(data.get("flows")) or []
    return FlowCatalog(
        business_year=year,
        reprt_code=require_str(data, "reprt_code", "flows.json"),
        report_kind=require_str(data, "report_kind", "flows.json"),
        companies=tuple(_parse_company(item, n) for n, item in enumerate(companies)),
        flows=tuple(_parse_flow(item, n) for n, item in enumerate(flows)),
    )


def select_companies(catalog: FlowCatalog, spec: str) -> tuple[CompanySpec, ...]:
    if spec == "all":
        return catalog.companies
    wanted = [code.strip() for code in spec.split(",") if code.strip()]
    known = {company.corp_code: company for company in catalog.companies}
    unknown = [code for code in wanted if code not in known]
    if unknown or not wanted:
        msg = f"unknown or empty --companies: {unknown or spec}"
        raise HarnessError(msg)
    return tuple(known[code] for code in wanted)


def select_flows(catalog: FlowCatalog, spec: str, target: str) -> tuple[FlowSpec, ...]:
    surface = "local" if target == "local" else "remote"
    if spec == "all":
        return tuple(flow for flow in catalog.flows if flow.surface == surface)
    wanted = {name.strip() for name in spec.split(",") if name.strip()}
    unknown = sorted(wanted - {flow.flow_id for flow in catalog.flows})
    if unknown or not wanted:
        msg = f"unknown or empty --flows: {unknown or spec}"
        raise HarnessError(msg)
    # Catalog order, so a flow that reuses another's output always runs after it.
    selected = tuple(flow for flow in catalog.flows if flow.flow_id in wanted)
    for flow in selected:
        if flow.surface != surface:
            msg = f"flow {flow.flow_id} is a {flow.surface} flow; it cannot run on --target {target}"
            raise HarnessError(msg)
        if flow.reuses is not None and flow.reuses not in wanted:
            msg = f"flow {flow.flow_id} reuses {flow.reuses}'s output; select both"
            raise HarnessError(msg)
    return selected


class _HasCorpCode(Protocol):
    @property
    def corp_code(self) -> str: ...


# (flow, corp_code, rep): the key of one recorded sample in flows.jsonl
type PlannedSample = tuple[str, str, int]


def run_order[C: _HasCorpCode](
    flows: Sequence[FlowSpec], companies: Sequence[C], reps: range
) -> list[tuple[C, int, FlowSpec]]:
    """(company, rep, flow) in the order ``run`` measures them."""
    return [
        (company, rep, flow)
        for company in companies
        for rep in reps
        for flow in flows
        if flow.applies_to(company.corp_code)
    ]


def planned_samples(
    flows: Sequence[FlowSpec], companies: Sequence[_HasCorpCode], reps: range
) -> list[PlannedSample]:
    """Every sample a ``run`` of this plan records, one per (flow, company, rep)."""
    return [
        (flow.flow_id, company.corp_code, rep)
        for company, rep, flow in run_order(flows, companies, reps)
    ]


# ---------------------------------------------------------------- resolution


@dataclass(frozen=True, slots=True)
class ResolvedCompany:
    corp_code: str
    name: str
    rcept_no: str
    attachment_id: str

    def placeholders(self, catalog: FlowCatalog) -> dict[str, object]:
        return {
            "name": self.name,
            "corp_code": self.corp_code,
            "rcept_no": self.rcept_no,
            "attachment_id": self.attachment_id,
            "bsns_year": catalog.business_year,
            "reprt_code": catalog.reprt_code,
            "report_kind": catalog.report_kind,
        }


def load_resolved(path: Path) -> dict[str, ResolvedCompany]:
    if not path.is_file():
        msg = f"resolved flows file not found: {path} (run the 'resolve' subcommand first)"
        raise HarnessError(msg)
    data = require_dict(json.loads(path.read_text(encoding="utf-8")), str(path))
    companies = require_dict(data.get("companies"), f"{path}.companies")
    resolved: dict[str, ResolvedCompany] = {}
    for corp_code, raw in companies.items():
        entry = require_dict(raw, f"{path}.companies.{corp_code}")
        resolved[corp_code] = ResolvedCompany(
            corp_code=corp_code,
            name=require_str(entry, "name", corp_code),
            rcept_no=require_str(entry, "rcept_no", corp_code),
            attachment_id=require_str(entry, "attachment_id", corp_code),
        )
    return resolved


def substitute(value: object, values: dict[str, object]) -> object:
    """Replace every string that is exactly "{name}" by that value (type kept)."""
    if isinstance(value, str) and value.startswith("{") and value.endswith("}"):
        name = value[1:-1]
        if name in values:
            return values[name]
        msg = f"unknown placeholder {value}"
        raise HarnessError(msg)
    if isinstance(value, dict):
        return {key: substitute(item, values) for key, item in value.items()}
    if isinstance(value, list):
        return [substitute(item, values) for item in value]
    return value


# ---------------------------------------------------------------- callers


class ToolCaller(Protocol):
    def call(self, tool: str, args: dict[str, object]) -> Envelope: ...

    def set_output_dir(self, path: Path) -> None: ...

    def close(self) -> None: ...


class _InnerCaller(Protocol):
    def call(self, tool: str, args: dict[str, object]) -> object: ...


class LocalToolCaller:
    """runner.LocalCaller: the in-process stdio server's call_tool."""

    def __init__(self, output_dir: Path, key: str) -> None:
        output_dir.mkdir(parents=True, exist_ok=True)
        self._inner = fitness_runner().LocalCaller(HARNESS_ROOT, output_dir, key)

    def call(self, tool: str, args: dict[str, object]) -> Envelope:
        return cast("Envelope", self._inner.call(tool, args))

    def set_output_dir(self, path: Path) -> None:
        self._inner.set_output_dir(path)

    def close(self) -> None:
        return


class RemoteToolCaller:
    """runner.RemoteCaller (or its ASGI variant): JSON-RPC tools/call, Bearer key."""

    def __init__(self, inner: object) -> None:
        self._inner = inner

    def call(self, tool: str, args: dict[str, object]) -> Envelope:
        return cast("Envelope", cast("_InnerCaller", self._inner).call(tool, args))

    def set_output_dir(self, path: Path) -> None:
        del path

    def close(self) -> None:
        close = getattr(self._inner, "close", None)
        if callable(close):
            close()


class AsgiBridge:
    """Synchronous ``post`` onto an in-process ASGI app.

    The app's lifespan and one ``httpx2.AsyncClient`` over ``ASGITransport``
    live in a single long-running task on a private event loop thread, as a
    deployed server keeps them across requests; each ``post`` is scheduled
    onto that loop and awaited from the calling thread.
    """

    def __init__(self, app: Starlette) -> None:
        self._loop = asyncio.new_event_loop()
        self._thread = threading.Thread(
            target=self._loop.run_forever, name="perf-asgi", daemon=True
        )
        self._thread.start()
        self._ready: concurrent.futures.Future[None] = concurrent.futures.Future()
        self._stop: asyncio.Event | None = None
        self._client: httpx2.AsyncClient | None = None
        self._serving = asyncio.run_coroutine_threadsafe(self._serve(app), self._loop)
        self._ready.result(timeout=120)

    async def _serve(self, app: Starlette) -> None:
        try:
            stop = asyncio.Event()
            async with (
                app.router.lifespan_context(app),
                httpx2.AsyncClient(
                    transport=httpx2.ASGITransport(app=app),
                    base_url=ASGI_BASE_URL,
                    timeout=httpx2.Timeout(600.0),
                ) as client,
            ):
                self._client = client
                self._stop = stop
                self._ready.set_result(None)
                await stop.wait()
        except BaseException as error:
            if not self._ready.done():
                self._ready.set_exception(error)
            raise

    def post(
        self, url: str, *, content: bytes, headers: dict[str, str]
    ) -> httpx2.Response:
        client = self._client
        if client is None:
            msg = "ASGI bridge is not serving"
            raise HarnessError(msg)
        future = asyncio.run_coroutine_threadsafe(
            client.post(url, content=content, headers=headers), self._loop
        )
        return future.result()

    def close(self) -> None:
        stop = self._stop
        if stop is not None:
            self._loop.call_soon_threadsafe(stop.set)
        self._serving.result(timeout=120)
        self._loop.call_soon_threadsafe(self._loop.stop)
        self._thread.join(timeout=120)
        self._loop.close()


def asgi_tool_caller(app: Starlette, key: str) -> RemoteToolCaller:
    """runner.RemoteCaller with its HTTP client swapped for the ASGI bridge.

    Subclassing keeps RemoteCaller.call - request shape, Bearer key header and
    envelope parsing - the single definition shared with the remote target.
    """
    base = cast("type", fitness_runner().RemoteCaller)

    class AsgiRemoteCaller(base):  # type: ignore[misc,valid-type]
        _client: httpx2.Client | AsgiBridge

        def __init__(self) -> None:
            super().__init__(ASGI_URL, key)
            self._client.close()
            self._client = AsgiBridge(app)

    return RemoteToolCaller(AsgiRemoteCaller())


def code_under_test_root() -> Path:
    """Project root of the imported dart_crawler (editable install: <root>/src/dart_crawler)."""
    import dart_crawler

    package_file = dart_crawler.__file__
    if package_file is None:
        msg = "dart_crawler has no __file__; cannot locate the code under test"
        raise HarnessError(msg)
    root = Path(package_file).resolve().parents[2]
    if not (root / "pyproject.toml").is_file():
        msg = f"dart_crawler is not an editable project checkout: {package_file}"
        raise HarnessError(msg)
    return root


def normalized_root(path: str | Path) -> str:
    """A tree path for comparison: resolved, and case-folded on Windows."""
    return os.path.normcase(os.path.realpath(path))


def code_root_mismatch(actual: str | Path, expected: str | Path) -> str | None:
    """None when both paths name the same tree, else a readable reason."""
    if normalized_root(actual) == normalized_root(expected):
        return None
    return f"code under test is {actual}, expected {expected}"


def load_asgi_app(root: Path) -> Starlette:
    from starlette.applications import Starlette

    path = root / "api" / "mcp.py"
    spec = importlib.util.spec_from_file_location("perf_code_under_test_api_mcp", path)
    if spec is None or spec.loader is None:
        msg = f"cannot load ASGI entry point {path}"
        raise HarnessError(msg)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    app = getattr(module, "app", None)
    if not isinstance(app, Starlette):
        msg = f"{path} does not define a Starlette 'app'"
        raise HarnessError(msg)
    return app


def redact_secret(text: str, key: str, limit: int = MESSAGE_LIMIT) -> str:
    """Mask every occurrence of ``key`` (and a trailing key prefix), then truncate.

    Redacting first means truncation can never cut a key in half and leave its
    prefix behind; the trailing-prefix mask covers text that was already cut
    upstream.
    """
    if key:
        text = text.replace(key, "***")
        longest = min(len(key) - 1, len(text))
        for length in range(longest, _MIN_TRAILING_KEY_PREFIX - 1, -1):
            if text.endswith(key[:length]):
                text = text[:-length] + "***"
                break
    return text[:limit]


def safe_call(
    caller: ToolCaller, tool: str, args: dict[str, object], key: str
) -> Envelope:
    """Call a tool, turning a raised exception into a failure envelope."""
    try:
        return caller.call(tool, args)
    except Exception as error:  # noqa: BLE001 - the harness records, never crashes
        message = redact_secret(f"{type(error).__name__}: {error}", key)
        return {
            "ok": False,
            "data": None,
            "error": {"code": "HARNESS_EXCEPTION", "message": message},
            "warnings": [],
        }


def warning_codes(envelope: Envelope) -> list[str]:
    codes: list[str] = []
    for warning in as_list(envelope.get("warnings")) or []:
        code = (as_dict(warning) or {}).get("code")
        if isinstance(code, str):
            codes.append(code)
    return codes


def error_fields(envelope: Envelope, key: str) -> tuple[str | None, str | None]:
    error = as_dict(envelope.get("error"))
    if error is None:
        return None, None
    code = error.get("code")
    message = error.get("message")
    text = redact_secret(str(message), key) if message is not None else None
    return (str(code) if code is not None else None), text


# ---------------------------------------------------------------- resolve


def _data_list(envelope: Envelope, what: str, key: str) -> list[dict[str, object]]:
    if envelope.get("ok") is not True:
        code, message = error_fields(envelope, key)
        msg = f"{what} failed: {code} {message}"
        raise HarnessError(msg)
    items = as_list(envelope.get("data"))
    if items is None:
        msg = f"{what} returned no list"
        raise HarnessError(msg)
    return [entry for entry in (as_dict(item) for item in items) if entry is not None]


def _pick_filing(filings: list[dict[str, object]], year: int) -> dict[str, object]:
    candidates = [
        filing
        for filing in filings
        if filing.get("fiscal_year") == year
        and filing.get("report_period") == "FY"
        and filing.get("withdrawn") is not True
    ]
    if not candidates:
        msg = f"no FY{year} representative filing"
        raise HarnessError(msg)
    return max(
        candidates,
        key=lambda filing: (
            str(filing.get("receipt_date")),
            str(filing.get("rcept_no")),
        ),
    )


def _pick_separate_attachment(
    attachments: list[dict[str, object]],
) -> dict[str, object]:
    separate = [
        attachment for attachment in attachments if attachment.get("standalone") is True
    ]
    if not separate:
        msg = "no separate (standalone) attachment"
        raise HarnessError(msg)
    audit = [
        attachment
        for attachment in separate
        if "감사보고서" in str(attachment.get("title"))
    ]
    return (audit or separate)[0]


def pinned_id_mismatches(
    company: CompanySpec, rcept_no: str, attachment_id: str
) -> list[str]:
    """Differences between discovered ids and the ids pinned in flows.json."""
    problems: list[str] = []
    if company.expected_rcept_no is not None and rcept_no != company.expected_rcept_no:
        problems.append(
            f"{company.corp_code}: rcept_no {rcept_no} != pinned {company.expected_rcept_no}"
        )
    if (
        company.expected_attachment_id is not None
        and attachment_id != company.expected_attachment_id
    ):
        problems.append(
            f"{company.corp_code}: attachment_id {attachment_id} != pinned {company.expected_attachment_id}"
        )
    return problems


def check_resolved_pins(
    catalog: FlowCatalog, resolved: Mapping[str, ResolvedCompany]
) -> None:
    """Refuse a resolved file whose ids disagree with any pinned company."""
    problems = [
        problem
        for company in catalog.companies
        if company.corp_code in resolved
        for problem in pinned_id_mismatches(
            company,
            resolved[company.corp_code].rcept_no,
            resolved[company.corp_code].attachment_id,
        )
    ]
    if problems:
        msg = "resolved ids disagree with flows.json: " + "; ".join(problems)
        raise HarnessError(msg)


def resolve_company(
    caller: ToolCaller, catalog: FlowCatalog, company: CompanySpec, key: str
) -> dict[str, object]:
    companies = _data_list(
        safe_call(
            caller,
            "search_companies",
            {"company_query": company.name, "report_kind": catalog.report_kind},
            key,
        ),
        f"search_companies({company.name})",
        key,
    )
    if company.corp_code not in {entry.get("corp_code") for entry in companies}:
        msg = f"search_companies({company.name}) did not return {company.corp_code}"
        raise HarnessError(msg)
    filing = _pick_filing(
        _data_list(
            safe_call(
                caller,
                "list_report_filings",
                {"corp_code": company.corp_code, "report_kind": catalog.report_kind},
                key,
            ),
            f"list_report_filings({company.corp_code})",
            key,
        ),
        catalog.business_year,
    )
    rcept_no = str(filing.get("rcept_no"))
    attachment = _pick_separate_attachment(
        _data_list(
            safe_call(caller, "list_report_attachments", {"rcept_no": rcept_no}, key),
            f"list_report_attachments({rcept_no})",
            key,
        )
    )
    attachment_id = str(attachment.get("attachment_id"))
    problems = pinned_id_mismatches(company, rcept_no, attachment_id)
    if problems:
        raise HarnessError("; ".join(problems))
    return {
        "name": company.name,
        "rcept_no": rcept_no,
        "attachment_id": attachment_id,
        "attachment_title": attachment.get("title"),
        "report_name": filing.get("report_name"),
        "receipt_date": filing.get("receipt_date"),
        "fiscal_year": filing.get("fiscal_year"),
    }


def resolve_command(namespace: argparse.Namespace) -> int:
    catalog = load_catalog()
    key = read_api_key()
    if not key:
        print("OPEN_DART_API_KEY not found in .env or environment", file=sys.stderr)
        return EXIT_USAGE
    out_path = Path(namespace.out).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    caller = LocalToolCaller(out_path.parent / "resolve_scratch", key)
    monitor = UpstreamMonitor(namespace.max_upstream)
    monitor.install()
    resolved: dict[str, object] = {}
    errors: list[str] = []
    try:
        for company in catalog.companies:
            try:
                resolved[company.corp_code] = resolve_company(
                    caller, catalog, company, key
                )
            except HarnessError as error:
                errors.append(str(error))
            print(
                f"[resolve] {company.corp_code} {company.name}: {resolved.get(company.corp_code) or 'FAILED'}"
            )
    finally:
        monitor.uninstall()
    print(f"[resolve] upstream requests: {monitor.count}")
    if errors:
        for problem in errors:
            print(f"resolve error: {problem}", file=sys.stderr)
        return EXIT_USAGE
    document = {
        "schema_version": 1,
        "resolved_at_utc": utc_now(),
        "business_year": catalog.business_year,
        "reprt_code": catalog.reprt_code,
        "report_kind": catalog.report_kind,
        "upstream_request_count": monitor.count,
        "companies": resolved,
    }
    out_path.write_text(
        json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[resolve] wrote {out_path}")
    return 0


# ---------------------------------------------------------------- environment


def _git(root: Path, *args: str) -> str | None:
    git = shutil.which("git")
    if git is None:
        return None
    completed = subprocess.run(  # noqa: S603
        [git, "-C", str(root), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    return completed.stdout if completed.returncode == 0 else None


def git_state(root: Path) -> dict[str, object]:
    sha = _git(root, "rev-parse", "HEAD")
    status = _git(root, "status", "--porcelain", "--untracked-files=no")
    untracked = _git(root, "ls-files", "--others", "--exclude-standard")
    return {
        "git_sha": sha.strip() if sha else None,
        "git_dirty": None if status is None else bool(status.strip()),
        "git_untracked_count": None
        if untracked is None
        else len(untracked.splitlines()),
    }


def _one_get(client: httpx2.Client) -> float | str:
    started = time.perf_counter_ns()
    try:
        response = client.get(OPENDART_HOME)
        response.read()
    except httpx2.HTTPError as error:
        return type(error).__name__
    return ms(time.perf_counter_ns() - started)


def measure_opendart_rtt() -> dict[str, object]:
    """Median of RTT_SAMPLES unauthenticated GETs of the OpenDART home page."""
    samples: list[float] = []
    errors: list[str] = []
    with httpx2.Client(timeout=httpx2.Timeout(10.0)) as client:
        for _ in range(RTT_SAMPLES):
            outcome = _one_get(client)
            if isinstance(outcome, str):
                errors.append(outcome)
            else:
                samples.append(outcome)
    ordered = sorted(samples)
    median = None
    if ordered:
        middle = len(ordered) // 2
        median = (
            ordered[middle]
            if len(ordered) % 2
            else round((ordered[middle - 1] + ordered[middle]) / 2, 3)
        )
    return {
        "opendart_rtt_ms_median": median,
        "opendart_rtt_samples_ms": samples,
        "opendart_rtt_errors": errors,
    }


def write_env(
    run_dir: Path, session: dict[str, object], base: dict[str, object]
) -> None:
    path = run_dir / "env.json"
    if path.is_file():
        document = require_dict(json.loads(path.read_text(encoding="utf-8")), str(path))
        sessions = as_list(document.get("sessions")) or []
        sessions.append(session)
        document["sessions"] = sessions
    else:
        document = {
            **base,
            "opendart_rtt_ms_median": session.get("opendart_rtt_ms_median"),
            "sessions": [session],
        }
    path.write_text(
        json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8"
    )


# ---------------------------------------------------------------- run


@dataclass(slots=True)
class Session:
    run_id: str
    arm: str | None
    target: str
    run_dir: Path
    catalog: FlowCatalog
    caller: ToolCaller
    key: str
    monitor: UpstreamMonitor | None
    timers: StageTimers | None
    session_id: str
    first_flow_pending: bool = True
    first_call_pending: bool = True

    @property
    def calls_path(self) -> Path:
        return self.run_dir / "calls.jsonl"

    @property
    def flows_path(self) -> Path:
        return self.run_dir / "flows.jsonl"

    @property
    def probes_path(self) -> Path:
        return self.run_dir / "probes.jsonl"

    def upstream(self) -> int | None:
        return None if self.monitor is None else self.monitor.count

    def budget_exceeded(self) -> bool:
        return self.monitor is not None and self.monitor.exceeded


@dataclass(slots=True)
class FlowState:
    flow: FlowSpec
    company: ResolvedCompany
    rep: int
    values: dict[str, object]
    output_dir: Path | None
    server_ns: int = 0
    loader_ns: int = 0
    valid: bool = True
    reasons: list[str] = field(default_factory=list)
    warning_codes: set[str] = field(default_factory=set)
    possibly_cold: bool = False
    call_count: int = 0
    page_count: int = 0
    columns: tuple[str, ...] | None = None
    rows: list[JsonRow] = field(default_factory=list)
    reported_total_rows: int | None = None
    returned_rows: int = 0
    output_path: str | None = None

    def invalidate(self, reason: str) -> None:
        self.valid = False
        self.reasons.append(reason)


type Verifier = Callable[[FlowState, object, dict[str, object]], str | None]


def _contains(data: object, key: str, expected: str) -> str | None:
    values = {(as_dict(item) or {}).get(key) for item in as_list(data) or []}
    return None if expected in values else f"{key} {expected} not returned"


def _verify_ids(state: FlowState, data: object, args: dict[str, object]) -> str | None:
    del args
    entry = as_dict(data) or {}
    if entry.get("rcept_no") != state.company.rcept_no:
        return f"rcept_no {entry.get('rcept_no')} != pinned {state.company.rcept_no}"
    if entry.get("attachment_id") != state.company.attachment_id:
        return f"attachment_id {entry.get('attachment_id')} != pinned"
    return None


def _verify_output_file(
    state: FlowState, data: object, args: dict[str, object]
) -> str | None:
    del args
    entry = as_dict(data) or {}
    raw = entry.get("output_path") or entry.get("absolute_path")
    if not isinstance(raw, str):
        return "no output path returned"
    path = Path(raw).resolve()
    if not path.is_file():
        return f"output file missing: {path.name}"
    if state.output_dir is not None and state.output_dir.resolve() not in path.parents:
        return "output file outside the flow's output dir"
    state.output_path = str(path)
    return None


def _verify_page(state: FlowState, data: object, args: dict[str, object]) -> str | None:
    page = as_dict(data)
    request = as_dict(args.get("request")) or {}
    if page is None:
        return "no page returned"
    if page.get("domain") != request.get("domain"):
        return f"page domain {page.get('domain')} != {request.get('domain')}"
    columns = tuple(str(column) for column in as_list(page.get("columns")) or [])
    if state.columns is None:
        state.columns = columns
        total = page.get("total_rows")
        state.reported_total_rows = total if isinstance(total, int) else None
    elif columns != state.columns:
        return f"page {state.page_count} columns differ from page 0"
    rows = [
        row
        for row in (as_dict(item) for item in as_list(page.get("rows")) or [])
        if row is not None
    ]
    state.rows.extend(rows)
    returned = page.get("returned_rows")
    state.returned_rows += returned if isinstance(returned, int) else len(rows)
    state.page_count += 1
    return None


VERIFIERS: Final[dict[str, Verifier]] = {
    "corp_code": lambda state, data, _args: _contains(
        data, "corp_code", state.company.corp_code
    ),
    "rcept_no": lambda state, data, _args: _contains(
        data, "rcept_no", state.company.rcept_no
    ),
    "attachment_id": lambda state, data, _args: _contains(
        data, "attachment_id", state.company.attachment_id
    ),
    "ids": _verify_ids,
    "output_file": _verify_output_file,
    "page": _verify_page,
}


def _returned_rows(data: object) -> int | None:
    entry = as_dict(data)
    if entry is not None:
        for key in ("returned_rows", "total_rows", "section_count"):
            value = entry.get(key)
            if isinstance(value, int):
                return value
        return None
    items = as_list(data)
    return None if items is None else len(items)


def _record_call(
    session: Session,
    state: FlowState,
    step: StepSpec,
    *,
    envelope: Envelope,
    elapsed_ns: int,
    page_index: int,
    cold: bool,
    upstream: int | None,
) -> None:
    data = envelope.get("data")
    entry = as_dict(data) or {}
    code, message = error_fields(envelope, session.key)
    output_path = entry.get("output_path") or entry.get("absolute_path")
    append_jsonl(
        session.calls_path,
        {
            "run_id": session.run_id,
            "arm": session.arm,
            "session_id": session.session_id,
            "target": session.target,
            "flow": state.flow.flow_id,
            "corp_code": state.company.corp_code,
            "rep": state.rep,
            "step": step.step,
            "tool": step.tool,
            "call_index": state.call_count,
            "page_index": page_index,
            "elapsed_ms": ms(elapsed_ns),
            "ok": envelope.get("ok"),
            "error_code": code,
            "error_message": message,
            "warning_codes": warning_codes(envelope),
            "data_sha256": sha256_text(canonical(data)),
            "returned_rows": _returned_rows(data),
            "has_next_cursor": entry.get("next_cursor") is not None,
            "output_path": output_path if isinstance(output_path, str) else None,
            "possibly_cold": cold,
            "upstream_requests": upstream,
        },
    )


def _timed_call(
    session: Session, tool: str, args: dict[str, object]
) -> tuple[Envelope, int, bool]:
    cold = session.first_call_pending and session.target in REMOTE_TARGETS
    session.first_call_pending = False
    started = time.perf_counter_ns()
    envelope = safe_call(session.caller, tool, args, session.key)
    return envelope, time.perf_counter_ns() - started, cold


def _check_envelope(
    state: FlowState,
    step: StepSpec,
    envelope: Envelope,
    args: dict[str, object],
    key: str,
) -> None:
    codes = warning_codes(envelope)
    state.warning_codes.update(codes)
    if envelope.get("ok") is not True:
        code, _ = error_fields(envelope, key)
        state.invalidate(f"{step.step}: ok=false ({code})")
        return
    missing = [code for code in step.expect_warnings if code not in codes]
    if missing:
        state.invalidate(f"{step.step}: expected warnings missing {missing}")
    for kind in step.verify:
        reason = VERIFIERS[kind](state, envelope.get("data"), args)
        if reason is not None:
            state.invalidate(f"{step.step}: {reason}")


def _next_cursor(step: StepSpec, envelope: Envelope) -> str | None:
    if not step.follow_cursor:
        return None
    cursor = (as_dict(envelope.get("data")) or {}).get("next_cursor")
    return cursor if isinstance(cursor, str) and cursor else None


def run_step(session: Session, state: FlowState, step: StepSpec) -> None:
    args = cast("dict[str, object]", substitute(step.args, state.values))
    page_index = 0
    while True:
        before = session.upstream()
        envelope, elapsed_ns, cold = _timed_call(session, step.tool, args)
        after = session.upstream()
        state.server_ns += elapsed_ns
        state.possibly_cold = state.possibly_cold or cold
        upstream = None if before is None or after is None else after - before
        _record_call(
            session,
            state,
            step,
            envelope=envelope,
            elapsed_ns=elapsed_ns,
            page_index=page_index,
            cold=cold,
            upstream=upstream,
        )
        state.call_count += 1
        _check_envelope(state, step, envelope, args, session.key)
        if session.budget_exceeded():
            state.invalidate("upstream request budget exceeded")
        cursor = _next_cursor(step, envelope)
        if not state.valid or cursor is None:
            return
        page_index += 1
        if page_index >= MAX_PAGES:
            state.invalidate(f"{step.step}: more than {MAX_PAGES} pages")
            return
        args = copy.deepcopy(args)
        request = as_dict(args.get("request"))
        if request is None:
            state.invalidate(f"{step.step}: follow_cursor needs a 'request' argument")
            return
        request["cursor"] = cursor


def _loader_value(sheet: WriteOnlyWorksheet, value: object) -> object:
    # A client loader must keep text as text: openpyxl would store "=..." as
    # a formula and rejects control characters, so such strings get an
    # explicit string cell (control characters removed from the xlsx only;
    # rows_sha256 is computed from the unmodified JSON rows).
    if isinstance(value, str) and (
        value.startswith("=") or ILLEGAL_CHARACTERS_RE.search(value)
    ):
        cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value))
        cell.data_type = "s"
        return cell
    return value


def write_reference_xlsx(
    path: Path, columns: Sequence[str], rows: Sequence[JsonRow]
) -> None:
    """The reference loader: every page's columns + rows into one sheet."""
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("data")
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_loader_value(sheet, row.get(column)) for column in columns])
    workbook.save(path)


def run_loader(session: Session, state: FlowState) -> tuple[str | None, str | None]:
    """Write the reference xlsx (timed) and hash the rows (untimed)."""
    columns = state.columns or ()
    if (
        state.reported_total_rows is not None
        and state.returned_rows != state.reported_total_rows
    ):
        state.invalidate(
            f"returned rows {state.returned_rows} != total_rows {state.reported_total_rows}"
        )
        return None, None
    xlsx_dir = session.run_dir / "xlsx"
    xlsx_dir.mkdir(parents=True, exist_ok=True)
    path = xlsx_dir / f"{state.flow.flow_id}_{state.company.corp_code}_{state.rep}.xlsx"
    started = time.perf_counter_ns()
    write_reference_xlsx(path, columns, state.rows)
    state.loader_ns = time.perf_counter_ns() - started
    rows_hash = sha256_text(canonical({"columns": list(columns), "rows": state.rows}))
    return str(path), rows_hash


def prepare_output_dir(
    session: Session, flow: FlowSpec, corp_code: str, rep: int
) -> tuple[Path | None, str | None]:
    if flow.output_dir is None:
        return None, None
    local_out = session.run_dir / "local_out"
    reused = flow.reuses
    if reused is not None:
        directory = local_out / f"{reused}_{corp_code}_r{rep}"
        return (
            directory,
            None if directory.is_dir() else f"reused output dir missing ({reused})",
        )
    directory = local_out / f"{flow.flow_id}_{corp_code}_r{rep}"
    if directory.exists() and any(directory.iterdir()):
        return directory, "output dir not empty (reuse of a run id?)"
    directory.mkdir(parents=True, exist_ok=True)
    return directory, None


def _flow_record(
    session: Session,
    state: FlowState,
    *,
    first: bool,
    upstream: int | None,
    started: str,
) -> dict[str, object]:
    return {
        "run_id": session.run_id,
        "arm": session.arm,
        "session_id": session.session_id,
        "target": session.target,
        "flow": state.flow.flow_id,
        "corp_code": state.company.corp_code,
        "rep": state.rep,
        "representative": state.flow.representative,
        "started_utc": started,
        "total_ms": ms(state.server_ns + state.loader_ns),
        "server_ms": ms(state.server_ns),
        "loader_ms": ms(state.loader_ns),
        "first_in_session": first,
        "possibly_cold": state.possibly_cold,
        "valid": state.valid,
        "invalid_reasons": state.reasons,
        "call_count": state.call_count,
        "page_count": state.page_count,
        "returned_rows": state.returned_rows if state.flow.loader else None,
        "warning_codes": sorted(state.warning_codes),
        "upstream_request_count": upstream,
    }


def run_flow(
    session: Session, flow: FlowSpec, company: ResolvedCompany, rep: int
) -> dict[str, object]:
    started = utc_now()
    first = session.first_flow_pending
    session.first_flow_pending = False
    output_dir, problem = prepare_output_dir(session, flow, company.corp_code, rep)
    state = FlowState(
        flow=flow,
        company=company,
        rep=rep,
        values=company.placeholders(session.catalog),
        output_dir=output_dir,
    )
    context: dict[str, object] = {
        "flow": flow.flow_id,
        "corp_code": company.corp_code,
        "rep": rep,
    }
    if session.monitor is not None:
        session.monitor.set_context(context)
    if session.timers is not None:
        session.timers.reset()
    before = session.upstream()
    if problem is not None:
        state.invalidate(problem)
    elif output_dir is not None:
        session.caller.set_output_dir(output_dir)
    for step in flow.steps if state.valid else ():
        run_step(session, state, step)
        if not state.valid:
            break
    xlsx_path: str | None = None
    rows_hash: str | None = None
    if flow.loader and state.valid:
        xlsx_path, rows_hash = run_loader(session, state)
    after = session.upstream()
    record = _flow_record(
        session,
        state,
        first=first,
        upstream=None if before is None or after is None else after - before,
        started=started,
    )
    record["xlsx_path"] = xlsx_path if flow.loader else state.output_path
    record["rows_sha256"] = rows_hash
    values_hash = (
        xlsx_sha256_of_values(Path(state.output_path))
        if state.valid and state.output_path and not flow.loader
        else None
    )
    record["xlsx_sha256_of_values"] = values_hash
    # The identity two runs of this flow must share: its data hash plus the
    # flow's sorted unique warning codes (compare_runs recomputes it the same way).
    record["identity_sha256"] = flow_identity(
        rows_hash or values_hash, state.warning_codes
    )
    append_jsonl(session.flows_path, record)
    write_probes(session, context)
    return record


def write_probes(session: Session, context: dict[str, object]) -> None:
    if session.timers is None:
        return
    base = {
        "run_id": session.run_id,
        "arm": session.arm,
        "session_id": session.session_id,
        **context,
    }
    if session.monitor is not None:
        for request in session.monitor.drain_records():
            append_jsonl(
                session.probes_path,
                {
                    "run_id": session.run_id,
                    "arm": session.arm,
                    "session_id": session.session_id,
                    **request,
                },
            )
    append_jsonl(
        session.probes_path,
        {**base, "kind": "stages", "stages": session.timers.snapshot()},
    )


def _open_caller(
    target: str, run_dir: Path, key: str, url: str
) -> tuple[ToolCaller, Path | None, str | None]:
    """Build the target's caller; returns (caller, code root, url for env.json)."""
    if target == "local":
        caller = LocalToolCaller(run_dir / "local_out", key)
        return caller, code_under_test_root(), None
    if target == "remote":
        return (
            RemoteToolCaller(fitness_runner().RemoteCaller(url, key)),
            None,
            redact_url(url),
        )
    root = code_under_test_root()
    # The in-process ASGI surface signs page cursors with this secret; a
    # deployment sets its own, a harness session needs any 32+ byte value.
    os.environ.setdefault(CURSOR_ENV_NAME, secrets.token_urlsafe(48))
    app = load_asgi_app(root)
    return (
        asgi_tool_caller(app, key),
        root,
        f"asgi:{(root / 'api' / 'mcp.py').as_posix()}#/api/mcp",
    )


@contextmanager
def open_session(
    namespace: argparse.Namespace,
    catalog: FlowCatalog,
    run_dir: Path,
    key: str,
    *,
    planned: Sequence[PlannedSample],
) -> Iterator[Session]:
    target: str = namespace.target
    caller, root, env_url = _open_caller(target, run_dir, key, namespace.url)
    monitor: UpstreamMonitor | None = None
    timers: StageTimers | None = None
    if target in IN_PROCESS_TARGETS:
        monitor = UpstreamMonitor(
            namespace.max_upstream, record_requests=namespace.probes
        )
        monitor.install()
        if namespace.probes:
            timers = StageTimers()
            timers.install()
    elif namespace.probes:
        print(
            "note: --probes has no in-process code to probe on --target remote; ignored"
        )
    session_id = f"{datetime.now(UTC).strftime('%Y%m%dT%H%M%S%fZ')}-{os.getpid()}"
    base: dict[str, object] = {
        "run_id": namespace.run_id,
        "arm": namespace.arm,
        "timestamp_utc": utc_now(),
        "target": target,
        "url": env_url,
        "code_root": None if root is None else str(root),
        **(
            git_state(root)
            if root is not None
            else {"git_sha": None, "git_dirty": None}
        ),
        "python": platform.python_version(),
        "platform": platform.platform(),
        "harness_root": str(HARNESS_ROOT),
        "probes": bool(namespace.probes),
    }
    session_info = {
        "session_id": session_id,
        "started_utc": utc_now(),
        "pid": os.getpid(),
        **base,
        # compare_runs checks the recorded samples against this plan, so a
        # session that stopped early cannot pass as a complete one.
        "planned_samples": [
            {"flow": flow, "corp_code": corp_code, "rep": rep}
            for flow, corp_code, rep in planned
        ],
        **measure_opendart_rtt(),
    }
    write_env(run_dir, session_info, base)
    session = Session(
        run_id=namespace.run_id,
        arm=namespace.arm,
        target=target,
        run_dir=run_dir,
        catalog=catalog,
        caller=caller,
        key=key,
        monitor=monitor,
        timers=timers,
        session_id=session_id,
    )
    try:
        yield session
    finally:
        if timers is not None:
            timers.uninstall()
        if monitor is not None:
            monitor.uninstall()
        caller.close()


def _run_plan(
    namespace: argparse.Namespace,
) -> tuple[FlowCatalog, tuple[FlowSpec, ...], list[ResolvedCompany]]:
    catalog = load_catalog()
    flows = select_flows(catalog, namespace.flows, namespace.target)
    companies = select_companies(catalog, namespace.companies)
    resolved = load_resolved(
        Path(namespace.resolved)
        if namespace.resolved
        else Path(namespace.out) / RESOLVED_NAME
    )
    missing = [
        company.corp_code for company in companies if company.corp_code not in resolved
    ]
    if missing:
        msg = f"companies missing from the resolved file: {missing}"
        raise HarnessError(msg)
    check_resolved_pins(catalog, resolved)
    if namespace.expect_code_root:
        problem = code_root_mismatch(code_under_test_root(), namespace.expect_code_root)
        if problem is not None:
            raise HarnessError(problem)
    return catalog, flows, [resolved[company.corp_code] for company in companies]


def run_command(namespace: argparse.Namespace) -> int:
    if not _RUN_ID_PATTERN.match(namespace.run_id):
        print(
            "--run-id may contain only letters, digits, '.', '_' and '-'",
            file=sys.stderr,
        )
        return EXIT_USAGE
    catalog, flows, companies = _run_plan(namespace)
    key = read_api_key()
    if not key:
        print("OPEN_DART_API_KEY not found in .env or environment", file=sys.stderr)
        return EXIT_USAGE
    run_dir = Path(namespace.out).resolve() / namespace.run_id
    if namespace.arm:
        run_dir /= namespace.arm
    if namespace.arm is None and (run_dir / "flows.jsonl").exists():
        print(f"run id already used: {run_dir}", file=sys.stderr)
        return EXIT_USAGE
    run_dir.mkdir(parents=True, exist_ok=True)
    invalid = 0
    reps = range(namespace.rep_offset + 1, namespace.rep_offset + namespace.reps + 1)
    planned = planned_samples(flows, companies, reps)
    with open_session(namespace, catalog, run_dir, key, planned=planned) as session:
        for company, rep, flow in run_order(flows, companies, reps):
            record = run_flow(session, flow, company, rep)
            invalid += 0 if record["valid"] else 1
            print(
                f"[{flow.flow_id} {company.corp_code} r{rep}] total={record['total_ms']}ms "
                f"server={record['server_ms']}ms loader={record['loader_ms']}ms "
                f"valid={record['valid']} upstream={record['upstream_request_count']} "
                f"{record['invalid_reasons'] or ''}"
            )
            if session.budget_exceeded():
                print(
                    f"aborted: upstream budget {namespace.max_upstream} exceeded",
                    file=sys.stderr,
                )
                return EXIT_BUDGET
        print(
            f"--- {run_dir} (upstream requests this session: {session.upstream()}) ---"
        )
    return EXIT_INVALID_FLOW if invalid else 0


# ---------------------------------------------------------------- ab


def _arm_upstream(run_root: Path) -> int:
    total = 0
    for arm in ("A", "B"):
        path = run_root / arm / "flows.jsonl"
        if not path.is_file():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            count = (as_dict(json.loads(line)) or {}).get("upstream_request_count")
            total += count if isinstance(count, int) else 0
    return total


def _check_project(path: Path) -> Path:
    resolved = path.resolve()
    if (
        not (resolved / "pyproject.toml").is_file()
        or not (resolved / "api" / "mcp.py").is_file()
    ):
        msg = f"not a Dart it project tree (pyproject.toml + api/mcp.py): {resolved}"
        raise HarnessError(msg)
    return resolved


def _ab_command_line(
    uv: str,
    project: Path,
    namespace: argparse.Namespace,
    *,
    corp_codes: Sequence[str],
    rep: int,
    arm: str,
    budget: int,
    resolved: Path,
) -> list[str]:
    command = [
        uv,
        "run",
        "--frozen",
        "--project",
        str(project),
        "python",
        str(HARNESS_FILE),
        "run",
        "--target",
        namespace.target,
        "--flows",
        namespace.flows,
        "--companies",
        ",".join(corp_codes),
        "--reps",
        "1",
        "--rep-offset",
        str(rep - 1),
        "--run-id",
        namespace.run_id,
        "--arm",
        arm,
        "--out",
        str(Path(namespace.out).resolve()),
        "--resolved",
        str(resolved),
        "--max-upstream",
        str(budget),
        "--url",
        namespace.url,
        "--expect-code-root",
        str(project),
    ]
    if namespace.probes:
        command.append("--probes")
    return command


def ab_child_env(environ: Mapping[str, str]) -> dict[str, str]:
    """The parent environment minus anything that could mix the two trees.

    The parent runs inside this repo's venv; each child's ``uv run --project``
    must select its own project's venv and import only its own ``src``.
    """
    env = {
        name: value
        for name, value in environ.items()
        if name.upper() not in AB_CHILD_ENV_DROP
    }
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def ab_schedule(corp_codes: Sequence[str], reps: int) -> list[tuple[str, int, str]]:
    """(corp_code, rep, arm) in execution order, counterbalanced ABBA.

    Within a company the arm that goes first alternates with the rep (even rep
    index A->B, odd B->A) and the pattern flips for every other company, so
    neither arm systematically runs right after the other has warmed upstream
    caches.
    """
    schedule: list[tuple[str, int, str]] = []
    for company_index, corp_code in enumerate(corp_codes):
        for rep_index in range(reps):
            first, second = (
                ("A", "B") if (company_index + rep_index) % 2 == 0 else ("B", "A")
            )
            schedule.append((corp_code, rep_index + 1, first))
            schedule.append((corp_code, rep_index + 1, second))
    return schedule


def ab_block_schedule(reps: int) -> list[tuple[int, str]]:
    """(rep, arm) in execution order for ``--session-scope block``.

    Each block is one rep of every company in one session per arm; the arm
    that goes first alternates with the block (even A->B, odd B->A).
    """
    schedule: list[tuple[int, str]] = []
    for rep_index in range(reps):
        first, second = ("A", "B") if rep_index % 2 == 0 else ("B", "A")
        schedule.append((rep_index + 1, first))
        schedule.append((rep_index + 1, second))
    return schedule


# (corp_codes run by one child, rep, arm)
type AbJob = tuple[tuple[str, ...], int, str]


def ab_jobs(corp_codes: Sequence[str], reps: int, session_scope: str) -> list[AbJob]:
    """The ``run`` children of an ab run, in execution order."""
    if session_scope == "block":
        return [(tuple(corp_codes), rep, arm) for rep, arm in ab_block_schedule(reps)]
    return [
        ((corp_code,), rep, arm) for corp_code, rep, arm in ab_schedule(corp_codes, reps)
    ]


@dataclass(frozen=True, slots=True)
class AbPlan:
    uv: str
    projects: dict[str, Path]
    run_root: Path
    resolved: Path
    env: dict[str, str]


def _run_ab_jobs(
    plan: AbPlan, namespace: argparse.Namespace, companies: Sequence[CompanySpec]
) -> tuple[int, list[dict[str, object]]]:
    """One ``run --reps 1`` subprocess per job of ``ab_jobs``, in ABBA order."""
    runs: list[dict[str, object]] = []
    status = 0
    jobs = ab_jobs(
        [company.corp_code for company in companies],
        namespace.reps,
        namespace.session_scope,
    )
    for position, (corp_codes, rep, arm) in enumerate(jobs, start=1):
        budget = namespace.max_upstream - _arm_upstream(plan.run_root)
        if budget <= 0:
            print("aborted: upstream budget exhausted", file=sys.stderr)
            return EXIT_BUDGET, runs
        command = _ab_command_line(
            plan.uv,
            plan.projects[arm],
            namespace,
            corp_codes=corp_codes,
            rep=rep,
            arm=arm,
            budget=budget,
            resolved=plan.resolved,
        )
        completed = subprocess.run(command, env=plan.env, check=False)  # noqa: S603
        companies_field: dict[str, object] = (
            {"corp_codes": list(corp_codes)}
            if namespace.session_scope == "block"
            else {"corp_code": corp_codes[0]}
        )
        runs.append(
            {
                "position": position,
                **companies_field,
                "rep": rep,
                "arm": arm,
                "project": str(plan.projects[arm]),
                "exit_code": completed.returncode,
            }
        )
        if completed.returncode == EXIT_BUDGET:
            return EXIT_BUDGET, runs
        if completed.returncode != 0 and status == 0:
            status = completed.returncode
    return status, runs


def ab_command(namespace: argparse.Namespace) -> int:
    uv = shutil.which("uv")
    if uv is None:
        print("uv not found on PATH", file=sys.stderr)
        return EXIT_USAGE
    if namespace.target not in IN_PROCESS_TARGETS:
        print(
            "ab compares two project trees in-process; --target remote measures one "
            "deployment for both arms (use --target remote-asgi)",
            file=sys.stderr,
        )
        return EXIT_USAGE
    catalog = load_catalog()
    select_flows(catalog, namespace.flows, namespace.target)
    companies = select_companies(catalog, namespace.companies)
    projects = {
        "A": _check_project(Path(namespace.a_project)),
        "B": _check_project(Path(namespace.b_project)),
    }
    run_root = Path(namespace.out).resolve() / namespace.run_id
    if run_root.exists() and any(run_root.iterdir()):
        print(f"run id already used: {run_root}", file=sys.stderr)
        return EXIT_USAGE
    resolved = (
        Path(namespace.resolved)
        if namespace.resolved
        else Path(namespace.out) / RESOLVED_NAME
    ).resolve()
    check_resolved_pins(catalog, load_resolved(resolved))
    run_root.mkdir(parents=True, exist_ok=True)
    plan = AbPlan(
        uv=uv,
        projects=projects,
        run_root=run_root,
        resolved=resolved,
        env=ab_child_env(os.environ),
    )
    status, runs = _run_ab_jobs(plan, namespace, companies)
    ab_record = {
        "run_id": namespace.run_id,
        "created_utc": utc_now(),
        "projects": {arm: str(path) for arm, path in projects.items()},
        "target": namespace.target,
        "flows": namespace.flows,
        "companies": [company.corp_code for company in companies],
        "reps": namespace.reps,
        "session_scope": namespace.session_scope,
        "order": [
            f"block:r{run['rep']}:{run['arm']}"
            if namespace.session_scope == "block"
            else f"{run['corp_code']}:r{run['rep']}:{run['arm']}"
            for run in runs
        ],
        "arm_order": "".join(str(run["arm"]) for run in runs),
        "subprocesses": runs,
        "upstream_request_count": _arm_upstream(run_root),
    }
    (run_root / "ab.json").write_text(
        json.dumps(ab_record, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"--- ab run {run_root}: compare with scripts/perf/compare_runs.py --ab {run_root} ---"
    )
    return status


# ---------------------------------------------------------------- CLI


def _add_common(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--target", required=True, choices=TARGETS)
    parser.add_argument(
        "--flows",
        required=True,
        help="comma-separated flow ids, or 'all' for the target's surface",
    )
    parser.add_argument(
        "--companies", default="all", help="'all' or comma-separated corp_codes"
    )
    parser.add_argument("--reps", type=int, default=3)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--url", default=prod_url())
    parser.add_argument("--probes", action="store_true")
    parser.add_argument("--max-upstream", type=int, default=3000)
    parser.add_argument("--out", default=str(DEFAULT_OUT))
    parser.add_argument(
        "--resolved", default=None, help=f"default: <out>/{RESOLVED_NAME}"
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Dart it end-to-end latency harness")
    commands = parser.add_subparsers(dest="command", required=True)
    resolve = commands.add_parser(
        "resolve", help="untimed rcept_no / attachment_id discovery"
    )
    resolve.add_argument("--out", default=str(DEFAULT_OUT / RESOLVED_NAME))
    resolve.add_argument("--max-upstream", type=int, default=300)
    run = commands.add_parser("run", help="timed flows in one session")
    _add_common(run)
    run.add_argument("--arm", choices=("A", "B"), default=None, help=argparse.SUPPRESS)
    run.add_argument("--rep-offset", type=int, default=0, help=argparse.SUPPRESS)
    run.add_argument(
        "--expect-code-root",
        default=None,
        help="refuse to run unless dart_crawler is imported from this project tree",
    )
    ab = commands.add_parser(
        "ab", help="alternate run subprocesses between two project trees"
    )
    _add_common(ab)
    ab.add_argument("--a-project", required=True)
    ab.add_argument("--b-project", required=True)
    ab.add_argument(
        "--session-scope",
        choices=("rep", "block"),
        default="rep",
        help="rep: one child per company x rep x arm; "
        "block: one child per rep x arm running every company",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    # httpx2 logs every request URL at INFO, and a product request URL carries
    # the API key in its query string; keep those lines from ever printing.
    logging.getLogger("httpx2").setLevel(logging.WARNING)
    namespace = build_parser().parse_args(argv)
    handlers: dict[str, Callable[[argparse.Namespace], int]] = {
        "resolve": resolve_command,
        "run": run_command,
        "ab": ab_command,
    }
    try:
        return handlers[namespace.command](namespace)
    except HarnessError as error:
        print(f"error: {error}", file=sys.stderr)
        return EXIT_USAGE


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
