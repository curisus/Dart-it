"""Deterministic oracles over runner records.

Usage:
  uv run python scripts/fitness/oracles.py --run-dir <out>/<run_id>/<scenario> --checks checks.json [--out oracles.json]

checks.json = [
  {"name": "...", "type": "envelope_contract", "record": "<step>"},
  {"type": "expect_ok", "record": "<step>"},
  {"type": "expect_error", "record": "<step>", "code": "INVALID_INPUT"},
  {"type": "expect_warnings", "record": "<step>", "codes": ["PARTIAL_COLLECTION"], "mode": "subset|exact|none"},
  {"type": "row_count_consistency", "record": "<step>", "rows_key": "accounts", "count_key": "returned_row_count"},
  {"type": "official_reconciliation", "api_record": "<step>", "sections_record": "<step>", "accounts": ["자산총계","부채총계","자본총계"]},
  {"type": "balance_identity", "record": "<step>"},
  {"type": "page_sum", "records": ["<step>", "<step>_p1"]},
  {"type": "equal_data", "records": ["<a>", "<b>"], "ignore_keys": ["parser_version"]},
  {"type": "xlsx_report", "path": "<xlsx>", "expect_section_count": 47},
  {"type": "xlsx_query", "path": "<xlsx>", "expect_total_rows": 141, "expect_domain": "get_financial_statements"},
  {"type": "md_utf8_order", "path": "<md>", "expect_titles": ["재무상태표", "손익계산서"]},
  {"type": "note_grep", "record": "<step>", "pattern": "전환사채|RCPS", "excerpt": 300},
  {"type": "sla_p95", "records": ["<step>"], "limit_ms": 10000},
  {"type": "elapsed_max", "record": "<step>", "limit_ms": 120000}
]
Records are addressed by their step name (file <seq>_<step>.json in the scenario dir).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

ERROR_CODES = {
    "CONFIG_ERROR", "INVALID_INPUT", "NOT_FOUND", "UPSTREAM_AUTH", "UPSTREAM_RATE_LIMIT", "UPSTREAM_UNAVAILABLE",
    "UPSTREAM_LAYOUT_CHANGED", "PARSE_FAILED", "CORE_STATEMENT_MISSING", "OUTPUT_WRITE_FAILED", "VALIDATION_FAILED",
}
WARNING_CODES = {
    "PARTIAL_COLLECTION", "IMAGE_CONTENT_SKIPPED", "AMOUNT_MISMATCH", "COMPARISON_UNAVAILABLE", "FALLBACK_SOURCE_USED",
    "ORIGINAL_FILING_SOURCE_USED", "VIEWER_DISCOVERY_SKIPPED", "EXISTING_FILE_REUSED",
}
SCALES = (1, 1_000, 1_000_000, 100_000_000)
TOTAL_ACCOUNTS = ("자산총계", "부채총계", "자본총계")
INDENT_CHARS = (" ", "\t", "\u00a0", "\u3000")


def load_records(run_dir: Path) -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for file in sorted(run_dir.glob("*.json")):
        if file.name.endswith(".data.json") or file.name in ("oracles.json", "checks.json"):
            continue
        record = json.loads(file.read_text(encoding="utf-8"))
        if "step" not in record or "data_path" not in record:
            continue
        record["_data"] = json.loads(Path(record["data_path"]).read_text(encoding="utf-8"))
        records[record["step"]] = record
    return records


def _rec(records: dict[str, dict[str, Any]], name: str) -> dict[str, Any]:
    if name not in records:
        raise KeyError(f"record not found: {name} (have {sorted(records)})")
    return records[name]


def _num(text: Any) -> Decimal | None:
    if isinstance(text, bool):
        return None
    if isinstance(text, (int, float)):
        return Decimal(str(text))
    if not isinstance(text, str):
        return None
    cleaned = text.strip().replace(",", "").replace(" ", "")
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = cleaned.strip("()")
    if cleaned.startswith("-"):
        negative, cleaned = True, cleaned[1:]
    if not re.fullmatch(r"\d+(\.\d+)?", cleaned):
        return None
    value = Decimal(cleaned)
    return -value if negative else value


def _compact(text: str) -> str:
    return re.sub(r"\s+", "", text)


def envelope_contract(rec: dict[str, Any]) -> tuple[bool, str]:
    ok = rec["ok"]
    data = rec["_data"]
    code = rec["error_code"]
    if ok is True:
        good = data is not None and code is None
        detail = "ok=true, data present, error null" if good else f"ok=true but data={type(data).__name__} code={code}"
    elif ok is False:
        good = data is None and code in ERROR_CODES
        detail = f"ok=false code={code}" + ("" if good else " (data present or unknown code)")
    else:
        good, detail = False, f"ok={ok!r}"
    unknown = [w for w in rec["warning_codes"] if w not in WARNING_CODES]
    if unknown:
        good, detail = False, detail + f"; unknown warning codes {unknown}"
    return good, detail


def expect_ok(rec: dict[str, Any]) -> tuple[bool, str]:
    return rec["ok"] is True, f"ok={rec['ok']} code={rec['error_code']} msg={str(rec['error_message'])[:120]}"


def expect_error(rec: dict[str, Any], code: str) -> tuple[bool, str]:
    return rec["ok"] is False and rec["error_code"] == code, (
        f"ok={rec['ok']} code={rec['error_code']} next_action={str(rec['next_action'])[:120]}"
    )


def expect_warnings(rec: dict[str, Any], codes: list[str], mode: str = "subset") -> tuple[bool, str]:
    have = set(rec["warning_codes"])
    want = set(codes)
    if mode == "none":
        return not have, f"warnings={sorted(have)}"
    if mode == "exact":
        return have == want, f"warnings={sorted(have)} expected={sorted(want)}"
    return want <= have, f"warnings={sorted(have)} expected_superset_of={sorted(want)}"


def row_count_consistency(rec: dict[str, Any], rows_key: str, count_key: str = "returned_row_count") -> tuple[bool, str]:
    data = rec["_data"]
    if not isinstance(data, dict):
        return False, "data is not an object"
    rows = data.get(rows_key)
    declared = data.get(count_key)
    grouped = isinstance(rows, list) and bool(rows) and all(isinstance(r, dict) and "rows" in r for r in rows)
    if grouped and rows_key in ("topics", "events", "groups"):
        actual = sum(len(r.get("rows") or []) for r in rows)
        per_group_ok = all(r.get("row_count") == len(r.get("rows") or []) for r in rows)
        return declared == actual and per_group_ok, f"{count_key}={declared} sum_rows={actual} per_group_ok={per_group_ok}"
    actual = len(rows) if isinstance(rows, list) else None
    return declared == actual, f"{count_key}={declared} len({rows_key})={actual}"


def _balance_sheet_rows(sections: list[dict[str, Any]]) -> list[list[str]]:
    rows: list[list[str]] = []
    for section in sections:
        if section.get("kind") != "balance_sheet":
            continue
        for block in section.get("blocks") or []:
            table = block.get("table")
            if isinstance(table, dict):
                rows.extend([str(c) for c in row] for row in table.get("rows") or [])
    return rows


def official_reconciliation(api_rec: dict[str, Any], sec_rec: dict[str, Any], accounts: list[str]) -> tuple[bool, str]:
    api = api_rec["_data"]
    official: dict[str, str] = {}
    for row in (api.get("accounts") or []) if isinstance(api, dict) else []:
        amount = row.get("thstrm_amount")
        if row.get("sj_div") == "BS" and isinstance(amount, str) and amount.strip():
            official.setdefault(_compact(str(row.get("account_nm", ""))), amount)
    sections = (sec_rec["_data"] or {}).get("sections") or []
    rows = _balance_sheet_rows(sections)
    results = []
    matched = 0
    for account in accounts:
        target = _num(official.get(account))
        if target is None:
            results.append(f"{account}: API 없음")
            continue
        hit = None
        for row in rows:
            if row and _compact(row[0]) == account:
                for cell in row[1:]:
                    value = _num(cell)
                    if value is None:
                        continue
                    for scale in SCALES:
                        if value * scale == target:
                            hit = (cell, scale)
                            break
                    if hit:
                        break
                break
        if hit:
            matched += 1
            results.append(f"{account}: 원문 {hit[0]} x{hit[1]} == API {official[account]}")
        else:
            results.append(f"{account}: 불일치 (API {official[account]}, rows={len(rows)})")
    return matched == len(accounts) and matched > 0, "; ".join(results)


def balance_identity(rec: dict[str, Any]) -> tuple[bool, str]:
    data = rec["_data"]
    amounts: dict[str, Decimal] = {}
    for row in (data.get("accounts") or []) if isinstance(data, dict) else []:
        if row.get("sj_div") == "BS":
            name = _compact(str(row.get("account_nm", "")))
            value = _num(row.get("thstrm_amount"))
            if name in TOTAL_ACCOUNTS and value is not None and name not in amounts:
                amounts[name] = value
    if len(amounts) < 3:
        return False, f"계정 부족 {sorted(amounts)}"
    total = amounts["부채총계"] + amounts["자본총계"]
    return amounts["자산총계"] == total, f"자산 {amounts['자산총계']} vs 부채+자본 {total}"


def page_sum(recs: list[dict[str, Any]]) -> tuple[bool, str]:
    datas = [r["_data"] for r in recs]
    if not datas or any(not isinstance(d, dict) for d in datas):
        return False, "missing page data"
    total = datas[0].get("total_rows")
    returned = sum(d.get("returned_rows", 0) for d in datas)
    indexes = [d.get("page_index") for d in datas]
    fingerprints = {d.get("request_fingerprint") for d in datas}
    last_cursor = datas[-1].get("next_cursor")
    rows_total = sum(len(d.get("rows") or []) for d in datas)
    max_bytes = max(r["bytes"] for r in recs)
    good = (returned == total == rows_total and indexes == list(range(len(datas)))
            and len(fingerprints) == 1 and last_cursor is None and max_bytes < 3_500_000)
    return good, (f"total_rows={total} sum_returned={returned} sum_len_rows={rows_total} page_index={indexes} "
                  f"fingerprints={len(fingerprints)} last_cursor={last_cursor!r} max_bytes={max_bytes}")


def equal_data(recs: list[dict[str, Any]], ignore_keys: list[str] | None = None) -> tuple[bool, str]:
    ignore = set(ignore_keys or [])

    def strip(value: Any) -> Any:
        if isinstance(value, dict):
            return {k: strip(v) for k, v in value.items() if k not in ignore}
        if isinstance(value, list):
            return [strip(v) for v in value]
        return value

    texts = [json.dumps(strip(r["_data"]), ensure_ascii=False, sort_keys=True) for r in recs]
    good = all(t == texts[0] for t in texts)
    return good, "identical" if good else _first_diff(texts[0], texts[1])


def _first_diff(a: str, b: str) -> str:
    for i, (x, y) in enumerate(zip(a, b, strict=False)):
        if x != y:
            return f"differs at char {i}: ...{a[max(0, i - 60):i + 60]}... vs ...{b[max(0, i - 60):i + 60]}..."
    return f"length differs {len(a)} vs {len(b)}"


def xlsx_report(path: str, expect_section_count: int | None = None) -> tuple[bool, str]:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        names = wb.sheetnames
        problems: list[str] = []
        if names[0] != "수집정보":
            problems.append(f"first sheet {names[0]!r} != 수집정보")
        if expect_section_count is not None and len(names) - 1 != expect_section_count:
            problems.append(f"sheets {len(names) - 1} != section_count {expect_section_count}")
        meta = {str(row[0].value): row[1].value for row in wb["수집정보"].iter_rows(min_row=1, max_col=2) if row[0].value is not None}
        for key in ("validation_status", "source_coverage_status"):
            if meta.get(key) != "passed":
                problems.append(f"수집정보 {key}={meta.get(key)!r}")
        formula = comma_numeric = numeric = merged = bad_anchor = indented = note_lists = 0
        for ws in wb.worksheets[1:]:
            for rng in ws.merged_cells.ranges:
                merged += 1
                al = ws.cell(rng.min_row, rng.min_col).alignment
                if al.horizontal != "center" or al.vertical != "center" or not al.wrap_text:
                    bad_anchor += 1
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if cell.data_type == "f":
                        formula += 1
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        numeric += 1
                        if "#,##0" in (cell.number_format or ""):
                            comma_numeric += 1
                    if isinstance(v, str):
                        if v[:1] in INDENT_CHARS:
                            indented += 1
                        if re.fullmatch(r"\d{1,2}(,\d{1,2})+", v):
                            note_lists += 1
        if formula:
            problems.append(f"formula cells {formula}")
        if merged and bad_anchor:
            problems.append(f"merged anchors not centered/wrapped {bad_anchor}/{merged}")
        if numeric and comma_numeric == 0:
            problems.append("no numeric cell carries #,##0 format")
        detail = (f"sheets={len(names)} meta_validation={meta.get('validation_status')} coverage={meta.get('source_coverage_status')} "
                  f"merged={merged} bad_anchor={bad_anchor} numeric={numeric} comma_fmt={comma_numeric} formula={formula} "
                  f"indented_strings={indented} note_number_lists_as_text={note_lists}")
        return not problems, detail + ("" if not problems else " | PROBLEMS: " + "; ".join(problems))
    finally:
        wb.close()


def xlsx_query(path: str, expect_total_rows: int | None = None, expect_domain: str | None = None) -> tuple[bool, str]:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        names = wb.sheetnames
        problems: list[str] = []
        data_sheets = [n for n in names if n == "data" or n.startswith("data_")]
        if not data_sheets or names[-2:] != ["metadata", "warnings"]:
            problems.append(f"sheet layout {names}")
        header = [c.value for c in wb[data_sheets[0]][1]] if data_sheets else []
        rows = sum(max(wb[n].max_row - 1, 0) for n in data_sheets)
        if expect_total_rows is not None and rows != expect_total_rows:
            problems.append(f"data rows {rows} != total_rows {expect_total_rows}")
        meta: dict[str, Any] = {}
        if "metadata" in names:
            meta = {str(r[0].value): r[1].value for r in wb["metadata"].iter_rows(min_row=1, max_col=2) if r[0].value is not None}
        if expect_domain and meta.get("domain") != expect_domain:
            problems.append(f"metadata domain {meta.get('domain')!r}")
        formula = 0
        unprefixed = 0
        for n in names:
            for row in wb[n].iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula += 1
                    if isinstance(cell.value, str) and cell.value[:1] in "=+-@" and not cell.quotePrefix:
                        unprefixed += 1
        if formula:
            problems.append(f"formula cells {formula}")
        if unprefixed:
            problems.append(f"formula-like strings without quotePrefix {unprefixed}")
        detail = f"sheets={names} header={header[:8]} data_rows={rows} metadata_keys={sorted(meta)[:10]} domain={meta.get('domain')} formula={formula}"
        return not problems, detail + ("" if not problems else " | PROBLEMS: " + "; ".join(problems))
    finally:
        wb.close()


def md_utf8_order(path: str, expect_titles: list[str] | None = None) -> tuple[bool, str]:
    raw = Path(path).read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError as error:
        return False, f"not utf-8: {error}"
    positions = [(t, text.find(t)) for t in (expect_titles or [])]
    missing = [t for t, p in positions if p < 0]
    found = [p for _, p in positions if p >= 0]
    order_ok = found == sorted(found)
    headings = len(re.findall(r"^#{1,6} ", text, flags=re.MULTILINE))
    tables = len(re.findall(r"^\|", text, flags=re.MULTILINE))
    return not missing and order_ok, f"chars={len(text)} headings={headings} table_lines={tables} missing={missing} order_ok={order_ok}"


def note_grep(rec: dict[str, Any], pattern: str, excerpt: int = 300) -> tuple[bool, str]:
    regex = re.compile(pattern)
    hits: list[dict[str, Any]] = []
    sections = (rec["_data"] or {}).get("sections") or []
    for section in sections:
        texts: list[str] = []
        for block in section.get("blocks") or []:
            if block.get("text"):
                texts.append(str(block["text"]))
            table = block.get("table")
            if isinstance(table, dict):
                for row in table.get("rows") or []:
                    texts.append(" | ".join(str(c) for c in row))
        joined = "\n".join(texts)
        per_section = 0
        for match in regex.finditer(joined):
            start = max(0, match.start() - excerpt // 3)
            hits.append({"section_id": section.get("section_id"), "title": section.get("title"),
                         "keyword": match.group(0), "excerpt": joined[start:start + excerpt]})
            per_section += 1
            if per_section >= 3:
                break
    by_section = sorted({h["section_id"] for h in hits if h["section_id"]})
    return bool(hits), json.dumps({"hit_count": len(hits), "sections": by_section, "hits": hits[:12]}, ensure_ascii=False)


def sla_p95(recs: list[dict[str, Any]], limit_ms: int) -> tuple[bool, str]:
    times = sorted(r["elapsed_ms"] for r in recs)
    if not times:
        return False, "no records"
    rank = max(1, int(round(0.95 * len(times) + 0.5)))
    p95 = times[min(rank, len(times)) - 1]
    return p95 <= limit_ms, f"n={len(times)} p95={p95}ms max={times[-1]}ms limit={limit_ms}ms"


def elapsed_max(rec: dict[str, Any], limit_ms: int) -> tuple[bool, str]:
    return rec["elapsed_ms"] <= limit_ms, f"elapsed={rec['elapsed_ms']}ms limit={limit_ms}ms"


CHECKS = {
    "envelope_contract": lambda recs, c: envelope_contract(_rec(recs, c["record"])),
    "expect_ok": lambda recs, c: expect_ok(_rec(recs, c["record"])),
    "expect_error": lambda recs, c: expect_error(_rec(recs, c["record"]), c["code"]),
    "expect_warnings": lambda recs, c: expect_warnings(_rec(recs, c["record"]), c.get("codes", []), c.get("mode", "subset")),
    "row_count_consistency": lambda recs, c: row_count_consistency(_rec(recs, c["record"]), c["rows_key"], c.get("count_key", "returned_row_count")),
    "official_reconciliation": lambda recs, c: official_reconciliation(_rec(recs, c["api_record"]), _rec(recs, c["sections_record"]), c.get("accounts", list(TOTAL_ACCOUNTS))),
    "balance_identity": lambda recs, c: balance_identity(_rec(recs, c["record"])),
    "page_sum": lambda recs, c: page_sum([_rec(recs, n) for n in c["records"]]),
    "equal_data": lambda recs, c: equal_data([_rec(recs, n) for n in c["records"]], c.get("ignore_keys")),
    "xlsx_report": lambda recs, c: xlsx_report(c["path"], c.get("expect_section_count")),
    "xlsx_query": lambda recs, c: xlsx_query(c["path"], c.get("expect_total_rows"), c.get("expect_domain")),
    "md_utf8_order": lambda recs, c: md_utf8_order(c["path"], c.get("expect_titles")),
    "note_grep": lambda recs, c: note_grep(_rec(recs, c["record"]), c["pattern"], c.get("excerpt", 300)),
    "sla_p95": lambda recs, c: sla_p95([_rec(recs, n) for n in c["records"]], c["limit_ms"]),
    "elapsed_max": lambda recs, c: elapsed_max(_rec(recs, c["record"]), c["limit_ms"]),
}


def run_checks(run_dir: Path, checks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records = load_records(run_dir)
    results = []
    for check in checks:
        name = check.get("name") or f"{check['type']}:{check.get('record') or check.get('records') or check.get('path')}"
        try:
            passed, detail = CHECKS[check["type"]](records, check)
        except Exception as error:  # noqa: BLE001
            passed, detail = False, f"oracle error {type(error).__name__}: {error}"
        results.append({"name": name, "type": check["type"], "passed": bool(passed), "detail": detail})
        print(f"[{'PASS' if passed else 'FAIL'}] {name}: {str(detail)[:300]}")
    return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", required=True)
    parser.add_argument("--checks", required=True)
    parser.add_argument("--out", default=None)
    ns = parser.parse_args(argv)
    run_dir = Path(ns.run_dir)
    checks = json.loads(Path(ns.checks).read_text(encoding="utf-8"))
    results = run_checks(run_dir, checks)
    out = Path(ns.out) if ns.out else run_dir / "oracles.json"
    existing = json.loads(out.read_text(encoding="utf-8")) if out.exists() else []
    out.write_text(json.dumps(existing + results, ensure_ascii=False, indent=2), encoding="utf-8")
    passed = sum(1 for r in results if r["passed"])
    print(f"--- {passed}/{len(results)} checks passed -> {out} ---")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
