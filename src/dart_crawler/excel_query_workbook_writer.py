from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.excel_native_projection import (
    ExcelProjectionFailure,
    ProjectedExcelCell,
    project_excel_cell,
)
from dart_crawler.excel_page_models import ExcelScalar
from dart_crawler.excel_query_workbook_plan import (
    ExcelWorkbookPlan,
    WorkbookRow,
    metadata_rows,
    warning_rows,
)


@dataclass(frozen=True, slots=True)
class WorkbookWritten:
    pass


type WorkbookWriteOutcome = WorkbookWritten | ExcelProjectionFailure


class _WorksheetCreator(Protocol):
    def create_sheet(self, title: str) -> Worksheet: ...


def write_excel_query_workbook(
    path: Path,
    plan: ExcelWorkbookPlan,
) -> WorkbookWriteOutcome:
    workbook = Workbook()
    try:
        first_sheet = workbook.worksheets[0]
        first_sheet.title = plan.data_sheet_names[0]
        for sheet_index, sheet_name in enumerate(plan.data_sheet_names):
            sheet = (
                first_sheet
                if sheet_index == 0
                else workbook.create_sheet(sheet_name)
            )
            if not isinstance(sheet, Worksheet):
                return ExcelProjectionFailure("excel_cell_projection_failed")
            outcome = _write_data_sheet(sheet, plan, sheet_index)
            if isinstance(outcome, ExcelProjectionFailure):
                return outcome
        metadata_sheet = _create_sheet(workbook, "metadata")
        metadata_outcome = _write_rows(metadata_sheet, metadata_rows(plan))
        if isinstance(metadata_outcome, ExcelProjectionFailure):
            return metadata_outcome
        warning_sheet = _create_sheet(workbook, "warnings")
        warning_outcome = _write_rows(warning_sheet, warning_rows(plan))
        if isinstance(warning_outcome, ExcelProjectionFailure):
            return warning_outcome
        workbook.save(path)
        return WorkbookWritten()
    finally:
        workbook.close()


def _create_sheet(creator: _WorksheetCreator, title: str) -> Worksheet:
    return creator.create_sheet(title)


def _write_data_sheet(
    sheet: Worksheet,
    plan: ExcelWorkbookPlan,
    sheet_index: int,
) -> WorkbookWriteOutcome:
    header_outcome = _write_rows(sheet, (tuple(plan.dataset.columns),))
    if isinstance(header_outcome, ExcelProjectionFailure):
        return header_outcome
    start = sheet_index * plan.options.data_rows_per_sheet
    end = min(start + plan.options.data_rows_per_sheet, plan.dataset.total_rows)
    for source_index in range(start, end):
        row = plan.dataset.rows[source_index]
        outcome = _write_rows(
            sheet,
            (tuple(row[column] for column in plan.dataset.columns),),
            start_row=source_index - start + 2,
        )
        if isinstance(outcome, ExcelProjectionFailure):
            return outcome
    return WorkbookWritten()


def _write_rows(
    sheet: Worksheet,
    rows: tuple[WorkbookRow, ...],
    *,
    start_row: int = 1,
) -> WorkbookWriteOutcome:
    for row_offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            projected = project_excel_cell(value)
            if isinstance(projected, ExcelProjectionFailure):
                return projected
            cell = sheet.cell(start_row + row_offset, column)
            if isinstance(cell, MergedCell):
                return ExcelProjectionFailure("excel_cell_projection_failed")
            _set_cell(cell, projected)
    return WorkbookWritten()


def _set_cell(cell: Cell, projected: ProjectedExcelCell) -> None:
    value: ExcelScalar = projected.value
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
        cell.quotePrefix = True
