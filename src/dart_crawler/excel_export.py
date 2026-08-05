"""Atomic XLSX generation from the normalized document model."""

from __future__ import annotations

import os
import re
import tempfile
import zipfile
from enum import StrEnum, unique
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict

from dart_crawler.document_model import BlockKind, ParsedDocument, SectionKind
from dart_crawler.result import (
    ErrorCode,
    Result,
    WarningCode,
    WarningInfo,
    error_info,
)
from dart_crawler.value_parser import parse_cell_value


@unique
class CollectionStatus(StrEnum):
    """Whether every parsed section was available without image-only content."""

    COMPLETE = "complete"
    PARTIAL = "partial"


class ExportContext(BaseModel):
    """Metadata and normalized document selected for one export."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    company_name: str
    report_date: str | None
    report_title: str
    receipt_date: str
    rcept_no: str
    attachment_id: str
    correction_chain: tuple[str, ...]
    source_url: str
    parser_version: str
    document: ParsedDocument
    comparison_warnings: tuple[WarningInfo, ...] = ()
    collection_warnings: tuple[WarningInfo, ...] = ()


class ExportedFile(BaseModel):
    """Observable result of one XLSX export attempt."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    output_path: Path
    collection_status: CollectionStatus
    reused: bool
    missing_sections: tuple[str, ...]
    source_sha256: str


class ExcelExportService:
    """Write searchable workbook sheets without overwriting different files."""

    def __init__(self, output_dir: Path) -> None:
        self._output_dir = output_dir

    def export(self, context: ExportContext) -> Result[ExportedFile]:
        """Validate core statements and atomically create or reuse a workbook."""
        missing = _missing_core_sections(context.document)
        if missing:
            return Result.failure(
                error_info(
                    ErrorCode.CORE_STATEMENT_MISSING,
                    "핵심 재무제표가 누락되어 엑셀을 만들지 않았습니다.",
                    retryable=False,
                    details={"missing_sections": list(missing)},
                ),
                next_action="재무상태표, 손익·포괄손익, 자본변동표, 현금흐름표가 포함된 원문을 선택하세요.",
            )
        status = (
            CollectionStatus.PARTIAL
            if any(
                block.kind is BlockKind.IMAGE
                for section in context.document.sections
                for block in section.blocks
            )
            else CollectionStatus.COMPLETE
        )
        warnings: list[WarningInfo] = list(
            context.comparison_warnings + context.collection_warnings
        )
        if status is CollectionStatus.PARTIAL:
            warnings.append(
                WarningInfo(
                    code=WarningCode.PARTIAL_COLLECTION,
                    message="이미지 전용 구역을 제외하고 엑셀을 생성했습니다.",
                )
            )
            warnings.append(
                WarningInfo(
                    code=WarningCode.IMAGE_CONTENT_SKIPPED,
                    message="OCR을 수행하지 않아 이미지 내용을 자리표시자로 남겼습니다.",
                )
            )
        output_path = self._output_path(context, status)
        try:
            self._output_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            return Result.failure(
                error_info(
                    ErrorCode.OUTPUT_WRITE_FAILED,
                    "출력 폴더를 만들 수 없습니다.",
                    retryable=False,
                    details={"reason": str(exc)},
                )
            )
        reused = _matching_existing_file(output_path, context)
        if reused:
            warnings.append(
                WarningInfo(
                    code=WarningCode.EXISTING_FILE_REUSED,
                    message="접수번호, 첨부 식별자, 원문 SHA-256이 같은 기존 파일을 재사용했습니다.",
                )
            )
            return Result.success(
                ExportedFile(
                    output_path=output_path,
                    collection_status=status,
                    reused=True,
                    missing_sections=(),
                    source_sha256=context.document.source_sha256,
                ),
                warnings=tuple(warnings),
            )
        output_path = _next_available_path(output_path, context)
        try:
            _write_workbook(output_path, context, status)
        except OSError as exc:
            return Result.failure(
                error_info(
                    ErrorCode.OUTPUT_WRITE_FAILED,
                    "엑셀 파일을 저장할 수 없습니다.",
                    retryable=False,
                    details={"reason": str(exc)},
                )
            )
        return Result.success(
            ExportedFile(
                output_path=output_path,
                collection_status=status,
                reused=False,
                missing_sections=(),
                source_sha256=context.document.source_sha256,
            ),
            warnings=tuple(warnings),
        )

    def _output_path(self, context: ExportContext, status: CollectionStatus) -> Path:
        report_date = context.report_date or context.receipt_date
        partial_suffix = "_부분수집" if status is CollectionStatus.PARTIAL else ""
        filename = (
            f"{context.company_name}_{report_date} {context.report_title}_"
            f"{context.receipt_date}{partial_suffix}.xlsx"
        )
        return self._output_dir / _safe_filename(filename)


def _missing_core_sections(document: ParsedDocument) -> tuple[str, ...]:
    labels = {
        SectionKind.BALANCE_SHEET: "재무상태표",
        SectionKind.INCOME: "손익·포괄손익",
        SectionKind.EQUITY: "자본변동표",
        SectionKind.CASH_FLOW: "현금흐름표",
    }
    missing = []
    for kind, label in labels.items():
        sections = [section for section in document.sections if section.kind is kind]
        if not sections or not any(
            block.kind is BlockKind.TABLE
            for section in sections
            for block in section.blocks
        ):
            missing.append(label)
    return tuple(missing)


def _safe_filename(filename: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", filename).rstrip(" .")
    stem, suffix = os.path.splitext(cleaned)
    if stem.upper() in {"CON", "PRN", "AUX", "NUL"}:
        stem = "_" + stem
    return stem[:220] + suffix


def _next_available_path(path: Path, context: ExportContext) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for number in range(2, 10_000):
        candidate = path.with_name(f"{stem}_{number}{suffix}")
        if not candidate.exists():
            return candidate
    return path.with_name(f"{stem}_{context.rcept_no}{suffix}")


def _matching_existing_file(path: Path, context: ExportContext) -> bool:
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "수집정보" not in workbook.sheetnames:
            workbook.close()
            return False
        values = {
            str(row[0].value): str(row[1].value)
            for row in workbook["수집정보"].iter_rows(min_col=1, max_col=2)
            if row[0].value is not None and row[1].value is not None
        }
        workbook.close()
    except (OSError, KeyError, ValueError, zipfile.BadZipFile):
        return False
    return (
        values.get("rcept_no") == context.rcept_no
        and values.get("attachment_id") == context.attachment_id
        and values.get("source_sha256") == context.document.source_sha256
    )


def _write_workbook(
    path: Path, context: ExportContext, status: CollectionStatus
) -> None:
    workbook = Workbook()
    metadata_sheet = workbook.worksheets[0]
    metadata_sheet.title = "수집정보"
    metadata = _metadata(context, status)
    for row_number, (key, value) in enumerate(metadata.items(), start=1):
        metadata_sheet.cell(row=row_number, column=1, value=key)
        metadata_sheet.cell(row=row_number, column=2, value=parse_cell_value(value))
    used_names = {"수집정보"}
    for section in context.document.sections:
        sheet_name = _sheet_name(section.title, used_names)
        used_names.add(sheet_name)
        sheet = workbook.create_sheet(sheet_name)
        row_number = 1
        for block in section.blocks:
            if block.kind in {BlockKind.HEADING, BlockKind.PARAGRAPH}:
                sheet.cell(row=row_number, column=1, value=parse_cell_value(block.text))
                row_number += 1
            elif block.kind is BlockKind.TABLE:
                table_start_row = row_number
                for source_row in block.rows:
                    for column_number, value in enumerate(source_row, start=1):
                        sheet.cell(
                            row=row_number,
                            column=column_number,
                            value=parse_cell_value(value),
                        )
                    row_number += 1
                for start_row, start_column, end_row, end_column in block.merged_ranges:
                    sheet.merge_cells(
                        start_row=table_start_row + start_row - 1,
                        start_column=start_column,
                        end_row=table_start_row + end_row - 1,
                        end_column=end_column,
                    )
            elif block.kind is BlockKind.IMAGE:
                sheet.cell(row=row_number, column=1, value="[이미지 내용 생략]")
                row_number += 1
    with tempfile.NamedTemporaryFile(
        dir=path.parent,
        prefix=f".{path.stem}.",
        suffix=".tmp",
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _metadata(context: ExportContext, status: CollectionStatus) -> dict[str, str]:
    section_statuses = [
        f"{section.title}:"
        f"{'이미지 전용' if any(block.kind is BlockKind.IMAGE for block in section.blocks) else '수집'}"
        for section in context.document.sections
    ]
    all_warnings = context.comparison_warnings + context.collection_warnings
    warning_codes = [warning.code.value for warning in all_warnings]
    if status is CollectionStatus.PARTIAL:
        warning_codes.extend(
            [
                WarningCode.PARTIAL_COLLECTION.value,
                WarningCode.IMAGE_CONTENT_SKIPPED.value,
            ]
        )
    return {
        "company_name": context.company_name,
        "report_title": context.report_title,
        "report_date": context.report_date or context.receipt_date,
        "receipt_date": context.receipt_date,
        "rcept_no": context.rcept_no,
        "attachment_id": context.attachment_id,
        "correction_chain": ",".join(context.correction_chain),
        "source_url": context.source_url,
        "source_type": context.document.source_type,
        "source_sha256": context.document.source_sha256,
        "parser_version": context.parser_version,
        "collection_status": status.value,
        "discovered_section_count": str(len(context.document.sections)),
        "created_section_count": str(len(context.document.sections)),
        "missing_section_count": "0",
        "missing_sections": "",
        "section_statuses": " | ".join(section_statuses),
        "warning_count": str(len(warning_codes)),
        "warning_codes": ",".join(warning_codes),
    }


def _sheet_name(title: str, used_names: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", title).strip() or "구역"
    candidate = base[:31]
    number = 2
    while candidate in used_names:
        suffix = f"_{number}"
        candidate = base[: 31 - len(suffix)] + suffix
        number += 1
    return candidate
