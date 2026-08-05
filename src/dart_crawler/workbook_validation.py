"""Validate normalized documents and persisted XLSX workbooks."""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol, assert_never

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.document_model import BlockKind, ParsedDocument
from dart_crawler.document_validation import ValidationSummary
from dart_crawler.result import ErrorCode, JsonValue, Result, error_info
from dart_crawler.source_coverage import source_coverage_metadata
from dart_crawler.value_parser import parse_cell_value

METADATA_SHEET = "수집정보"
IMAGE_PLACEHOLDER = "[이미지 내용 생략]"

type CellValue = int | float | str
type CellCoordinate = tuple[int, int]


class ValidationContext(Protocol):
    """Export fields required to validate workbook identity and contents."""

    @property
    def rcept_no(self) -> str: ...

    @property
    def attachment_id(self) -> str: ...

    @property
    def document(self) -> ParsedDocument: ...


@dataclass(frozen=True, slots=True)
class _SheetExpectation:
    cells: dict[CellCoordinate, CellValue]
    merges: frozenset[str]


def validate_workbook(
    path: Path,
    context: ValidationContext,
    *,
    collection_status: str,
    summary: ValidationSummary,
) -> Result[ValidationSummary]:
    """Reopen a saved workbook and compare its cells and merges to the source model."""
    try:
        workbook = load_workbook(path, data_only=False)
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        return _workbook_failure("workbook_open_failed", reason=str(exc))
    try:
        expected_sheets = _sheet_expectations(context.document)
        structure_validation = _validate_sheet_structure(
            workbook, len(expected_sheets), summary
        )
        if not structure_validation.ok:
            return structure_validation
        metadata_validation = _validate_metadata(
            workbook,
            context,
            collection_status=collection_status,
            summary=summary,
        )
        if not metadata_validation.ok:
            return metadata_validation
        for sheet_name, expectation in zip(
            workbook.sheetnames[1:], expected_sheets, strict=True
        ):
            sheet_validation = _validate_sheet(
                workbook[sheet_name], expectation, summary
            )
            if not sheet_validation.ok:
                return sheet_validation
        return Result.success(summary)
    except KeyError as exc:
        return _workbook_failure("sheet_missing", reason=str(exc))
    finally:
        workbook.close()


def _validate_sheet_structure(
    workbook: Workbook,
    expected_content_sheet_count: int,
    summary: ValidationSummary,
) -> Result[ValidationSummary]:
    if workbook.sheetnames[:1] != [METADATA_SHEET]:
        return _workbook_failure("metadata_sheet_not_first")
    expected_count = expected_content_sheet_count + 1
    if len(workbook.sheetnames) != expected_count:
        return _workbook_failure(
            "sheet_count_mismatch",
            expected=expected_count,
            actual=len(workbook.sheetnames),
        )
    return Result.success(summary)


def _validate_metadata(
    workbook: Workbook,
    context: ValidationContext,
    *,
    collection_status: str,
    summary: ValidationSummary,
) -> Result[ValidationSummary]:
    metadata = {
        str(row[0].value): str(row[1].value)
        for row in workbook[METADATA_SHEET].iter_rows(min_col=1, max_col=2)
        if row[0].value is not None and row[1].value is not None
    }
    expected_metadata = {
        "rcept_no": context.rcept_no,
        "attachment_id": context.attachment_id,
        "source_sha256": context.document.source_sha256,
        "collection_status": collection_status,
        "validation_status": "passed",
        "validated_cell_count": str(summary.checked_cell_count),
        "validated_merge_count": str(summary.checked_merge_count),
    }
    if context.document.source_coverage is not None:
        expected_metadata.update(
            source_coverage_metadata(context.document.source_coverage)
        )
    for key, expected_value in expected_metadata.items():
        if metadata.get(key) != expected_value:
            return _workbook_failure(
                "metadata_mismatch",
                key=key,
                expected=expected_value,
                actual=metadata.get(key, ""),
            )
    return Result.success(summary)


def _validate_sheet(
    sheet: Worksheet,
    expectation: _SheetExpectation,
    summary: ValidationSummary,
) -> Result[ValidationSummary]:
    actual_merges = {str(item) for item in sheet.merged_cells.ranges}
    if actual_merges != expectation.merges:
        return _workbook_failure(
            "merge_mismatch",
            sheet=sheet.title,
            expected=",".join(sorted(expectation.merges)),
            actual=",".join(sorted(actual_merges)),
        )
    for (row_number, column_number), expected_value in expectation.cells.items():
        cell = sheet.cell(row=row_number, column=column_number)
        if cell.data_type == "f":
            return _workbook_failure(
                "formula_cell_detected",
                sheet=sheet.title,
                cell=cell.coordinate,
            )
        normalized_expected = _normalized_expected(expected_value)
        if cell.value != normalized_expected:
            return _workbook_failure(
                "cell_value_mismatch",
                sheet=sheet.title,
                cell=cell.coordinate,
                expected=str(normalized_expected),
                actual=str(cell.value),
            )
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        for column_number, cell in enumerate(row, start=1):
            if cell.data_type == "f":
                return _workbook_failure(
                    "formula_cell_detected",
                    sheet=sheet.title,
                    cell=cell.coordinate,
                )
            coordinate = (row_number, column_number)
            if coordinate in expectation.cells or cell.value is None:
                continue
            expected_cell_value = _normalized_expected(
                expectation.cells.get(coordinate)
            )
            if cell.value != expected_cell_value:
                return _workbook_failure(
                    "cell_value_mismatch",
                    sheet=sheet.title,
                    cell=cell.coordinate,
                    expected=str(expected_cell_value),
                    actual=str(cell.value),
                )
    return Result.success(summary)


def _normalized_expected(value: CellValue | None) -> CellValue | None:
    return None if value == "" else value


def _sheet_expectations(document: ParsedDocument) -> tuple[_SheetExpectation, ...]:
    expectations: list[_SheetExpectation] = []
    for section in document.sections:
        cells: dict[CellCoordinate, CellValue] = {}
        merges: set[str] = set()
        row_number = 1
        for block in section.blocks:
            if block.kind in {BlockKind.HEADING, BlockKind.PARAGRAPH}:
                cells[(row_number, 1)] = parse_cell_value(block.text)
                row_number += 1
                continue
            if block.kind is BlockKind.IMAGE:
                cells[(row_number, 1)] = IMAGE_PLACEHOLDER
                row_number += 1
                continue
            if block.kind is not BlockKind.TABLE:
                assert_never(block.kind)
            table_start_row = row_number
            for source_row in block.rows:
                for column_number, value in enumerate(source_row, start=1):
                    cells[(row_number, column_number)] = parse_cell_value(value)
                row_number += 1
            for start_row, start_column, end_row, end_column in block.merged_ranges:
                start = f"{get_column_letter(start_column)}{table_start_row + start_row - 1}"
                end = f"{get_column_letter(end_column)}{table_start_row + end_row - 1}"
                merges.add(f"{start}:{end}")
        expectations.append(
            _SheetExpectation(cells=cells, merges=frozenset(merges))
        )
    return tuple(expectations)


def _workbook_failure(
    issue: str,
    **details: JsonValue,
) -> Result[ValidationSummary]:
    return Result.failure(
        error_info(
            ErrorCode.VALIDATION_FAILED,
            "저장된 엑셀을 원문과 대조하는 검증에 실패했습니다.",
            retryable=False,
            details={"issue": issue, **details},
        ),
        next_action="검증 실패 원인을 확인한 뒤 엑셀 생성을 다시 실행하세요.",
    )
