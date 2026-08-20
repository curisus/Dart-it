from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.document_model import BlockKind, DocumentBlock
from dart_crawler.document_validation import ValidationSummary
from dart_crawler.result import WarningCode
from dart_crawler.source_coverage import source_coverage_metadata
from dart_crawler.value_parser import parse_cell_value, thousands_number_format
from dart_crawler.workbook_layout import apply_workbook_layout
from dart_crawler.workbook_metadata import (
    ATTACHMENT_ID_KEY,
    COLLECTION_STATUS_KEY,
    RCEPT_NO_KEY,
    SOURCE_RCEPT_NO_KEY,
    SOURCE_SHA256_KEY,
    VALIDATED_CELL_COUNT_KEY,
    VALIDATED_MERGE_COUNT_KEY,
    VALIDATION_STATUS_KEY,
    shared_metadata_expectations,
)
from dart_crawler.workbook_validation import IMAGE_PLACEHOLDER, METADATA_SHEET

if TYPE_CHECKING:
    from dart_crawler.excel_export import CollectionStatus, ExportContext


def write_workbook(
    path: Path,
    context: ExportContext,
    status: CollectionStatus,
    validation_summary: ValidationSummary,
) -> None:
    workbook = Workbook()
    metadata_sheet = workbook.worksheets[0]
    metadata_sheet.title = METADATA_SHEET
    metadata = _metadata(context, status, validation_summary)
    for row_number, (key, value) in enumerate(metadata.items(), start=1):
        metadata_sheet.cell(row=row_number, column=1, value=key)
        metadata_sheet.cell(row=row_number, column=2, value=parse_cell_value(value))
    used_names = {METADATA_SHEET}
    for section in context.document.sections:
        sheet_name = _sheet_name(section.title, used_names)
        used_names.add(sheet_name)
        _write_section(workbook.create_sheet(sheet_name), section.blocks)
    apply_workbook_layout(workbook)
    with tempfile.NamedTemporaryFile(
        dir=path.parent,
        prefix=f".{path.stem}.",
        suffix=".tmp",
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _write_section(sheet: Worksheet, blocks: tuple[DocumentBlock, ...]) -> None:
    row_number = 1
    for block in blocks:
        if block.kind in {BlockKind.HEADING, BlockKind.PARAGRAPH}:
            if parse_cell_value(block.text) != "":
                _write_cell(sheet, row_number, 1, block.text)
            row_number += 1
        elif block.kind is BlockKind.TABLE:
            table_start_row = row_number
            for source_row in block.rows:
                for column_number, cell_text in enumerate(source_row, start=1):
                    _write_cell(sheet, row_number, column_number, cell_text)
                row_number += 1
            for start_row, start_column, end_row, end_column in block.merged_ranges:
                sheet.merge_cells(
                    start_row=table_start_row + start_row - 1,
                    start_column=start_column,
                    end_row=table_start_row + end_row - 1,
                    end_column=end_column,
                )
        elif block.kind is BlockKind.IMAGE:
            sheet.cell(row=row_number, column=1, value=IMAGE_PLACEHOLDER)
            row_number += 1


def _write_cell(sheet: Worksheet, row: int, column: int, source_text: str) -> None:
    cell = sheet.cell(row=row, column=column, value=parse_cell_value(source_text))
    number_format = thousands_number_format(source_text)
    if number_format is not None:
        cell.number_format = number_format


def _metadata(
    context: ExportContext,
    status: CollectionStatus,
    validation_summary: ValidationSummary,
) -> dict[str, str]:
    shared_metadata = shared_metadata_expectations(
        context,
        collection_status=status.value,
        summary=validation_summary,
    )
    coverage = context.document.source_coverage
    coverage_metadata = (
        source_coverage_metadata(coverage) if coverage is not None else {}
    )
    section_statuses = [
        f"{section.title}:{_section_status(section.blocks)}"
        for section in context.document.sections
    ]
    all_warnings = context.comparison_warnings + context.collection_warnings
    warning_codes = [warning.code.value for warning in all_warnings]
    if status.value == "partial":
        warning_codes.extend(
            [
                WarningCode.PARTIAL_COLLECTION.value,
                WarningCode.IMAGE_CONTENT_SKIPPED.value,
            ]
        )
    return {
        "company_name": context.company_name,
        "report_title": context.report_title,
        "report_date": context.report_date or context.receipt_date,
        "receipt_date": context.receipt_date,
        RCEPT_NO_KEY: shared_metadata[RCEPT_NO_KEY],
        SOURCE_RCEPT_NO_KEY: shared_metadata[SOURCE_RCEPT_NO_KEY],
        ATTACHMENT_ID_KEY: shared_metadata[ATTACHMENT_ID_KEY],
        "correction_chain": ",".join(context.correction_chain),
        "source_url": context.source_url,
        "source_type": context.document.source_type,
        SOURCE_SHA256_KEY: shared_metadata[SOURCE_SHA256_KEY],
        "parser_version": context.parser_version,
        COLLECTION_STATUS_KEY: shared_metadata[COLLECTION_STATUS_KEY],
        VALIDATION_STATUS_KEY: shared_metadata[VALIDATION_STATUS_KEY],
        VALIDATED_CELL_COUNT_KEY: shared_metadata[VALIDATED_CELL_COUNT_KEY],
        VALIDATED_MERGE_COUNT_KEY: shared_metadata[VALIDATED_MERGE_COUNT_KEY],
        "discovered_section_count": str(len(context.document.sections)),
        "created_section_count": str(len(context.document.sections)),
        "missing_section_count": "0",
        "missing_sections": "",
        "section_statuses": " | ".join(section_statuses),
        "warning_count": str(len(warning_codes)),
        "warning_codes": ",".join(warning_codes),
        **coverage_metadata,
    }


def _section_status(blocks: tuple[DocumentBlock, ...]) -> str:
    image_count = sum(block.kind is BlockKind.IMAGE for block in blocks)
    if image_count == 0:
        return "수집"
    if image_count == len(blocks):
        return "이미지 전용"
    return "부분수집"


def _sheet_name(title: str, used_names: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", title).strip() or "구역"
    candidate = base[:31]
    number = 2
    while candidate in used_names:
        suffix = f"_{number}"
        candidate = base[: 31 - len(suffix)] + suffix
        number += 1
    return candidate
