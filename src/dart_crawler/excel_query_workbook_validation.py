from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.excel_native_projection import (
    ExcelProjectionFailure,
    ProjectedExcelCell,
    project_excel_cell,
    reopened_values_match,
)
from dart_crawler.excel_page_models import ExcelScalar
from dart_crawler.excel_query_workbook_plan import (
    ExcelWorkbookPlan,
    WorkbookRow,
    metadata_rows,
    warning_rows,
)


@dataclass(frozen=True, slots=True)
class WorkbookValidated:
    pass


@dataclass(frozen=True, slots=True)
class WorkbookValidationFailure:
    reason: str = "excel_workbook_validation_failed"


type WorkbookValidationOutcome = WorkbookValidated | WorkbookValidationFailure


def validate_excel_query_workbook(
    path: Path,
    plan: ExcelWorkbookPlan,
) -> WorkbookValidationOutcome:
    if _contains_forbidden_archive_parts(path):
        return WorkbookValidationFailure()
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (KeyError, OSError, SyntaxError, ValueError, zipfile.BadZipFile):
        return WorkbookValidationFailure()
    try:
        if tuple(workbook.sheetnames) != plan.sheet_names:
            return WorkbookValidationFailure()
        if any(
            sheet.sheet_state != "visible" or bool(sheet.merged_cells)
            for sheet in workbook.worksheets
        ):
            return WorkbookValidationFailure()
        for sheet_index, name in enumerate(plan.data_sheet_names):
            if not _validate_data_sheet(workbook[name], plan, sheet_index):
                return WorkbookValidationFailure()
        if not _validate_rows(workbook["metadata"], metadata_rows(plan)):
            return WorkbookValidationFailure()
        if not _validate_rows(workbook["warnings"], warning_rows(plan)):
            return WorkbookValidationFailure()
        return WorkbookValidated()
    except KeyError:
        return WorkbookValidationFailure()
    finally:
        workbook.close()


def _contains_forbidden_archive_parts(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except (OSError, zipfile.BadZipFile):
        return True
    forbidden_prefixes = ("xl/charts/", "xl/tables/")
    return any(
        name.startswith(forbidden_prefixes)
        for name in names
    )


def _validate_data_sheet(
    sheet: Worksheet,
    plan: ExcelWorkbookPlan,
    sheet_index: int,
) -> bool:
    start = sheet_index * plan.options.data_rows_per_sheet
    end = min(start + plan.options.data_rows_per_sheet, plan.dataset.total_rows)
    expected_rows = end - start + 1
    if sheet.max_row != expected_rows:
        return False
    expected_columns = max(1, len(plan.dataset.columns))
    if sheet.max_column != expected_columns:
        return False
    if not _validate_rows(
        sheet,
        (tuple(plan.dataset.columns),),
        validate_dimensions=False,
    ):
        return False
    for source_index in range(start, end):
        source = plan.dataset.rows[source_index]
        expected = tuple(source[column] for column in plan.dataset.columns)
        if not _validate_rows(
            sheet,
            (expected,),
            start_row=source_index - start + 2,
            validate_dimensions=False,
        ):
            return False
    return True


def _validate_rows(
    sheet: Worksheet,
    rows: tuple[WorkbookRow, ...],
    *,
    start_row: int = 1,
    validate_dimensions: bool = True,
) -> bool:
    if validate_dimensions:
        if sheet.max_row != max(1, len(rows)):
            return False
        expected_columns = max((len(row) for row in rows), default=0)
        if sheet.max_column != max(1, expected_columns):
            return False
    for row_offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            projected = project_excel_cell(value)
            if isinstance(projected, ExcelProjectionFailure):
                return False
            if not _cell_matches(
                sheet.cell(start_row + row_offset, column), projected
            ):
                return False
    return True


def _cell_matches(
    cell: Cell | MergedCell,
    projected: ProjectedExcelCell,
) -> bool:
    if isinstance(cell, MergedCell):
        return False
    if cell.data_type == "f":
        return False
    expected = projected.value
    actual_value = cell.value
    if not isinstance(actual_value, (str, int, float, bool, type(None))):
        return False
    actual: ExcelScalar = actual_value
    if isinstance(expected, str):
        expected_type = "inlineStr" if expected == "" else "s"
        if cell.data_type != expected_type or not cell.quotePrefix:
            return False
    return reopened_values_match(expected, actual)
