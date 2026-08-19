"""Column sizing, text wrapping, and merged-cell centering for workbooks.

The writer applies :func:`apply_workbook_layout` right before saving, and the
workbook validator recomputes the same layout from the reloaded file with
:func:`compute_sheet_layout` to confirm the display attributes were written.
Both sides must stay symmetric, mirroring the writer↔validation cell contract.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING
from unicodedata import east_asian_width

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

MAX_TEXT_COLUMN_WIDTH = 60.0
COLUMN_WIDTH_PADDING = 2.0
MAX_COLUMN_WIDTH = 255.0
MERGED_ANCHOR_ALIGNMENT = Alignment(
    horizontal="center", vertical="center", wrap_text=True
)
WRAPPED_TEXT_ALIGNMENT = Alignment(wrap_text=True)

type LayoutCoordinate = tuple[int, int]


@dataclass(frozen=True, slots=True)
class SheetLayout:
    """Display layout derived only from sheet contents and merged ranges."""

    column_widths: dict[str, float]
    merge_anchors: frozenset[LayoutCoordinate]
    wrapped_cells: frozenset[LayoutCoordinate]


def compute_sheet_layout(sheet: Worksheet) -> SheetLayout:
    """Derive column widths, wrap targets, and merge anchors from one sheet."""
    merge_anchors: set[LayoutCoordinate] = set()
    multi_column_anchors: set[LayoutCoordinate] = set()
    for merged_range in sheet.merged_cells.ranges:
        anchor = (int(merged_range.min_row), int(merged_range.min_col))
        merge_anchors.add(anchor)
        if merged_range.max_col > merged_range.min_col:
            multi_column_anchors.add(anchor)
    text_widths: dict[int, float] = {}
    number_widths: dict[int, float] = {}
    wrapped_cells: set[LayoutCoordinate] = set()
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        for column_number, cell in enumerate(row, start=1):
            value = cell.value
            if value is None:
                continue
            coordinate = (row_number, column_number)
            width = _display_width(_display_text(value, str(cell.number_format)))
            is_text = isinstance(value, str)
            if (
                is_text
                and width > MAX_TEXT_COLUMN_WIDTH
                and coordinate not in merge_anchors
            ):
                wrapped_cells.add(coordinate)
            if coordinate in multi_column_anchors:
                continue
            bucket = text_widths if is_text else number_widths
            bucket[column_number] = max(bucket.get(column_number, 0.0), width)
    column_widths: dict[str, float] = {}
    for column_number in sorted(text_widths.keys() | number_widths.keys()):
        text_width = min(text_widths.get(column_number, 0.0), MAX_TEXT_COLUMN_WIDTH)
        content_width = max(text_width, number_widths.get(column_number, 0.0))
        column_widths[get_column_letter(column_number)] = min(
            content_width + COLUMN_WIDTH_PADDING, MAX_COLUMN_WIDTH
        )
    return SheetLayout(
        column_widths=column_widths,
        merge_anchors=frozenset(merge_anchors),
        wrapped_cells=frozenset(wrapped_cells),
    )


def apply_workbook_layout(workbook: Workbook) -> None:
    """Fit columns, wrap long text, and center merged anchors on every sheet."""
    for sheet in workbook.worksheets:
        layout = compute_sheet_layout(sheet)
        for column_letter, width in layout.column_widths.items():
            sheet.column_dimensions[column_letter].width = width
        for row_number, column_number in layout.merge_anchors:
            cell = sheet.cell(row=row_number, column=column_number)
            cell.alignment = MERGED_ANCHOR_ALIGNMENT
        for row_number, column_number in layout.wrapped_cells:
            cell = sheet.cell(row=row_number, column=column_number)
            cell.alignment = WRAPPED_TEXT_ALIGNMENT


def _display_text(value: object, number_format: str) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, bool) or not isinstance(value, int | float):
        return str(value)
    positive_format, _, negative_format = number_format.partition(";")
    has_thousands_grouping = positive_format.startswith("#,##0")
    if not has_thousands_grouping and not positive_format.startswith("0"):
        return str(value)
    decimal_places = len(positive_format.partition(".")[2])
    has_parenthesized_negative = (
        value < 0
        and negative_format.startswith("(")
        and negative_format.endswith(")")
    )
    display_value = abs(value) if has_parenthesized_negative else value
    if decimal_places == 0:
        formatted_value = (
            f"{display_value:,}" if has_thousands_grouping else str(display_value)
        )
    elif has_thousands_grouping:
        formatted_value = f"{display_value:,.{decimal_places}f}"
    else:
        formatted_value = f"{display_value:.{decimal_places}f}"
    if has_parenthesized_negative:
        return f"({formatted_value})"
    return formatted_value


def _display_width(text: str) -> float:
    return sum(2.0 if east_asian_width(char) in {"W", "F"} else 1.0 for char in text)
