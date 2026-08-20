from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from dart_crawler.workbook_metadata import (
    ATTACHMENT_ID_KEY,
    RCEPT_NO_KEY,
    SOURCE_SHA256_KEY,
)
from dart_crawler.workbook_validation import METADATA_SHEET

if TYPE_CHECKING:
    from dart_crawler.excel_export import ExportContext


def safe_filename(filename: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", filename).rstrip(" .")
    stem, suffix = os.path.splitext(cleaned)
    if stem.upper() in {"CON", "PRN", "AUX", "NUL"}:
        stem = "_" + stem
    return stem[:220] + suffix


def next_available_path(path: Path, context: ExportContext) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for number in range(2, 10_000):
        candidate = path.with_name(f"{stem}_{number}{suffix}")
        if not candidate.exists():
            return candidate
    return path.with_name(f"{stem}_{context.rcept_no}{suffix}")


def matching_existing_file(path: Path, context: ExportContext) -> bool:
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if METADATA_SHEET not in workbook.sheetnames:
            workbook.close()
            return False
        values = {
            str(row[0].value): str(row[1].value)
            for row in workbook[METADATA_SHEET].iter_rows(min_col=1, max_col=2)
            if row[0].value is not None and row[1].value is not None
        }
        workbook.close()
    except (OSError, KeyError, ValueError, zipfile.BadZipFile):
        return False
    return (
        values.get(RCEPT_NO_KEY) == context.rcept_no
        and values.get(ATTACHMENT_ID_KEY) == context.attachment_id
        and values.get(SOURCE_SHA256_KEY) == context.document.source_sha256
    )
