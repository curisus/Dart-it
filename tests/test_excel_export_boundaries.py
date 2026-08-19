from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.cell import Cell

from dart_crawler import excel_export
from dart_crawler.document_model import ParsedDocument
from dart_crawler.document_validation import ValidationSummary, validate_document
from dart_crawler.excel_export import ExcelExportService, ExportContext
from dart_crawler.result import ErrorCode, Result, WarningCode, error_info
from dart_crawler.workbook_validation import validate_workbook
from dart_crawler.xml_parser import parse_xml_document


def _document(*, internal_whitespace: bool = False) -> ParsedDocument:
    balance_label = "자산   총계" if internal_whitespace else "자산총계"
    markup = (
        "<document>"
        "<heading>재무상태표</heading>"
        f"<table><tr><td>{balance_label}</td><td>1,000</td></tr></table>"
        "<heading>손익 및 포괄손익계산서</heading>"
        "<table><tr><td>매출</td><td>10</td></tr></table>"
        "<heading>자본변동표</heading>"
        "<table><tr><td>자본</td><td>5</td></tr></table>"
        "<heading>현금흐름표</heading>"
        "<table><tr><td>현금</td><td>6</td></tr></table>"
        "</document>"
    )
    result = parse_xml_document(markup.encode())
    assert result.ok is True
    assert result.data is not None
    return result.data


def _context(document: ParsedDocument) -> ExportContext:
    return ExportContext(
        company_name="Sample Company",
        report_date="2026-03-10",
        report_title="감사보고서",
        receipt_date="20260310",
        rcept_no="20260310002820",
        source_rcept_no="20260310002820",
        attachment_id="opendart:20260310002820:audit.xml",
        correction_chain=("20260310002820",),
        source_url="https://dart.example/report",
        parser_version="0.1.0",
        document=document,
    )


def _remove_source_receipt_metadata(path: Path) -> None:
    workbook = load_workbook(path)
    metadata = workbook["수집정보"]
    row_number = next(
        row[0].row
        for row in metadata.iter_rows(min_col=1, max_col=1)
        if row[0].value == "source_rcept_no"
    )
    metadata.delete_rows(row_number)
    workbook.save(path)
    workbook.close()


def _assert_second_export_is_not_reused(
    tmp_path: Path,
    original_context: ExportContext,
    changed_context: ExportContext,
) -> None:
    service = ExcelExportService(tmp_path)
    original = service.export(original_context)
    assert original.ok is True
    assert original.data is not None

    changed = service.export(changed_context)

    assert changed.ok is True
    assert changed.data is not None
    assert changed.error is None
    assert changed.data.reused is False
    assert changed.data.output_path != original.data.output_path
    assert all(
        warning.code is not WarningCode.EXISTING_FILE_REUSED
        for warning in changed.warnings
    )


def test_export_does_not_reuse_when_rcept_no_alone_differs(tmp_path: Path) -> None:
    original = _context(_document())
    changed = original.model_copy(update={"rcept_no": "20260310002821"})

    _assert_second_export_is_not_reused(tmp_path, original, changed)


def test_export_does_not_reuse_when_attachment_id_alone_differs(
    tmp_path: Path,
) -> None:
    original = _context(_document())
    changed = original.model_copy(
        update={"attachment_id": "opendart:20260310002820:other.xml"}
    )

    _assert_second_export_is_not_reused(tmp_path, original, changed)


def test_export_does_not_reuse_when_source_sha256_alone_differs(
    tmp_path: Path,
) -> None:
    original = _context(_document())
    changed_document = replace(original.document, source_sha256="c" * 64)
    changed = original.model_copy(update={"document": changed_document})

    _assert_second_export_is_not_reused(tmp_path, original, changed)


def test_export_reuses_legacy_workbook_without_source_receipt_metadata(
    tmp_path: Path,
) -> None:
    # Given
    context = _context(_document())
    service = ExcelExportService(tmp_path)
    original = service.export(context)
    assert original.ok is True
    assert original.data is not None
    _remove_source_receipt_metadata(original.data.output_path)

    # When
    result = service.export(context)

    # Then
    assert result.ok is True
    assert result.data is not None
    assert result.data.reused is True
    assert result.data.output_path == original.data.output_path
    assert list(tmp_path.glob("*_2.xlsx")) == []


def test_export_reuses_ambiguous_legacy_workbook_without_source_receipt_metadata(
    tmp_path: Path,
) -> None:
    # Given
    requested_rcept_no = "20260310002820"
    member_name = "20260309001719/audit.xml"
    context = _context(_document()).model_copy(
        update={
            "source_rcept_no": requested_rcept_no,
            "attachment_id": (
                f"opendart:{requested_rcept_no}:{member_name}"
            ),
        }
    )
    service = ExcelExportService(tmp_path)
    original = service.export(context)
    assert original.ok is True
    assert original.data is not None
    _remove_source_receipt_metadata(original.data.output_path)

    # When
    result = service.export(context)

    # Then
    assert result.ok is True
    assert result.data is not None
    assert result.data.reused is True
    assert result.data.output_path == original.data.output_path
    assert list(tmp_path.glob("*_2.xlsx")) == []


def test_post_write_validation_requires_source_receipt_metadata(
    tmp_path: Path,
) -> None:
    # Given
    context = _context(_document())
    exported = ExcelExportService(tmp_path).export(context)
    assert exported.ok is True
    assert exported.data is not None
    _remove_source_receipt_metadata(exported.data.output_path)
    document_validation = validate_document(context.document)
    assert document_validation.ok is True
    assert document_validation.data is not None

    # When
    result = validate_workbook(
        exported.data.output_path,
        context,
        collection_status=exported.data.collection_status.value,
        summary=document_validation.data,
    )

    # Then
    assert result.ok is False
    assert result.error is not None
    assert result.error.details["issue"] == "metadata_mismatch"
    assert result.error.details["key"] == "source_rcept_no"


def test_export_does_not_reuse_when_existing_workbook_revalidation_fails(
    tmp_path: Path,
) -> None:
    context = _context(_document())
    service = ExcelExportService(tmp_path)
    original = service.export(context)
    assert original.ok is True
    assert original.data is not None
    workbook = load_workbook(original.data.output_path)
    workbook["재무상태표"]["A2"] = "변조된 값"
    workbook.save(original.data.output_path)
    workbook.close()

    result = service.export(context)

    assert result.ok is True
    assert result.data is not None
    assert result.error is None
    assert result.data.reused is False
    assert result.data.output_path != original.data.output_path
    assert all(
        warning.code is not WarningCode.EXISTING_FILE_REUSED
        for warning in result.warnings
    )


def test_export_removes_workbook_when_post_write_validation_fails(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    def reject_workbook(
        path: Path,
        context: ExportContext,
        *,
        collection_status: str,
        summary: ValidationSummary,
        allow_legacy_source_receipt_omission: bool,
    ) -> Result[ValidationSummary]:
        assert path.suffix == ".xlsx"
        assert context.rcept_no
        assert collection_status == "complete"
        assert summary.checked_cell_count > 0
        assert allow_legacy_source_receipt_omission is False
        return Result[ValidationSummary].failure(
            error_info(
                ErrorCode.VALIDATION_FAILED,
                "저장된 엑셀 검증에 실패했습니다.",
                retryable=False,
            )
        )

    monkeypatch.setattr(excel_export, "validate_workbook", reject_workbook)

    result = ExcelExportService(tmp_path).export(_context(_document()))

    assert result.ok is False
    assert result.data is None
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert list(tmp_path.glob("*.xlsx")) == []


def test_internal_whitespace_collapse_matches_writer_and_validator(
    tmp_path: Path,
) -> None:
    document = _document(internal_whitespace=True)
    context = _context(document)

    result = ExcelExportService(tmp_path).export(context)

    assert result.ok is True
    assert result.data is not None
    workbook = load_workbook(result.data.output_path)
    balance_sheet = next(
        sheet for sheet in workbook.worksheets if sheet.title == "재무상태표"
    )
    assert balance_sheet.cell(row=2, column=1).value == "자산 총계"
    workbook.close()
    summary = validate_document(document)
    assert summary.ok is True
    assert summary.data is not None
    validation = validate_workbook(
        result.data.output_path,
        context,
        collection_status="complete",
        summary=summary.data,
    )
    assert validation.ok is True
    assert validation.data is not None
    assert validation.error is None


@pytest.mark.parametrize("formula", ["+SUM(B2:B3)", "@SUM(B2:B3)"])
def test_validate_workbook_rejects_non_equals_formula_cells(
    formula: str,
    tmp_path: Path,
) -> None:
    document = _document()
    context = _context(document)
    exported = ExcelExportService(tmp_path).export(context)
    assert exported.ok is True
    assert exported.data is not None
    workbook = load_workbook(exported.data.output_path)
    balance_sheet = next(
        sheet for sheet in workbook.worksheets if sheet.title == "재무상태표"
    )
    formula_cell = balance_sheet.cell(row=2, column=1)
    assert isinstance(formula_cell, Cell)
    formula_cell.value = formula
    formula_cell.data_type = "f"
    workbook.save(exported.data.output_path)
    workbook.close()
    summary = validate_document(document)
    assert summary.ok is True
    assert summary.data is not None

    result = validate_workbook(
        exported.data.output_path,
        context,
        collection_status="complete",
        summary=summary.data,
    )

    assert result.ok is False
    assert result.data is None
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert result.error.details["issue"] == "formula_cell_detected"
