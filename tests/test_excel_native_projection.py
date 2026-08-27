import math
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.compat import safe_string
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.excel_export_result import Result as ExcelResult
from dart_crawler.excel_native_projection import (
    ExcelProjectionFailure,
    ProjectedExcelCell,
    project_excel_cell,
)
from dart_crawler.excel_page_models import ExcelRow, ExcelScalar
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.excel_query_workbook_plan import ExcelWorkbookOptions
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.result import ErrorCode, WarningCode, WarningInfo
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


@pytest.mark.parametrize(
    ("source", "stored_type"),
    [
        (True, bool),
        (False, bool),
        (2**53, int),
        (-(2**53), int),
        (2**53 + 1, float),
        (-(2**53 + 1), float),
        (1.25, float),
        (2.0, float),
        (-0.0, float),
        (None, type(None)),
    ],
)
def test_scalar_projection_preserves_required_native_type(
    source: ExcelScalar,
    stored_type: type[ExcelScalar],
) -> None:
    projected = project_excel_cell(source)

    assert isinstance(projected, ProjectedExcelCell)
    assert type(projected.value) is stored_type
    if isinstance(source, int) and not isinstance(source, bool) and abs(source) > 2**53:
        assert projected.value == float(source)
    elif isinstance(source, float) and not source:
        assert isinstance(projected.value, float)
        assert math.copysign(1.0, projected.value) == math.copysign(1.0, source)
    else:
        assert projected.value == source


def test_huge_integer_projection_is_a_typed_failure() -> None:
    projected = project_excel_cell(10**10_000)

    assert projected == ExcelProjectionFailure("excel_cell_projection_failed")


def test_string_projection_truncates_before_illegal_xml_validation() -> None:
    boundary = "가" * 32_767

    assert project_excel_cell("") == ProjectedExcelCell("")
    assert project_excel_cell(boundary) == ProjectedExcelCell(boundary)
    assert project_excel_cell(boundary + "끝") == ProjectedExcelCell(boundary)
    assert project_excel_cell("안전\x00") == ExcelProjectionFailure(
        "excel_cell_projection_failed"
    )
    assert project_excel_cell(boundary + "\x00") == ProjectedExcelCell(boundary)


def test_native_scalar_workbook_reopens_with_expected_values_and_types(
    tmp_path: Path,
) -> None:
    boundary = "나" * 32_767
    columns = (
        "boolean",
        "positive_limit",
        "negative_limit",
        "large_integer",
        "fractional_float",
        "integral_float",
        "negative_zero",
        "empty_string",
        "truncated_string",
        "unicode_string",
        "blank",
    )
    source_row: ExcelRow = {
        "boolean": True,
        "positive_limit": 2**53,
        "negative_limit": -(2**53),
        "large_integer": 2**53 + 1,
        "fractional_float": 1.25,
        "integral_float": 2.0,
        "negative_zero": -0.0,
        "empty_string": "",
        "truncated_string": boundary + "버려짐\x00",
        "unicode_string": "한글🙂漢字",
        "blank": None,
    }
    warning = WarningInfo(
        code=WarningCode.FALLBACK_SOURCE_USED,
        message="중첩 상세",
        details={"z": [2, {"나": True}], "a": "값"},
    )
    dataset = make_normalized_dataset(columns, (source_row,), warnings=(warning,))
    clock = RecordingClock(datetime(2026, 3, 4, 5, 6, 7, 8, tzinfo=UTC))

    result = publish_excel_dataset(
        dataset,
        tmp_path / "output",
        clock=clock,
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is True
    exported = result.data
    assert exported is not None
    assert clock.calls == 1
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        data = workbook["data"]
        assert isinstance(data, Worksheet)
        observed = {data.cell(1, index).value: data.cell(2, index) for index in range(1, 12)}
        assert type(observed["boolean"].value) is bool
        assert type(observed["positive_limit"].value) is int
        assert type(observed["negative_limit"].value) is int
        source_large_integer = source_row["large_integer"]
        assert isinstance(source_large_integer, int)
        assert not isinstance(source_large_integer, bool)
        assert safe_string(observed["large_integer"].value) == safe_string(
            float(source_large_integer)
        )
        assert observed["fractional_float"].value == 1.25
        assert observed["integral_float"].value == 2
        assert observed["negative_zero"].value == 0
        assert observed["empty_string"].value is None
        assert observed["truncated_string"].value == boundary
        assert observed["unicode_string"].value == "한글🙂漢字"
        assert observed["blank"].value is None
        warning_sheet = workbook["warnings"]
        assert isinstance(warning_sheet, Worksheet)
        assert warning_sheet.cell(2, 4).value == (
            '{"a":"값","z":[2,{"나":true}]}'
        )
        _assert_all_strings_are_literals(workbook.sheetnames, workbook)
    finally:
        workbook.close()


def test_huge_integer_failure_leaves_no_final_workbook(tmp_path: Path) -> None:
    dataset = make_normalized_dataset(("huge",), ({"huge": 10**10_000},))
    output_root = tmp_path / "output"

    result: ExcelResult[ExcelExportResult] = publish_excel_dataset(
        dataset,
        output_root,
        clock=RecordingClock(datetime(2026, 1, 1, tzinfo=UTC)),
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is False
    assert result.data is None
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert result.error.retryable is False
    assert result.error.details == {"reason": "excel_cell_projection_failed"}
    assert list(output_root.iterdir()) == []


def _assert_all_strings_are_literals(
    sheet_names: list[str],
    workbook: Workbook,
) -> None:
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert cell.data_type in {"s", "inlineStr"}
                    assert cell.quotePrefix is True
