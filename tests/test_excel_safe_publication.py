from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.excel_export_result import Result as ExcelResult
from dart_crawler.excel_page_models import ExcelRow
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.excel_query_workbook_plan import ExcelWorkbookOptions
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.normalized_excel_models import NormalizedExcelDataset
from tests.excel_publication_test_support import (
    FinalRacePublicationOps,
    LinkRacePublicationOps,
    RecordingPublicationOps,
)
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


def test_existing_final_and_lock_are_preserved_with_suffixes(tmp_path: Path) -> None:
    output_root = tmp_path / "output"
    output_root.mkdir()
    first_final = output_root / "search_companies.xlsx"
    _ = first_final.write_bytes(b"existing-final")

    first_result = _publish(_dataset(), output_root)

    assert first_result.ok is True
    assert first_result.data is not None
    assert first_result.data.filename == "search_companies_2.xlsx"
    assert first_final.read_bytes() == b"existing-final"

    second_lock = output_root / "search_companies_3.xlsx.lock"
    _ = second_lock.write_bytes(b"existing-lock")
    second_result = _publish(_dataset(), output_root)

    assert second_result.ok is True
    assert second_result.data is not None
    assert second_result.data.filename == "search_companies_4.xlsx"
    assert second_lock.read_bytes() == b"existing-lock"


def test_two_concurrent_exports_publish_distinct_complete_files(
    tmp_path: Path,
) -> None:
    output_root = tmp_path / "output"
    dataset = _dataset()
    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [executor.submit(_publish, dataset, output_root) for _ in range(2)]
        results = [future.result() for future in futures]

    assert all(result.ok for result in results)
    exports = [result.data for result in results]
    assert all(exported is not None for exported in exports)
    filenames = {exported.filename for exported in exports if exported is not None}
    assert filenames == {"search_companies.xlsx", "search_companies_2.xlsx"}
    assert list(output_root.glob("*.lock")) == []
    assert list(output_root.glob(".*.xlsx")) == []
    for exported in exports:
        assert exported is not None
        _assert_complete_workbook(Path(exported.absolute_path))


def test_temp_is_unpredictable_same_directory_and_published_by_link(
    tmp_path: Path,
) -> None:
    output_root = tmp_path / "output"
    operations = RecordingPublicationOps()

    result = _publish(_dataset(), output_root, file_ops=operations)

    assert result.ok is True
    assert len(operations.temp_paths) == 1
    assert len(operations.link_pairs) == 1
    temp_path = operations.temp_paths[0]
    source, destination = operations.link_pairs[0]
    assert source == temp_path
    assert source.parent == destination.parent == output_root.resolve()
    assert source != destination
    assert source.name.startswith(".search_companies.")
    assert source.suffix == ".xlsx"
    assert not source.exists()
    assert destination.exists()


def test_final_appearing_after_lock_is_never_overwritten_or_removed(
    tmp_path: Path,
) -> None:
    operations = FinalRacePublicationOps()

    result = _publish(_dataset(), tmp_path / "output", file_ops=operations)

    assert result.ok is True
    assert result.data is not None
    assert result.data.filename == "search_companies_2.xlsx"
    assert operations.raced_final is not None
    assert operations.raced_final.read_bytes() == b"other-owner"


def test_final_link_race_advances_suffix_without_cross_cleanup(
    tmp_path: Path,
) -> None:
    operations = LinkRacePublicationOps()

    result = _publish(_dataset(), tmp_path / "output", file_ops=operations)

    assert result.ok is True
    assert result.data is not None
    assert result.data.filename == "search_companies_2.xlsx"
    assert operations.raced_final is not None
    _assert_complete_workbook(operations.raced_final)
    _assert_complete_workbook(Path(result.data.absolute_path))


def _dataset() -> NormalizedExcelDataset:
    row: ExcelRow = {"sequence": 1, "label": "완전한 행"}
    return make_normalized_dataset(("sequence", "label"), (row,))


def _publish(
    dataset: NormalizedExcelDataset,
    output_root: Path,
    *,
    file_ops: RecordingPublicationOps | None = None,
) -> ExcelResult[ExcelExportResult]:
    clock = RecordingClock(datetime(2026, 6, 7, 8, 9, 10, 11, tzinfo=UTC))
    if file_ops is None:
        return publish_excel_dataset(
            dataset,
            output_root,
            clock=clock,
            options=ExcelWorkbookOptions(),
        )
    return publish_excel_dataset(
        dataset,
        output_root,
        clock=clock,
        options=ExcelWorkbookOptions(),
        file_ops=file_ops,
    )


def _assert_complete_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == ["data", "metadata", "warnings"]
        data = workbook["data"]
        assert isinstance(data, Worksheet)
        assert data.max_row == 2
        assert data.cell(2, 1).value == 1
        assert data.cell(2, 2).value == "완전한 행"
    finally:
        workbook.close()
