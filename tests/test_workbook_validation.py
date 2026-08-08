from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from dart_crawler.document_model import (
    BlockKind,
    DocumentBlock,
    DocumentSection,
    ParsedDocument,
    SectionKind,
    SourceCoverage,
)
from dart_crawler.document_validation import validate_document
from dart_crawler.html_parser import parse_html_document
from dart_crawler.workbook_layout import apply_workbook_layout
from dart_crawler.workbook_validation import validate_workbook


@dataclass(frozen=True, slots=True)
class _ValidationContext:
    rcept_no: str
    attachment_id: str
    document: ParsedDocument


def _complete_coverage() -> SourceCoverage:
    return SourceCoverage(
        source_text_token_count=0,
        captured_text_token_count=0,
        source_text_sha256="a" * 64,
        captured_text_sha256="a" * 64,
        source_table_count=1,
        captured_table_count=1,
        source_cell_count=3,
        captured_cell_count=3,
        source_image_count=0,
        captured_image_count=0,
    )


def test_validate_document_rejects_non_rectangular_table() -> None:
    document = ParsedDocument(
        sections=(
            DocumentSection(
                title="malformed",
                kind=SectionKind.OTHER,
                blocks=(
                    DocumentBlock(
                        BlockKind.TABLE,
                        rows=(("A",), ("B", "C")),
                    ),
                ),
            ),
        ),
        source_sha256="a" * 64,
        source_type="xml",
        source_coverage=_complete_coverage(),
    )

    result = validate_document(document)

    assert result.ok is False
    assert result.error is not None
    assert result.error.code.value == "VALIDATION_FAILED"


def test_validate_document_rejects_incomplete_source_coverage() -> None:
    coverage = SourceCoverage(
        source_text_token_count=2,
        captured_text_token_count=1,
        source_text_sha256="a" * 64,
        captured_text_sha256="b" * 64,
        source_table_count=1,
        captured_table_count=1,
        source_cell_count=2,
        captured_cell_count=1,
        source_image_count=0,
        captured_image_count=0,
    )
    document = ParsedDocument(
        sections=(
            DocumentSection(
                title="본문",
                kind=SectionKind.OTHER,
                blocks=(DocumentBlock(BlockKind.PARAGRAPH, text="captured"),),
            ),
        ),
        source_sha256="a" * 64,
        source_type="xml",
        source_coverage=coverage,
    )

    result = validate_document(document)

    assert result.ok is False
    assert result.error is not None
    assert result.error.code.value == "VALIDATION_FAILED"
    assert result.error.details["issue"] == "incomplete_source_coverage"


def test_validate_document_rejects_fingerprint_mismatch_with_equal_counts() -> None:
    coverage = SourceCoverage(
        source_text_token_count=1,
        captured_text_token_count=1,
        source_text_sha256="a" * 64,
        captured_text_sha256="b" * 64,
        source_table_count=0,
        captured_table_count=0,
        source_cell_count=0,
        captured_cell_count=0,
        source_image_count=0,
        captured_image_count=0,
    )
    document = ParsedDocument(
        sections=(
            DocumentSection(
                title="본문",
                kind=SectionKind.OTHER,
                blocks=(DocumentBlock(BlockKind.PARAGRAPH, text="captured"),),
            ),
        ),
        source_sha256="a" * 64,
        source_type="xml",
        source_coverage=coverage,
    )

    result = validate_document(document)

    assert result.ok is False
    assert result.error is not None
    assert result.error.details["issue"] == "incomplete_source_coverage"


def test_validate_workbook_rejects_formula_in_metadata_sheet(tmp_path: Path) -> None:
    parsed = parse_html_document(b"<heading>Header</heading><p>Tail</p>")
    assert parsed.ok is True
    assert parsed.data is not None
    validation = validate_document(parsed.data)
    assert validation.ok is True
    assert validation.data is not None
    context = _ValidationContext(
        rcept_no="20260310000001",
        attachment_id="opendart:20260310000001:audit.xml",
        document=parsed.data,
    )
    path = tmp_path / "formula_metadata.xlsx"
    workbook = Workbook()
    metadata = workbook.worksheets[0]
    metadata.title = "수집정보"
    metadata.append(("rcept_no", context.rcept_no))
    metadata.append(("evil", "=1+1"))
    content = workbook.create_sheet("Header")
    content["A1"] = "Header"
    content["A2"] = "Tail"
    workbook.save(path)
    workbook.close()

    result = validate_workbook(
        path,
        context,
        collection_status="complete",
        summary=validation.data,
    )

    assert result.ok is False
    assert result.error is not None
    assert result.error.details["issue"] == "formula_cell_detected"
    assert result.error.details["sheet"] == "수집정보"


def test_validate_workbook_rejects_missing_final_expected_cell(tmp_path: Path) -> None:
    parsed = parse_html_document(b"<heading>Header</heading><p>Tail</p>")
    assert parsed.ok is True
    assert parsed.data is not None
    validation = validate_document(parsed.data)
    assert validation.ok is True
    assert validation.data is not None
    context = _ValidationContext(
        rcept_no="20260310000001",
        attachment_id="opendart:20260310000001:audit.xml",
        document=parsed.data,
    )
    path = tmp_path / "truncated.xlsx"
    workbook = Workbook()
    metadata = workbook.worksheets[0]
    metadata.title = "수집정보"
    metadata_rows = {
        "rcept_no": context.rcept_no,
        "attachment_id": context.attachment_id,
        "source_sha256": context.document.source_sha256,
        "collection_status": "complete",
        "validation_status": "passed",
        "validated_cell_count": str(validation.data.checked_cell_count),
        "validated_merge_count": str(validation.data.checked_merge_count),
    }
    coverage = context.document.source_coverage
    assert coverage is not None
    metadata_rows.update(
        {
            "source_coverage_status": "passed",
            "source_text_token_count": str(coverage.source_text_token_count),
            "captured_text_token_count": str(coverage.captured_text_token_count),
            "source_table_count": str(coverage.source_table_count),
            "captured_table_count": str(coverage.captured_table_count),
            "source_cell_count": str(coverage.source_cell_count),
            "captured_cell_count": str(coverage.captured_cell_count),
            "source_image_count": str(coverage.source_image_count),
            "captured_image_count": str(coverage.captured_image_count),
        }
    )
    for key, value in metadata_rows.items():
        metadata.append((key, value))
    content = workbook.create_sheet("Header")
    content["A1"] = "Header"
    apply_workbook_layout(workbook)
    workbook.save(path)
    workbook.close()

    result = validate_workbook(
        path,
        context,
        collection_status="complete",
        summary=validation.data,
    )

    assert result.ok is False
    assert result.error is not None
    assert result.error.details["issue"] == "cell_value_mismatch"
    assert result.error.details["cell"] == "A2"
