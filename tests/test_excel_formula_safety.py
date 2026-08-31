from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from dart_crawler.excel_page_models import ExcelRow
from dart_crawler.excel_query_workbook_plan import ExcelWorkbookOptions
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.result import WarningCode, WarningInfo
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


def test_formula_prefixes_remain_literal_on_every_string_bearing_sheet(
    tmp_path: Path,
) -> None:
    prefixes = ("=", "+", "-", "@")
    columns = tuple(f"{prefix}header" for prefix in prefixes)
    source_row: ExcelRow = {
        column: f"{prefix}1+1"
        for column, prefix in zip(columns, prefixes, strict=True)
    }
    warnings = tuple(
        WarningInfo(
            code=WarningCode.FALLBACK_SOURCE_USED,
            message=f"{prefix}warning",
            details={"formula_like": f"{prefix}detail"},
        )
        for prefix in prefixes
    )
    dataset = make_normalized_dataset(columns, (source_row,), warnings=warnings)

    result = publish_excel_dataset(
        dataset,
        tmp_path / "output",
        clock=RecordingClock(datetime(2026, 4, 5, tzinfo=UTC)),
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is True
    exported = result.data
    assert exported is not None
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        data = workbook["data"]
        assert tuple(cell.value for cell in data[1]) == columns
        assert tuple(cell.value for cell in data[2]) == tuple(source_row.values())
        assert tuple(
            workbook["warnings"].cell(row, 3).value for row in range(2, 6)
        ) == tuple(f"{prefix}warning" for prefix in prefixes)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f"
                    if isinstance(cell.value, str):
                        assert cell.data_type in {"s", "inlineStr"}
                        assert cell.quotePrefix is True
    finally:
        workbook.close()
