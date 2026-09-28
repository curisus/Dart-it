import importlib
import json
import os
import sys
import threading
from collections.abc import Callable, Iterator
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import ModuleType
from typing import Final

import openpyxl
import pytest
import scripts.perf.perf_stats as perf_stats
import scripts.perf.stage_probes as stage_probes
import scripts.perf.xlsx_invariance as xlsx_invariance
from openpyxl.styles import Alignment, Font

_PERF_DIR: Final = Path(__file__).resolve().parents[1] / "scripts" / "perf"
_FAKE_MODULE: Final = "perf_probe_fake"
_FAKE_USER_MODULE: Final = "perf_probe_fake.user"

type Rows = list[list[object]]


def _workbook(
    path: Path, sheets: dict[str, Rows], *, merged: str | None = None
) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    if merged is not None:
        workbook.worksheets[0].merge_cells(merged)
    workbook.save(path)
    return path


def _metadata_rows(generated_at: str) -> Rows:
    return [
        ["key", "value"],
        ["domain", "get_financial_statements"],
        ["generated_at_utc", generated_at],
    ]


# ---------------------------------------------------------------- perf_stats


def test_median_handles_odd_and_even_sample_counts() -> None:
    assert perf_stats.median([3.0, 1.0, 2.0]) == 2.0
    assert perf_stats.median([4.0, 1.0, 3.0, 2.0]) == 2.5


def test_p90_interpolates_linearly_between_neighbours() -> None:
    samples = [float(value) for value in range(1, 11)]
    assert perf_stats.p90(samples) == pytest.approx(9.1)
    assert perf_stats.p90([5.0]) == 5.0


def test_summarize_reports_count_median_p90_and_range() -> None:
    summary = perf_stats.summarize([10.0, 30.0, 20.0])
    assert summary == perf_stats.LatencySummary(
        n=3, median_ms=20.0, p90_ms=28.0, min_ms=10.0, max_ms=30.0
    )


def test_statistics_reject_empty_or_non_finite_samples() -> None:
    with pytest.raises(ValueError, match="at least one"):
        perf_stats.median([])
    with pytest.raises(ValueError, match="finite"):
        perf_stats.p90([1.0, float("nan")])


def test_reduction_ratio_is_the_fraction_of_the_base_median_removed() -> None:
    assert perf_stats.reduction_ratio(1000.0, 400.0) == pytest.approx(0.6)
    assert perf_stats.reduction_ratio(1000.0, 1200.0) == pytest.approx(-0.2)
    with pytest.raises(ValueError, match="positive"):
        perf_stats.reduction_ratio(0.0, 1.0)


def _samples(
    corp_code: str,
    base_ms: tuple[float, ...],
    new_ms: tuple[float, ...],
    *,
    new_invalid: int = 0,
) -> perf_stats.CompanySamples:
    return perf_stats.CompanySamples(
        corp_code=corp_code, base_ms=base_ms, new_ms=new_ms, new_invalid=new_invalid
    )


def _assess(
    flow: str,
    companies: list[perf_stats.CompanySamples],
    *,
    representative: bool,
) -> perf_stats.FlowAssessment:
    return perf_stats.assess_flow(
        flow, companies, in_base=True, in_new=True, representative=representative
    )


def test_flow_reduction_is_the_median_of_each_companys_own_reduction() -> None:
    flow = _assess(
        "L-A",
        [
            _samples("X", (100.0, 120.0, 110.0), (50.0, 60.0, 55.0)),
            _samples("Y", (1000.0, 1000.0), (400.0, 400.0)),
        ],
        representative=True,
    )
    assert [company.reduction for company in flow.companies] == [
        pytest.approx(0.5),
        pytest.approx(0.6),
    ]
    assert flow.reduction == pytest.approx(0.55)
    assert flow.status is perf_stats.FlowStatus.COMPLETE
    assert flow.passed is True
    # the pooled median is informational only
    assert flow.pooled_base is not None
    assert flow.pooled_base.median_ms == 120.0


def test_flow_status_is_complete_incomplete_or_missing() -> None:
    complete = [_samples("X", (1.0, 2.0), (1.0, 2.0))]
    fewer_reps = [_samples("X", (1.0, 2.0), (1.0,))]
    invalid_rep = [_samples("X", (1.0,), (1.0,), new_invalid=1)]
    all_invalid = [perf_stats.CompanySamples("X", (), (), 3, 3)]
    status = perf_stats.flow_status
    assert status(complete, in_base=True, in_new=True) is perf_stats.FlowStatus.COMPLETE
    assert (
        status(fewer_reps, in_base=True, in_new=True)
        is perf_stats.FlowStatus.INCOMPLETE
    )
    assert (
        status(invalid_rep, in_base=True, in_new=True)
        is perf_stats.FlowStatus.INCOMPLETE
    )
    assert status(complete, in_base=True, in_new=False) is perf_stats.FlowStatus.MISSING
    assert (
        status(all_invalid, in_base=True, in_new=True) is perf_stats.FlowStatus.MISSING
    )


def _oracle(
    flows: list[perf_stats.FlowAssessment],
    *,
    invariance: perf_stats.Verdict = perf_stats.Verdict.PASS,
    allow_incomplete: bool = False,
) -> perf_stats.OracleVerdict:
    return perf_stats.oracle_verdict(
        flows, {"L-A"}, invariance=invariance, allow_incomplete=allow_incomplete
    )


def test_oracle_passes_when_representative_halves_and_secondary_holds() -> None:
    verdict = _oracle(
        [
            _assess("L-A", [_samples("X", (1000.0,), (500.0,))], representative=True),
            _assess("L-C", [_samples("X", (100.0,), (110.0,))], representative=False),
        ]
    )
    assert verdict.latency is perf_stats.Verdict.PASS
    assert verdict.overall is perf_stats.Verdict.PASS


def test_oracle_fails_on_small_representative_gain_or_secondary_regression() -> None:
    slow_gain = _oracle(
        [_assess("L-A", [_samples("X", (1000.0,), (510.0,))], representative=True)]
    )
    regression = _oracle(
        [
            _assess("L-A", [_samples("X", (1000.0,), (300.0,))], representative=True),
            _assess("L-C", [_samples("X", (100.0,), (111.0,))], representative=False),
        ]
    )
    assert slow_gain.latency is perf_stats.Verdict.FAIL
    assert regression.latency is perf_stats.Verdict.FAIL
    assert [flow.passed for flow in regression.flows] == [True, False]


def test_oracle_fails_on_missing_representative_and_on_data_change() -> None:
    no_representative = _oracle(
        [_assess("L-C", [_samples("X", (100.0,), (90.0,))], representative=False)]
    )
    changed_data = _oracle(
        [_assess("L-A", [_samples("X", (1000.0,), (100.0,))], representative=True)],
        invariance=perf_stats.Verdict.FAIL,
    )
    assert no_representative.missing_representative == ("L-A",)
    assert no_representative.overall is perf_stats.Verdict.FAIL
    assert changed_data.invariance is perf_stats.Verdict.FAIL
    assert changed_data.overall is perf_stats.Verdict.FAIL


def test_incomplete_flow_fails_unless_allowed_and_is_then_never_plain_pass() -> None:
    # W3 S1 with a real 60% speedup: company Y lost two of its three new reps.
    flows = [
        _assess(
            "L-A",
            [
                _samples("X", (100.0, 100.0, 100.0), (40.0, 40.0, 40.0)),
                _samples("Y", (1000.0, 1000.0, 1000.0), (400.0,), new_invalid=2),
            ],
            representative=True,
        )
    ]
    strict = _oracle(flows)
    allowed = _oracle(flows, allow_incomplete=True)
    assert flows[0].status is perf_stats.FlowStatus.INCOMPLETE
    assert strict.overall is perf_stats.Verdict.FAIL
    assert allowed.overall is perf_stats.Verdict.PASS_WITH_INCOMPLETE


def test_invariance_verdict_fails_on_difference_and_gates_missing_pairs() -> None:
    verdict = perf_stats.invariance_verdict
    assert verdict(["identical", "identical"]) is perf_stats.Verdict.PASS
    assert verdict(["identical", "different"]) is perf_stats.Verdict.FAIL
    assert verdict(["unstable_new"], allow_incomplete=True) is perf_stats.Verdict.FAIL
    assert verdict(["identical", "missing_new"]) is perf_stats.Verdict.FAIL
    assert (
        verdict(["identical", "missing_new"], allow_incomplete=True)
        is perf_stats.Verdict.PASS_WITH_INCOMPLETE
    )
    assert verdict([]) is perf_stats.Verdict.FAIL


def test_combine_requires_every_representative_to_pass_somewhere() -> None:
    local = perf_stats.CombineInput(
        "local", perf_stats.Verdict.PASS, frozenset({"L-A"})
    )
    remote = perf_stats.CombineInput(
        "remote", perf_stats.Verdict.PASS, frozenset({"R-A"})
    )
    partial = perf_stats.CombineInput(
        "remote", perf_stats.Verdict.PASS_WITH_INCOMPLETE, frozenset({"R-A"})
    )
    required = ("L-A", "R-A")
    assert (
        perf_stats.combine_verdicts([local, remote], required)[0]
        is perf_stats.Verdict.PASS
    )
    only_local, reasons = perf_stats.combine_verdicts([local], required)
    assert only_local is perf_stats.Verdict.FAIL
    assert reasons == ("representative R-A did not pass in any comparison",)
    assert (
        perf_stats.combine_verdicts([local, partial], required)[0]
        is perf_stats.Verdict.PASS_WITH_INCOMPLETE
    )


# ---------------------------------------------------------------- xlsx_invariance


def test_identical_workbooks_compare_identical_with_equal_value_hashes(
    tmp_path: Path,
) -> None:
    rows: Rows = [["account", "amount"], ["자산총계", 1234567], ["부채", -5.5]]
    a = _workbook(tmp_path / "a.xlsx", {"data": rows}, merged="A1:B1")
    b = _workbook(tmp_path / "b.xlsx", {"data": rows}, merged="A1:B1")
    report = xlsx_invariance.compare_workbooks(a, b)
    assert report.identical
    assert (
        report.a_sha256 == report.b_sha256 == xlsx_invariance.xlsx_sha256_of_values(a)
    )


def test_allow_listed_generation_time_does_not_break_invariance(tmp_path: Path) -> None:
    a = _workbook(
        tmp_path / "a.xlsx",
        {"data": [["x"]], "metadata": _metadata_rows("2026-09-28T01:00:00.000001Z")},
    )
    b = _workbook(
        tmp_path / "b.xlsx",
        {"data": [["x"]], "metadata": _metadata_rows("2026-09-28T02:30:00.999999Z")},
    )
    report = xlsx_invariance.compare_workbooks(a, b)
    assert report.identical
    assert xlsx_invariance.xlsx_sha256_of_values(
        a
    ) == xlsx_invariance.xlsx_sha256_of_values(b)


def test_non_allow_listed_metadata_change_is_a_difference(tmp_path: Path) -> None:
    a = _workbook(
        tmp_path / "a.xlsx", {"metadata": [["key", "value"], ["dataset_id", "aaa"]]}
    )
    b = _workbook(
        tmp_path / "b.xlsx", {"metadata": [["key", "value"], ["dataset_id", "bbb"]]}
    )
    report = xlsx_invariance.compare_workbooks(a, b)
    assert not report.identical
    assert report.summary == {"cell_value": 1}
    assert report.differences[0].location == "B2"


def test_value_type_changes_are_differences(tmp_path: Path) -> None:
    a = _workbook(tmp_path / "a.xlsx", {"data": [[1]]})
    b = _workbook(tmp_path / "b.xlsx", {"data": [["1"]]})
    assert xlsx_invariance.compare_workbooks(a, b).summary == {"cell_value": 1}


def test_sheet_order_merges_formats_and_widths_are_compared(tmp_path: Path) -> None:
    a = _workbook(
        tmp_path / "a.xlsx", {"one": [["v", 1]], "two": [["w"]]}, merged="A2:B2"
    )
    b = _workbook(tmp_path / "b.xlsx", {"two": [["w"]], "one": [["v", 1]]})
    workbook = openpyxl.load_workbook(b)
    sheet = workbook["one"]
    sheet["B1"].number_format = "#,##0"
    sheet.column_dimensions["A"].width = 30
    workbook.save(b)
    report = xlsx_invariance.compare_workbooks(a, b)
    assert not report.identical
    assert set(report.summary) == {
        "sheet_names",
        "merged_ranges",
        "number_format",
        "column_width",
    }


def test_alignment_wrap_bold_freeze_and_sheet_state_are_compared(
    tmp_path: Path,
) -> None:
    rows: Rows = [["title", "long text"], ["a", 1]]
    a = _workbook(tmp_path / "a.xlsx", {"data": rows, "other": [["x"]]})
    b = _workbook(tmp_path / "b.xlsx", {"data": rows, "other": [["x"]]})
    assert xlsx_invariance.compare_workbooks(a, b).identical
    workbook = openpyxl.load_workbook(b)
    sheet = workbook["data"]
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["B1"].alignment = Alignment(wrap_text=True)
    sheet["A2"].font = Font(bold=True)
    sheet.freeze_panes = "A2"
    workbook["other"].sheet_state = "hidden"
    workbook.save(b)
    report = xlsx_invariance.compare_workbooks(a, b)
    assert report.summary == {
        "alignment": 2,
        "font_bold": 1,
        "freeze_panes": 1,
        "sheet_state": 1,
    }
    assert report.a_sha256 != report.b_sha256


def test_flow_identity_covers_the_sorted_unique_warning_codes() -> None:
    identity = xlsx_invariance.flow_identity
    data = "d" * 64
    assert identity(data, ["B", "A", "A"]) == identity(data, ["A", "B"])
    assert identity(data, ["A"]) != identity(data, ["A", "B"])
    assert identity(data, []) != identity("e" * 64, [])
    assert identity(None, ["A"]) is None


def test_row_datasets_compare_columns_counts_and_rows() -> None:
    base = xlsx_invariance.RowDataset(columns=("a", "b"), rows=({"a": 1, "b": "x"},))
    same = xlsx_invariance.RowDataset(columns=("a", "b"), rows=({"b": "x", "a": 1},))
    changed = xlsx_invariance.RowDataset(
        columns=("a", "b"), rows=({"a": 2, "b": "x"}, {"a": 3, "b": "y"})
    )
    assert xlsx_invariance.compare_row_datasets(base, same).identical
    report = xlsx_invariance.compare_row_datasets(base, changed)
    assert report.summary == {"row_count": 1, "row": 1}
    assert report.a_sha256 == xlsx_invariance.rows_sha256(base)


def test_invariance_cli_exit_codes(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    a = _workbook(tmp_path / "a.xlsx", {"data": [["x"]]})
    b = _workbook(tmp_path / "b.xlsx", {"data": [["y"]]})
    assert xlsx_invariance.main([str(a), str(a)]) == 0
    assert xlsx_invariance.main([str(a), str(b)]) == 1
    assert xlsx_invariance.main([str(a), str(tmp_path / "missing.xlsx")]) == 2
    assert "identical: False" in capsys.readouterr().out


# ---------------------------------------------------------------- stage_probes


def test_redact_url_drops_query_fragment_and_userinfo() -> None:
    url = "https://user:pw@opendart.fss.or.kr:443/api/list.json?crtfc_key=SECRET&corp_code=1#frag"
    assert (
        stage_probes.redact_url(url) == "https://opendart.fss.or.kr:443/api/list.json"
    )
    assert "SECRET" not in stage_probes.redact_url(
        "https://dart.fss.or.kr/x?crtfc_key=SECRET"
    )


class _FakeResponse:
    def __init__(self, status_code: int, content: bytes) -> None:
        self.status_code = status_code
        self.content = content


class _FakeClient:
    def get(self, url: str, *, params: dict[str, str]) -> _FakeResponse:
        del url
        return _FakeResponse(200, b"x" * len(params))


def test_upstream_monitor_counts_logs_without_query_and_enforces_budget() -> None:
    monitor = stage_probes.UpstreamMonitor(
        2, record_requests=True, client_class=_FakeClient
    )
    monitor.install()
    try:
        monitor.set_context({"flow": "L-C"})
        client = _FakeClient()
        client.get(
            "https://opendart.fss.or.kr/api/fnlttSinglAcntAll.json?crtfc_key=K",
            params={"a": "1"},
        )
        client.get("https://opendart.fss.or.kr/api/company.json", params={})
        with pytest.raises(stage_probes.UpstreamBudgetExceededError):
            client.get("https://opendart.fss.or.kr/api/company.json", params={})
    finally:
        monitor.uninstall()
    records = monitor.drain_records()
    assert monitor.count == 2
    assert monitor.exceeded
    assert [record["path"] for record in records] == [
        "/api/fnlttSinglAcntAll.json",
        "/api/company.json",
    ]
    assert all("crtfc_key" not in str(record) for record in records)
    assert records[0]["flow"] == "L-C"
    assert records[0]["bytes"] == 1
    assert _FakeClient().get("u", params={}).status_code == 200


_FAKE_DEFINITION_SOURCE: Final = """
def work(value: int) -> int:
    return value if value <= 0 else work(value - 1) + 1


class Service:
    def run(self, value: int) -> int:
        return work(value)
"""
_FAKE_USER_SOURCE: Final = """
from perf_probe_fake import work


def call(value: int) -> int:
    return work(value)
"""


@pytest.fixture
def fake_modules(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> Iterator[tuple[ModuleType, ModuleType]]:
    # Real modules on disk, so calls resolve "work" through module globals the
    # way product code does (including a "from x import y" use site).
    package = tmp_path / _FAKE_MODULE
    package.mkdir()
    (package / "__init__.py").write_text(_FAKE_DEFINITION_SOURCE, encoding="utf-8")
    (package / "user.py").write_text(_FAKE_USER_SOURCE, encoding="utf-8")
    monkeypatch.syspath_prepend(str(tmp_path))
    try:
        yield (
            importlib.import_module(_FAKE_MODULE),
            importlib.import_module(_FAKE_USER_MODULE),
        )
    finally:
        sys.modules.pop(_FAKE_USER_MODULE, None)
        sys.modules.pop(_FAKE_MODULE, None)


def test_stage_timers_patch_definition_and_use_sites_then_restore(
    fake_modules: tuple[ModuleType, ModuleType],
) -> None:
    definition, user = fake_modules
    original_work = definition.work
    original_run = definition.Service.run
    timers = stage_probes.StageTimers(
        [f"{_FAKE_MODULE}:work", f"{_FAKE_MODULE}:Service.run"]
    )
    timers.install()
    try:
        assert user.work is definition.work
        assert user.work is not original_work
        assert definition.Service().run(3) == 3
        user_call: Callable[[int], int] = user.call
        assert user_call(2) == 2
    finally:
        timers.uninstall()
    snapshot = timers.snapshot()
    # recursion inside one stage is timed once per outermost call
    assert snapshot[f"{_FAKE_MODULE}:work"]["calls"] == 2
    assert snapshot[f"{_FAKE_MODULE}:Service.run"]["calls"] == 1
    assert user.work is original_work
    assert definition.work is original_work
    assert definition.Service.run is original_run


def test_stage_timers_accumulate_safely_across_threads(
    fake_modules: tuple[ModuleType, ModuleType],
) -> None:
    definition, _ = fake_modules
    timers = stage_probes.StageTimers([f"{_FAKE_MODULE}:work"])
    timers.install()
    barrier = threading.Barrier(8)

    def call(value: int) -> int:
        barrier.wait()
        work: Callable[[int], int] = definition.work
        return work(value)

    try:
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(call, range(8)))
    finally:
        timers.uninstall()
    assert results == list(range(8))
    assert timers.snapshot()[f"{_FAKE_MODULE}:work"]["calls"] == 8
    timers.reset()
    assert timers.snapshot() == {}


def test_stage_timers_reject_malformed_targets() -> None:
    with pytest.raises(ValueError, match="module:qualname"):
        stage_probes.StageTimers(["no_colon"]).install()


# ---------------------------------------------------------------- compare_runs / latency_harness
#
# Both scripts import their siblings by bare name, as they do when run from
# scripts/perf, so the tests load them the same way (they touch no network).


@pytest.fixture
def compare_runs(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    monkeypatch.syspath_prepend(str(_PERF_DIR))
    return importlib.import_module("compare_runs")


@pytest.fixture
def harness(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    monkeypatch.syspath_prepend(str(_PERF_DIR))
    return importlib.import_module("latency_harness")


def _record(
    flow: str,
    corp_code: str,
    rep: int,
    total_ms: float,
    *,
    valid: bool = True,
    warnings: tuple[str, ...] = (),
) -> dict[str, object]:
    return {
        "flow": flow,
        "corp_code": corp_code,
        "rep": rep,
        "total_ms": total_ms,
        "valid": valid,
        "first_in_session": rep == 1,
        "possibly_cold": False,
        "xlsx_sha256_of_values": "d" * 64 if valid else None,
        "rows_sha256": None,
        "warning_codes": list(warnings),
        "xlsx_path": None,
    }


def _write_run(
    run_dir: Path,
    records: list[dict[str, object]],
    *,
    target: str = "local",
    code_root: str | None = None,
) -> Path:
    run_dir.mkdir(parents=True)
    (run_dir / "flows.jsonl").write_text(
        "".join(json.dumps(record) + "\n" for record in records), encoding="utf-8"
    )
    env = {"target": target, "sessions": [{"code_root": code_root}]}
    (run_dir / "env.json").write_text(json.dumps(env), encoding="utf-8")
    return run_dir


def _write_ab(
    root: Path,
    base: list[dict[str, object]],
    new: list[dict[str, object]],
    *,
    arm_roots: tuple[str, str] | None = None,
) -> Path:
    projects = {"A": str(root / "tree_a"), "B": str(root / "tree_b")}
    roots = arm_roots or (projects["A"], projects["B"])
    _write_run(root / "A", base, code_root=roots[0])
    _write_run(root / "B", new, code_root=roots[1])
    (root / "ab.json").write_text(json.dumps({"projects": projects}), encoding="utf-8")
    return root


def _verdict(out_dir: Path) -> dict[str, object]:
    comparison = json.loads((out_dir / "comparison.json").read_text(encoding="utf-8"))
    verdict: dict[str, object] = comparison["verdict"]
    return verdict


def _statuses(verdict: dict[str, object]) -> dict[str, str]:
    flows = verdict["flows"]
    assert isinstance(flows, list)
    return {str(flow["flow"]): str(flow["status"]) for flow in flows}


type Records = list[dict[str, object]]


def _s1(new_x: float, new_y: float) -> tuple[Records, Records]:
    base = [_record("L-A", "X", rep, 100.0) for rep in (1, 2, 3)] + [
        _record("L-A", "Y", rep, 1000.0) for rep in (1, 2, 3)
    ]
    new = [
        *(_record("L-A", "X", rep, new_x) for rep in (1, 2, 3)),
        _record("L-A", "Y", 1, new_y),
        _record("L-A", "Y", 2, 50.0, valid=False),
        _record("L-A", "Y", 3, 50.0, valid=False),
    ]
    return base, new


def test_compare_s1_invalid_reps_make_the_flow_incomplete(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    # W3 S1: no company got faster, but the slow company lost two new reps.
    base, new = _s1(100.0, 1000.0)
    ab = _write_ab(tmp_path / "s1", base, new)
    assert compare_runs.main(["--ab", str(ab)]) == 1
    verdict = _verdict(ab)
    assert verdict["overall"] == "FAIL"
    assert _statuses(verdict) == {"L-A": "INCOMPLETE"}
    assert compare_runs.main(["--ab", str(ab), "--allow-incomplete"]) == 1


def test_compare_allow_incomplete_labels_a_real_speedup_pass_with_incomplete(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    base, new = _s1(40.0, 400.0)
    ab = _write_ab(tmp_path / "s1", base, new)
    assert compare_runs.main(["--ab", str(ab)]) == 1
    assert compare_runs.main(["--ab", str(ab), "--allow-incomplete"]) == 3
    verdict = _verdict(ab)
    assert verdict["overall"] == "PASS_WITH_INCOMPLETE"
    markdown = (ab / "comparison.md").read_text(encoding="utf-8")
    assert "PASS_WITH_INCOMPLETE is not a PASS" in markdown
    assert "invalid samples in new: L-A/Y x2" in markdown


def test_compare_s2_secondary_flow_invalid_in_both_arms_is_missing(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    base = [_record("L-A", "X", rep, 1000.0) for rep in (1, 2, 3)] + [
        _record("L-C", "X", rep, 100.0, valid=False) for rep in (1, 2, 3)
    ]
    new = [_record("L-A", "X", rep, 400.0) for rep in (1, 2, 3)] + [
        _record("L-C", "X", rep, 900.0, valid=False) for rep in (1, 2, 3)
    ]
    ab = _write_ab(tmp_path / "s2", base, new)
    assert compare_runs.main(["--ab", str(ab)]) == 1
    verdict = _verdict(ab)
    assert _statuses(verdict) == {"L-A": "COMPLETE", "L-C": "MISSING"}
    assert verdict["overall"] == "FAIL"


def test_compare_takes_the_representative_from_the_run_target(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    only_secondary = [_record("L-C", "X", 1, 100.0)]
    base = _write_run(tmp_path / "base", only_secondary)
    new = _write_run(tmp_path / "new", only_secondary)
    out = tmp_path / "out"
    arguments = ["--base", str(base), "--new", str(new), "--out", str(out)]
    assert compare_runs.main(arguments) == 1
    assert _verdict(out)["missing_representative"] == ["L-A"]
    remote = _write_run(tmp_path / "remote", only_secondary, target="remote-asgi")
    assert compare_runs.main(["--base", str(base), "--new", str(remote)]) == 2


def test_compare_detects_a_warning_code_difference(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    base = [_record("L-A", "X", 1, 1000.0, warnings=("FALLBACK_SOURCE_USED",))]
    new = [_record("L-A", "X", 1, 100.0)]
    ab = _write_ab(tmp_path / "warn", base, new)
    assert compare_runs.main(["--ab", str(ab)]) == 1
    verdict = _verdict(ab)
    assert verdict["invariance"] == "FAIL"
    assert verdict["latency"] == "PASS"


def test_compare_refuses_exclusions_for_ab_runs(
    compare_runs: ModuleType,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    records = [_record("L-A", "X", 1, 1.0)]
    ab = _write_ab(tmp_path / "ab", records, records)
    assert compare_runs.main(["--ab", str(ab), "--exclude-cold"]) == 2
    assert compare_runs.main(["--ab", str(ab), "--exclude-first-in-session"]) == 2
    assert "symmetric" in capsys.readouterr().err
    assert not (ab / "comparison.json").exists()


def test_compare_ab_verifies_each_arm_imported_its_own_tree(
    compare_runs: ModuleType,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    records = [_record("L-A", "X", 1, 1.0)]
    roots = (str(tmp_path / "somewhere_else"), str(tmp_path / "ab" / "tree_b"))
    ab = _write_ab(tmp_path / "ab", records, records, arm_roots=roots)
    assert compare_runs.main(["--ab", str(ab)]) == 2
    assert "arm A is not isolated" in capsys.readouterr().err


def _comparison_json(path: Path, overall: str, passed: dict[str, bool]) -> Path:
    flows = [
        {"flow": flow, "representative": True, "passed": ok, "status": "COMPLETE"}
        for flow, ok in passed.items()
    ]
    document = {"schema_version": 2, "verdict": {"overall": overall, "flows": flows}}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(document), encoding="utf-8")
    return path


def test_combine_needs_both_l_a_and_r_a_passing(
    compare_runs: ModuleType, tmp_path: Path
) -> None:
    local = _comparison_json(tmp_path / "l" / "comparison.json", "PASS", {"L-A": True})
    remote = _comparison_json(tmp_path / "r" / "comparison.json", "PASS", {"R-A": True})
    failed = _comparison_json(
        tmp_path / "f" / "comparison.json", "FAIL", {"R-A": False}
    )
    out = ["--out", str(tmp_path / "combined")]
    assert compare_runs.main(["combine", str(local), str(remote), *out]) == 0
    assert compare_runs.main(["combine", str(local), *out]) == 1
    assert compare_runs.main(["combine", str(local), str(failed), *out]) == 1
    combined_path = tmp_path / "combined" / "combined.json"
    assert json.loads(combined_path.read_text(encoding="utf-8"))["overall"] == "FAIL"


_FAKE_KEY: Final = "FAKEKEY0123456789abcdef0123456789abcdef0"


class _RaisingCaller:
    def __init__(self, message: str) -> None:
        self.message = message

    def call(self, tool: str, args: dict[str, object]) -> dict[str, object]:
        del tool, args
        raise RuntimeError(self.message)


def test_redaction_happens_before_truncation(harness: ModuleType) -> None:
    assert len(_FAKE_KEY) == 40
    for offset in range(240, 300):
        text = "x" * offset + "?crtfc_key=" + _FAKE_KEY + "&corp_code=1"
        envelope = harness.safe_call(_RaisingCaller(text), "tool", {}, _FAKE_KEY)
        _, recorded = harness.error_fields(
            {"error": {"code": "E", "message": text}}, _FAKE_KEY
        )
        for output in (envelope["error"]["message"], recorded):
            assert len(output) <= harness.MESSAGE_LIMIT
            assert _FAKE_KEY[:6] not in output
    # a message the fitness runner already cut inside the key
    already_cut = "y" * 280 + _FAKE_KEY[:20]
    assert _FAKE_KEY[:6] not in harness.redact_secret(already_cut, _FAKE_KEY)


def test_ab_schedule_is_counterbalanced_abba(harness: ModuleType) -> None:
    schedule = harness.ab_schedule(["c1", "c2"], 2)
    assert "".join(arm for _, _, arm in schedule) == "ABBABAAB"
    assert [(corp, rep) for corp, rep, _ in schedule[:4]] == [
        ("c1", 1),
        ("c1", 1),
        ("c1", 2),
        ("c1", 2),
    ]
    firsts = [arm for index, (_, _, arm) in enumerate(schedule) if index % 2 == 0]
    assert firsts.count("A") == firsts.count("B")


def test_code_root_mismatch_is_detected_after_normalization(
    harness: ModuleType, tmp_path: Path
) -> None:
    tree = tmp_path / "Tree"
    tree.mkdir()
    assert harness.code_root_mismatch(tree, tmp_path / "Tree" / ".." / "Tree") is None
    if os.name == "nt":
        assert harness.code_root_mismatch(tree, str(tree).upper()) is None
    assert harness.code_root_mismatch(tree, tmp_path / "other") is not None


def test_ab_child_env_drops_cross_tree_variables(harness: ModuleType) -> None:
    env = harness.ab_child_env(
        {
            "PATH": "p",
            "VIRTUAL_ENV": "v",
            "PYTHONPATH": "src",
            "PythonHome": "h",
            "UV_PROJECT_ENVIRONMENT": "e",
        }
    )
    assert env == {"PATH": "p", "PYTHONIOENCODING": "utf-8"}


def test_every_company_is_pinned_and_a_disagreeing_resolved_file_is_refused(
    harness: ModuleType,
) -> None:
    catalog = harness.load_catalog()
    assert all(company.expected_rcept_no for company in catalog.companies)
    assert all(company.expected_attachment_id for company in catalog.companies)
    pinned = {
        company.corp_code: harness.ResolvedCompany(
            corp_code=company.corp_code,
            name=company.name,
            rcept_no=company.expected_rcept_no,
            attachment_id=company.expected_attachment_id,
        )
        for company in catalog.companies
    }
    harness.check_resolved_pins(catalog, pinned)
    changed = catalog.companies[1]
    pinned[changed.corp_code] = harness.ResolvedCompany(
        corp_code=changed.corp_code,
        name=changed.name,
        rcept_no="20260101000000",
        attachment_id=changed.expected_attachment_id,
    )
    with pytest.raises(harness.HarnessError, match=changed.corp_code):
        harness.check_resolved_pins(catalog, pinned)
