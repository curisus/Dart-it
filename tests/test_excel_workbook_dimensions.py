from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.xml.constants import MAX_COLUMN, MAX_ROW

from dart_crawler.excel_page_models import ExcelRow
from dart_crawler.excel_query_workbook_plan import (
    EXCEL_DATA_ROWS_PER_SHEET,
    EXCEL_MAX_COLUMNS,
    ExcelWorkbookOptions,
)
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.result import ErrorCode
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


def test_production_excel_dimension_constants_match_native_limits() -> None:
    options = ExcelWorkbookOptions()

    assert options.data_rows_per_sheet == MAX_ROW - 1
    assert options.data_rows_per_sheet == EXCEL_DATA_ROWS_PER_SHEET
    assert EXCEL_MAX_COLUMNS == MAX_COLUMN


def test_injected_row_limit_splits_data_in_source_order(tmp_path: Path) -> None:
    columns = ("sequence", "label")
    rows = _ordered_rows(5)
    clock = _clock()

    result = publish_excel_dataset(
        make_normalized_dataset(columns, rows),
        tmp_path / "output",
        clock=clock,
        options=ExcelWorkbookOptions(data_rows_per_sheet=2),
    )

    assert result.ok is True
    exported = result.data
    assert exported is not None
    assert exported.sheet_names == (
        "data",
        "data_2",
        "data_3",
        "metadata",
        "warnings",
    )
    assert clock.calls == 1
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == list(exported.sheet_names)
        assert all(sheet.sheet_state == "visible" for sheet in workbook.worksheets)
        observed: list[tuple[int, str]] = []
        for name in ("data", "data_2", "data_3"):
            sheet = workbook[name]
            assert tuple(cell.value for cell in sheet[1]) == columns
            for row in sheet.iter_rows(min_row=2, max_col=2):
                sequence = row[0].value
                label = row[1].value
                assert isinstance(sequence, int)
                assert not isinstance(sequence, bool)
                assert isinstance(label, str)
                observed.append((sequence, label))
        assert observed == [(index, f"row-{index}") for index in range(5)]
        metadata = {
            row[0].value: row[1].value
            for row in workbook["metadata"].iter_rows(min_col=1, max_col=2)
        }
        assert metadata["data_sheet_names"] == '["data","data_2","data_3"]'
    finally:
        workbook.close()


def test_ten_thousand_rows_reopen_in_exact_order(tmp_path: Path) -> None:
    rows = _ordered_rows(10_000)

    result = publish_excel_dataset(
        make_normalized_dataset(("sequence", "label"), rows),
        tmp_path / "output",
        clock=_clock(),
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is True
    exported = result.data
    assert exported is not None
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        sheet = workbook["data"]
        assert sheet.max_row == 10_001
        assert [sheet.cell(row, 1).value for row in range(2, 10_002)] == list(
            range(10_000)
        )
        assert [sheet.cell(row, 2).value for row in range(2, 10_002)] == [
            f"row-{index}" for index in range(10_000)
        ]
    finally:
        workbook.close()


def test_excel_column_limit_is_accepted_at_exact_boundary(tmp_path: Path) -> None:
    columns = tuple(f"column_{index:05d}" for index in range(MAX_COLUMN))

    result = publish_excel_dataset(
        make_normalized_dataset(columns, ()),
        tmp_path / "output",
        clock=_clock(),
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is True
    exported = result.data
    assert exported is not None
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        sheet = workbook["data"]
        assert sheet.max_column == MAX_COLUMN
        assert sheet.cell(1, 1).value == "column_00000"
        assert sheet.cell(1, MAX_COLUMN).value == "column_16383"
    finally:
        workbook.close()


def test_excel_column_overflow_fails_before_output_creation(tmp_path: Path) -> None:
    columns = tuple(
        f"column_{index:05d}" for index in range(MAX_COLUMN + 1)
    )
    output_root = tmp_path / "output"
    clock = _clock()

    result = publish_excel_dataset(
        make_normalized_dataset(columns, ()),
        output_root,
        clock=clock,
        options=ExcelWorkbookOptions(),
    )

    assert result.ok is False
    assert result.data is None
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert result.error.retryable is False
    assert result.error.details == {"reason": "excel_column_limit_exceeded"}
    assert clock.calls == 0
    assert not output_root.exists()


def _ordered_rows(count: int) -> tuple[ExcelRow, ...]:
    rows: list[ExcelRow] = [
        {"sequence": index, "label": f"row-{index}"} for index in range(count)
    ]
    return tuple(rows)


def _clock() -> RecordingClock:
    return RecordingClock(datetime(2026, 5, 6, 7, 8, 9, 10, tzinfo=UTC))
