from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from dart_crawler.excel_export_result import ExcelCleanupWarningCode
from dart_crawler.excel_query_workbook_plan import ExcelWorkbookOptions
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.result import WarningCode, WarningInfo
from tests.excel_publication_test_support import CleanupFailurePublicationOps
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


@pytest.mark.parametrize(
    ("fail_temp", "fail_lock", "expected_codes"),
    [
        (True, False, (ExcelCleanupWarningCode.OUTPUT_TEMP_CLEANUP_FAILED,)),
        (False, True, (ExcelCleanupWarningCode.OUTPUT_LOCK_CLEANUP_FAILED,)),
        (
            True,
            True,
            (
                ExcelCleanupWarningCode.OUTPUT_TEMP_CLEANUP_FAILED,
                ExcelCleanupWarningCode.OUTPUT_LOCK_CLEANUP_FAILED,
            ),
        ),
    ],
)
def test_post_publish_cleanup_failures_preserve_final_and_warn_in_order(
    fail_temp: bool,
    fail_lock: bool,
    expected_codes: tuple[ExcelCleanupWarningCode, ...],
    tmp_path: Path,
) -> None:
    source_warning = WarningInfo(
        code=WarningCode.FALLBACK_SOURCE_USED,
        message="source warning",
        details={"source": "fallback"},
    )
    dataset = make_normalized_dataset(
        ("value",),
        ({"value": "완전한 값"},),
        warnings=(source_warning,),
    )
    operations = CleanupFailurePublicationOps(
        fail_temp=fail_temp,
        fail_lock=fail_lock,
    )

    result = publish_excel_dataset(
        dataset,
        tmp_path / "output",
        clock=RecordingClock(datetime(2026, 8, 9, tzinfo=UTC)),
        options=ExcelWorkbookOptions(),
        file_ops=operations,
    )

    assert result.ok is True
    assert result.data is not None
    assert tuple(warning.code for warning in result.warnings) == expected_codes
    assert all(warning.details == {} for warning in result.warnings)
    assert result.next_action is None
    final_path = Path(result.data.absolute_path)
    assert final_path.exists()
    assert len(operations.cleanup_paths) == 2
    temp_path, lock_path = operations.cleanup_paths
    assert not temp_path.name.endswith(".lock")
    assert lock_path.name.endswith(".lock")
    assert temp_path.exists() is fail_temp
    assert lock_path.exists() is fail_lock

    workbook = load_workbook(final_path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == ["data", "metadata", "warnings"]
        data = workbook["data"]
        warning_sheet = workbook["warnings"]
        assert isinstance(data, Worksheet)
        assert isinstance(warning_sheet, Worksheet)
        assert data.cell(2, 1).value == "완전한 값"
        assert tuple(cell.value for cell in warning_sheet[2]) == (
            1,
            WarningCode.FALLBACK_SOURCE_USED.value,
            "source warning",
            '{"source":"fallback"}',
        )
    finally:
        workbook.close()
    temp_path.unlink(missing_ok=True)
    lock_path.unlink(missing_ok=True)
