from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import dart_crawler.excel_query_workbook_writer as workbook_writer
from dart_crawler.excel_export_result import Result as ExcelResult
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.excel_query_workbook_plan import (
    ExcelWorkbookOptions,
    ExcelWorkbookPlan,
)
from dart_crawler.excel_query_workbook_writer import WorkbookWriteOutcome
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.result import ErrorCode
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.local_excel_export_test_support import RecordingClock


def test_malformed_workbook_xml_is_typed_and_fully_cleaned(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    original_writer = workbook_writer.write_excel_query_workbook

    def write_with_malformed_workbook_xml(
        path: Path,
        plan: ExcelWorkbookPlan,
    ) -> WorkbookWriteOutcome:
        outcome = original_writer(path, plan)
        with ZipFile(path) as source:
            members = tuple(
                (member.filename, source.read(member.filename))
                for member in source.infolist()
            )
        with ZipFile(path, "w", ZIP_DEFLATED) as destination:
            for name, content in members:
                destination.writestr(
                    name,
                    b"<broken" if name == "xl/workbook.xml" else content,
                )
        return outcome

    monkeypatch.setattr(
        workbook_writer,
        "write_excel_query_workbook",
        write_with_malformed_workbook_xml,
    )
    output_root = tmp_path / "output"

    result = _publish(output_root)

    _assert_validation_failure(result)
    assert list(output_root.iterdir()) == []


def test_single_cell_merge_is_rejected_before_publication(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    original_writer = workbook_writer.write_excel_query_workbook

    def write_with_single_cell_merge(
        path: Path,
        plan: ExcelWorkbookPlan,
    ) -> WorkbookWriteOutcome:
        outcome = original_writer(path, plan)
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            sheet = workbook["data"]
            assert isinstance(sheet, Worksheet)
            sheet.merge_cells("A1:A1")
            workbook.save(path)
        finally:
            workbook.close()
        return outcome

    monkeypatch.setattr(
        workbook_writer,
        "write_excel_query_workbook",
        write_with_single_cell_merge,
    )
    output_root = tmp_path / "output"

    result = _publish(output_root)

    _assert_validation_failure(result)
    assert list(output_root.iterdir()) == []


def _assert_validation_failure(
    result: ExcelResult[ExcelExportResult],
) -> None:
    assert result.ok is False
    assert result.data is None
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert result.error.retryable is False
    assert result.error.details == {
        "reason": "excel_workbook_validation_failed"
    }


def _publish(output_root: Path) -> ExcelResult[ExcelExportResult]:
    dataset = make_normalized_dataset(("value",), ())
    return publish_excel_dataset(
        dataset,
        output_root,
        clock=RecordingClock(datetime(2026, 7, 8, tzinfo=UTC)),
        options=ExcelWorkbookOptions(),
    )
