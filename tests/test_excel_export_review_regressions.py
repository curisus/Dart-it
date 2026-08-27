import json
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from typing import Protocol
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from mcp_types import CallToolResult
from openpyxl import load_workbook
from openpyxl.chart.bar_chart import BarChart
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

import dart_crawler.excel_query_workbook_writer as workbook_writer
from dart_crawler.domain import Company, Market, MatchConfidence
from dart_crawler.excel_export_result import Result as ExcelResult
from dart_crawler.excel_page_models import ExcelRow
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.excel_query_workbook_plan import (
    ExcelWorkbookOptions,
    ExcelWorkbookPlan,
)
from dart_crawler.excel_query_workbook_writer import WorkbookWriteOutcome
from dart_crawler.excel_safe_publication import publish_excel_dataset
from dart_crawler.mcp_server import mcp
from dart_crawler.remote_server import create_remote_server
from dart_crawler.result import ErrorCode, Result, WarningCode
from tests.excel_dispatch_fakes import RecordingExcelServiceFactory
from tests.excel_query_workbook_test_support import make_normalized_dataset
from tests.excel_service_fixture_factory import empty_excel_service_responses
from tests.local_excel_export_test_support import (
    RecordingClock,
    install_local_export_runtime,
)
from tests.local_excel_mcp_result_support import excel_export_result

_CORE_WARNING_CODES = (
    "PARTIAL_COLLECTION",
    "IMAGE_CONTENT_SKIPPED",
    "AMOUNT_MISMATCH",
    "COMPARISON_UNAVAILABLE",
    "FALLBACK_SOURCE_USED",
    "ORIGINAL_FILING_SOURCE_USED",
    "VIEWER_DISCOVERY_SKIPPED",
    "EXISTING_FILE_REUSED",
)
_CLEANUP_WARNING_CODES = (
    "OUTPUT_TEMP_CLEANUP_FAILED",
    "OUTPUT_LOCK_CLEANUP_FAILED",
)


class _ChartAdder(Protocol):
    def add_chart(self, chart: BarChart) -> None: ...


def _add_chart(sheet: _ChartAdder) -> None:
    sheet.add_chart(BarChart())


def test_zip_shaped_malformed_workbook_is_typed_and_fully_cleaned(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    original_writer = workbook_writer.write_excel_query_workbook

    def write_without_workbook_part(
        path: Path,
        plan: ExcelWorkbookPlan,
    ) -> WorkbookWriteOutcome:
        outcome = original_writer(path, plan)
        with ZipFile(path) as source:
            members = tuple(
                (member.filename, source.read(member.filename))
                for member in source.infolist()
                if member.filename != "xl/workbook.xml"
            )
        with ZipFile(path, "w", ZIP_DEFLATED) as destination:
            for name, content in members:
                destination.writestr(name, content)
        return outcome

    monkeypatch.setattr(
        workbook_writer,
        "write_excel_query_workbook",
        write_without_workbook_part,
    )
    output_root = tmp_path / "output"

    result = _publish(output_root)

    assert result.ok is False
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert list(output_root.iterdir()) == []


@pytest.mark.parametrize("artifact", ["table", "chart"])
def test_reopen_validation_rejects_injected_workbook_artifacts(
    artifact: str,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    original_writer = workbook_writer.write_excel_query_workbook

    def write_with_artifact(
        path: Path,
        plan: ExcelWorkbookPlan,
    ) -> WorkbookWriteOutcome:
        outcome = original_writer(path, plan)
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            sheet = workbook["data"]
            assert isinstance(sheet, Worksheet)
            if artifact == "table":
                sheet.add_table(Table(displayName="InjectedTable", ref="A1:A2"))
            else:
                _add_chart(sheet)
            workbook.save(path)
        finally:
            workbook.close()
        return outcome

    monkeypatch.setattr(
        workbook_writer,
        "write_excel_query_workbook",
        write_with_artifact,
    )
    output_root = tmp_path / "output"

    result = _publish(output_root, rows=({"value": 1},))

    assert result.ok is False
    assert result.error is not None
    assert result.error.code is ErrorCode.VALIDATION_FAILED
    assert list(output_root.iterdir()) == []


@pytest.mark.anyio
async def test_actual_local_export_preserves_success_next_action(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    company = Company(
        company_name="회사",
        corp_code="00123456",
        stock_code="123456",
        market=Market.KOSPI,
        ranking=1,
        match_confidence=MatchConfidence.EXACT,
    )
    factory = RecordingExcelServiceFactory(
        replace(
            empty_excel_service_responses(),
            search_companies=Result[tuple[Company, ...]].success(
                (company,),
                next_action="authoritative-success-action",
            ),
        )
    )
    _ = install_local_export_runtime(
        monkeypatch,
        tmp_path,
        factory,
        RecordingClock(datetime(2026, 1, 2, tzinfo=UTC)),
    )

    called = await mcp.call_tool(
        "export_query_excel",
        {
            "request": {
                "domain": "search_companies",
                "arguments": {"company_query": "회사"},
            }
        },
    )

    assert isinstance(called, CallToolResult)
    result = excel_export_result(called)
    assert result.ok is True
    assert result.warnings == ()
    assert result.next_action == (
        "authoritative-success-action"
    )


@pytest.mark.anyio
async def test_cleanup_warning_schema_is_local_to_new_export_tool() -> None:
    local_tools = {tool.name: tool for tool in await mcp.list_tools()}
    remote_tools = {
        tool.name: tool for tool in await create_remote_server().list_tools()
    }

    assert tuple(code.value for code in WarningCode) == _CORE_WARNING_CODES
    for tool in (
        local_tools["search_companies"],
        local_tools["export_report_excel"],
        remote_tools["search_companies"],
    ):
        schema_text = json.dumps(tool.output_schema, sort_keys=True)
        assert all(code not in schema_text for code in _CLEANUP_WARNING_CODES)
    export_schema = json.dumps(
        local_tools["export_query_excel"].output_schema,
        sort_keys=True,
    )
    assert all(code in export_schema for code in _CLEANUP_WARNING_CODES)


def _publish(
    output_root: Path,
    *,
    rows: tuple[ExcelRow, ...] = (),
) -> ExcelResult[ExcelExportResult]:
    dataset = make_normalized_dataset(("value",), rows)
    return publish_excel_dataset(
        dataset,
        output_root,
        clock=RecordingClock(datetime(2026, 7, 8, tzinfo=UTC)),
        options=ExcelWorkbookOptions(),
    )
