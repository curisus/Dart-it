from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from mcp_types import CallToolResult
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from pydantic import SecretStr

from dart_crawler.domain import Company, Market, MatchConfidence
from dart_crawler.excel_page_models import ExcelPage, ExcelRow, ExcelScalar
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.excel_query_service import ExcelQueryServiceFactory
from dart_crawler.http_client import HttpClient
from dart_crawler.mcp_server import mcp
from dart_crawler.query_limits import EXCEL_POLICY
from dart_crawler.result import JsonObject, Result, WarningCode, WarningInfo
from tests.excel_dispatch_fakes import RecordingExcelServiceFactory
from tests.excel_service_fixture_factory import empty_excel_service_responses
from tests.local_excel_export_test_support import (
    RecordingClock,
    install_local_export_runtime,
)
from tests.local_excel_mcp_result_support import structured_content
from tests.remote_server_test_support import RemoteRequest, call_tool_envelope


@pytest.mark.anyio
async def test_local_xlsx_equals_complete_actual_remote_page_traversal(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source_warning = WarningInfo(
        code=WarningCode.FALLBACK_SOURCE_USED,
        message="ordered warning",
        details={"unicode": "값", "nested": [1, {"key": True}]},
    )
    responses = replace(
        empty_excel_service_responses(),
        search_companies=Result[tuple[Company, ...]].success(
            _companies(5),
            warnings=(source_warning,),
        ),
    )
    local_factory = RecordingExcelServiceFactory(responses)
    remote_factory = RecordingExcelServiceFactory(responses)
    clock = RecordingClock(datetime(2026, 10, 11, 12, 13, 14, 15, tzinfo=UTC))
    _ = install_local_export_runtime(monkeypatch, tmp_path, local_factory, clock)
    monkeypatch.setenv("DART_MCP_CURSOR_SECRET", "s" * 32)

    def build_remote_factory(
        api_key: SecretStr,
        http_client: HttpClient,
    ) -> ExcelQueryServiceFactory:
        del api_key, http_client
        return remote_factory

    monkeypatch.setattr(
        "dart_crawler.remote_server._excel_page_factory",
        build_remote_factory,
    )
    local_call = await mcp.call_tool(
        "export_query_excel",
        {
            "request": {
                "domain": "search_companies",
                "arguments": {"company_query": "회사"},
            }
        },
    )
    assert isinstance(local_call, CallToolResult)
    local_result = Result[ExcelExportResult].model_validate(
        structured_content(local_call)
    )
    exported = local_result.data
    assert exported is not None

    remote_rows: list[ExcelRow] = []
    cursor: str | None = None
    terminal_page: ExcelPage | None = None
    while True:
        request: JsonObject = {
            "domain": "search_companies",
            "arguments": {"company_query": "회사"},
            "page_size": 2,
            "cursor": cursor,
        }
        envelope = await call_tool_envelope(
            "load_excel_page",
            {"request": request},
            RemoteRequest(headers=(("X-OpenDART-API-Key", "request-key"),)),
        )
        page_result = Result[ExcelPage].model_validate(envelope)
        page = page_result.data
        assert page is not None
        assert page.warnings == (source_warning,)
        remote_rows.extend(page.rows)
        terminal_page = page
        cursor = page.next_cursor
        if cursor is None:
            break

    assert terminal_page is not None
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        data = workbook["data"]
        columns = tuple(cell.value for cell in data[1])
        assert all(isinstance(column, str) for column in columns)
        workbook_rows: list[ExcelRow] = []
        for row_index in range(2, data.max_row + 1):
            row: ExcelRow = {}
            for column_index, raw_column in enumerate(columns, start=1):
                assert isinstance(raw_column, str)
                row[raw_column] = _scalar(data.cell(row_index, column_index))
            workbook_rows.append(row)
        assert workbook_rows == remote_rows
        metadata = {
            row[0].value: row[1].value
            for row in workbook["metadata"].iter_rows(min_col=1, max_col=2)
        }
        assert metadata["request_fingerprint"] == terminal_page.request_fingerprint
        assert metadata["source_fingerprint"] == terminal_page.source_fingerprint
        assert metadata["dataset_id"] == terminal_page.dataset_id
        assert metadata["total_rows"] == terminal_page.total_rows == 5
        assert tuple(cell.value for cell in workbook["warnings"][2]) == (
            1,
            "FALLBACK_SOURCE_USED",
            "ordered warning",
            '{"nested":[1,{"key":true}],"unicode":"값"}',
        )
    finally:
        workbook.close()
    assert local_factory.policies == [EXCEL_POLICY]
    assert len(local_factory.services) == 1
    assert len(local_factory.services[0].calls) == 1
    assert len(remote_factory.services) == 3
    assert remote_factory.policies == [EXCEL_POLICY] * 3
    assert all(len(service.calls) == 1 for service in remote_factory.services)
    assert clock.calls == 1


def _companies(count: int) -> tuple[Company, ...]:
    return tuple(
        Company(
            company_name=f"회사-{index}:=literal:한글🙂",
            corp_code=f"{index + 1:08d}",
            stock_code=f"{index + 1:06d}",
            market=Market.KOSPI,
            ranking=index + 1,
            match_confidence=MatchConfidence.EXACT,
        )
        for index in range(count)
    )


def _scalar(cell: Cell | MergedCell) -> ExcelScalar:
    assert isinstance(cell, Cell)
    value = cell.value
    assert isinstance(value, (str, int, float, bool, type(None)))
    return value
