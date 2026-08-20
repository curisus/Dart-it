from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from dart_crawler import workbook_layout
from dart_crawler.document_model import ParsedDocument
from dart_crawler.document_validation import validate_document
from dart_crawler.excel_export import ExcelExportService, ExportContext
from dart_crawler.value_parser import parse_cell_value, thousands_number_format
from dart_crawler.workbook_validation import validate_workbook
from dart_crawler.xml_parser import parse_xml_document


def _parse_document(markup: str) -> ParsedDocument:
    result = parse_xml_document(markup.encode())
    assert result.ok is True
    assert result.data is not None
    return result.data


def _document(*, include_image: bool = False) -> ParsedDocument:
    parts = [
        "<document>",
        "<heading>재무상태표</heading>",
        '<table><tr><td colspan="2">계정</td></tr>',
        "<tr><td>자산</td><td>1,000</td></tr></table>",
        "<heading>손익 및 포괄손익계산서</heading>",
        "<table><tr><td>매출</td><td>(10)</td></tr></table>",
        "<heading>자본변동표</heading>",
        "<table><tr><td>자본</td><td>5</td></tr></table>",
        "<heading>현금흐름표</heading>",
        "<table><tr><td>현금</td><td>6</td></tr></table>",
    ]
    if include_image:
        parts.extend(
            [
                "<heading>주석 1</heading>",
                "<image><img>note.png</img>",
                "<img-caption>이미지 설명문</img-caption></image>",
            ]
        )
    parts.append("</document>")
    return _parse_document("".join(parts))


def _context(document: ParsedDocument) -> ExportContext:
    return ExportContext(
        company_name="Sample Company",
        report_date="2026-03-10",
        report_title="감사보고서",
        receipt_date="20260310",
        rcept_no="20260310002820",
        source_rcept_no="20260310002820",
        attachment_id="opendart:20260310002820:audit.xml",
        correction_chain=("20260310002820",),
        source_url="https://dart.example/report",
        parser_version="0.1.0",
        document=document,
    )


def test_value_parser_handles_commas_parentheses_and_formula_text() -> None:
    assert parse_cell_value("1,000") == 1000
    assert parse_cell_value("(10)") == -10
    assert parse_cell_value("-") == "-"
    assert parse_cell_value("- 투자") == "- 투자"
    assert parse_cell_value("-SUM(A1:A2)") == "'-SUM(A1:A2)"
    assert parse_cell_value("=SUM(A1:A2)") == "'=SUM(A1:A2)"
    assert parse_cell_value("천원", unit_multiplier=1000) == "천원"


def test_value_parser_keeps_note_number_lists_as_text() -> None:
    assert parse_cell_value("26,34") == "26,34"
    assert parse_cell_value("4,36") == "4,36"
    assert parse_cell_value("4,34,36") == "4,34,36"
    assert parse_cell_value(",123") == ",123"
    assert parse_cell_value("12,3456") == "12,3456"
    assert parse_cell_value("1000,000") == "1000,000"


def test_value_parser_accepts_grouped_thousands_numbers() -> None:
    assert parse_cell_value("1,234") == 1234
    assert parse_cell_value("1,234,567") == 1234567
    assert parse_cell_value("(112,071)") == -112071
    assert parse_cell_value("1,234.56") == 1234.56
    assert parse_cell_value("123,456") == 123456


def test_thousands_number_format_follows_source_commas() -> None:
    assert thousands_number_format("1,234") == "#,##0"
    assert thousands_number_format("(112,071)") == "#,##0;(#,##0)"
    assert thousands_number_format("   1,000") == "#,##0"
    assert thousands_number_format("1,234.56") == "#,##0.00"
    assert thousands_number_format("1,234.5") == "#,##0.0"
    assert thousands_number_format("352") is None
    assert thousands_number_format("26,34") is None
    assert thousands_number_format("합계") is None
    assert thousands_number_format("") is None


@pytest.mark.parametrize(
    ("source_text", "expected_format"),
    [
        ("(112,071)", "#,##0;(#,##0)"),
        ("(1,234.56)", "#,##0.00;(#,##0.00)"),
        ("-112,071", "#,##0"),
        ("1,234", "#,##0"),
        ("1,234.5", "#,##0.0"),
        ("(352)", "0;(0)"),
        ("(12.5)", "0.0;(0.0)"),
        ("352", None),
    ],
)
def test_thousands_number_format_keeps_source_parentheses(
    source_text: str,
    expected_format: str | None,
) -> None:
    actual = thousands_number_format(source_text)

    assert actual == expected_format


@pytest.mark.parametrize(
    ("value", "number_format", "expected_display"),
    [
        (-112071, "#,##0;(#,##0)", "(112,071)"),
        (-1234.56, "#,##0.00;(#,##0.00)", "(1,234.56)"),
        (112071, "#,##0;(#,##0)", "112,071"),
        (-112071, "#,##0", "-112,071"),
        (-352, "0;(0)", "(352)"),
        (-12.5, "0.0;(0.0)", "(12.5)"),
    ],
)
def test_column_width_accounts_for_parenthesized_negative_display(
    value: float,
    number_format: str,
    expected_display: str,
) -> None:
    actual = workbook_layout._display_text(value, number_format)

    assert actual == expected_display


def test_value_parser_preserves_leading_indentation_for_text() -> None:
    assert parse_cell_value("   1. 현금및현금성자산") == "   1. 현금및현금성자산"
    assert parse_cell_value("\u00a0\u00a0계정") == "\u00a0\u00a0계정"
    assert parse_cell_value("   1,000") == 1000
    assert parse_cell_value("   ") == ""
    assert parse_cell_value("들여쓰기없음  ") == "들여쓰기없음"
    assert parse_cell_value(" =SUM(A1:A2)") == " =SUM(A1:A2)"


def test_export_preserves_blank_lines_and_indentation(tmp_path: Path) -> None:
    document = _parse_document(
        "<document>"
        "<heading>재무상태표</heading>"
        "<table><tr><td>   1. 현금및현금성자산</td><td>1,000</td></tr></table>"
        "<heading>손익 및 포괄손익계산서</heading>"
        "<table><tr><td>매출</td><td>(10)</td></tr></table>"
        "<heading>자본변동표</heading>"
        "<table><tr><td>자본</td><td>5</td></tr></table>"
        "<heading>현금흐름표</heading>"
        "<table><tr><td>현금</td><td>6</td></tr></table>"
        "<p>영업활동 설명</p>"
        "<p></p>"
        "<p>   들여쓴 문단</p>"
        "</document>"
    )

    result = ExcelExportService(tmp_path).export(_context(document))

    assert result.ok is True
    assert result.data is not None
    workbook = load_workbook(result.data.output_path)
    balance_sheet = workbook["재무상태표"]
    assert balance_sheet["A2"].value == "   1. 현금및현금성자산"
    cash_flow = workbook["현금흐름표"]
    assert cash_flow["A3"].value == "영업활동 설명"
    assert cash_flow["A4"].value is None
    assert cash_flow["A5"].value == "   들여쓴 문단"
    workbook.close()


def _formatted_document() -> ParsedDocument:
    return _parse_document(
        "<document>"
        "<heading>재무상태표</heading>"
        "<table><tr><td>자산총계</td><td>1,234,567</td></tr>"
        "<tr><td>손실충당금</td><td>(112,071)</td></tr>"
        "<tr><td>주석번호</td><td>26,34</td></tr>"
        "<tr><td>소액</td><td>352</td></tr>"
        "<tr><td>비율</td><td>1,234.56</td></tr>"
        "<tr><td>소수한자리</td><td>1,234.5</td></tr>"
        '<tr><td>병합금액</td><td colspan="2" rowspan="2">9,876,543</td></tr>'
        "<tr><td>다음행</td></tr></table>"
        "<heading>손익 및 포괄손익계산서</heading>"
        "<table><tr><td>매출</td><td>(10)</td></tr></table>"
        "<heading>자본변동표</heading>"
        "<table><tr><td>자본</td><td>5</td></tr></table>"
        "<heading>현금흐름표</heading>"
        "<table><tr><td>현금</td><td>6</td></tr></table>"
        "</document>"
    )


def test_export_applies_thousands_display_format(tmp_path: Path) -> None:
    result = ExcelExportService(tmp_path).export(_context(_formatted_document()))

    assert result.ok is True
    assert result.data is not None
    workbook = load_workbook(result.data.output_path)
    sheet = workbook["재무상태표"]
    assert sheet["B2"].value == 1234567
    assert sheet["B2"].number_format == "#,##0"
    assert sheet["B3"].value == -112071
    assert sheet["B3"].number_format == "#,##0;(#,##0)"
    assert sheet["B4"].value == "26,34"
    assert sheet["B4"].number_format == "General"
    assert sheet["B5"].value == 352
    assert sheet["B5"].number_format == "General"
    assert sheet["B6"].value == 1234.56
    assert sheet["B6"].number_format == "#,##0.00"
    assert sheet["B7"].value == 1234.5
    assert sheet["B7"].number_format == "#,##0.0"
    assert sheet["B8"].value == 9876543
    assert sheet["B8"].number_format == "#,##0"
    assert "B8:C9" in {str(item) for item in sheet.merged_cells.ranges}
    income_statement = workbook["손익 및 포괄손익계산서"]
    assert income_statement["B2"].value == -10
    assert income_statement["B2"].number_format == "0;(0)"
    workbook.close()


def _revalidate(
    document: ParsedDocument, path: Path
) -> tuple[bool, str, str]:
    summary = validate_document(document)
    assert summary.ok is True
    assert summary.data is not None
    result = validate_workbook(
        path,
        _context(document),
        collection_status="complete",
        summary=summary.data,
    )
    if result.ok:
        return True, "", ""
    assert result.error is not None
    return (
        False,
        str(result.error.details["issue"]),
        str(result.error.details.get("cell", "")),
    )


def test_validate_workbook_rejects_missing_thousands_format(tmp_path: Path) -> None:
    document = _formatted_document()
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    workbook["재무상태표"]["B2"].number_format = "General"
    workbook.save(path)
    workbook.close()

    ok, issue, cell = _revalidate(document, path)

    assert ok is False
    assert issue == "number_format_mismatch"
    assert cell == "B2"


def test_validate_workbook_rejects_unexpected_number_format(tmp_path: Path) -> None:
    document = _formatted_document()
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    workbook["재무상태표"]["A2"].number_format = "#,##0"
    workbook.save(path)
    workbook.close()

    ok, issue, cell = _revalidate(document, path)

    assert ok is False
    assert issue == "number_format_mismatch"
    assert cell == "A2"


def test_export_centers_merged_cells_and_sizes_columns(tmp_path: Path) -> None:
    result = ExcelExportService(tmp_path).export(_context(_formatted_document()))

    assert result.ok is True
    assert result.data is not None
    workbook = load_workbook(result.data.output_path)
    sheet = workbook["재무상태표"]
    anchor = sheet["B8"]
    assert anchor.alignment.horizontal == "center"
    assert anchor.alignment.vertical == "center"
    assert anchor.alignment.wrap_text is True
    assert sheet.column_dimensions["A"].width == pytest.approx(12.0, abs=0.05)
    assert sheet.column_dimensions["B"].width == pytest.approx(11.0, abs=0.05)
    assert workbook["수집정보"].column_dimensions["A"].width is not None
    workbook.close()


def _long_text_document() -> ParsedDocument:
    long_text = "가" * 80
    return _parse_document(
        "<document>"
        "<heading>재무상태표</heading>"
        "<table><tr><td>자산</td><td>1,000</td></tr></table>"
        "<heading>손익 및 포괄손익계산서</heading>"
        "<table><tr><td>매출</td><td>(10)</td></tr></table>"
        "<heading>자본변동표</heading>"
        "<table><tr><td>자본</td><td>5</td></tr></table>"
        "<heading>현금흐름표</heading>"
        "<table><tr><td>현금</td><td>6</td></tr></table>"
        f"<p>{long_text}</p>"
        "</document>"
    )


def test_export_wraps_long_text_and_caps_column_width(tmp_path: Path) -> None:
    result = ExcelExportService(tmp_path).export(_context(_long_text_document()))

    assert result.ok is True
    assert result.data is not None
    workbook = load_workbook(result.data.output_path)
    sheet = workbook["현금흐름표"]
    assert sheet["A3"].value == "가" * 80
    assert sheet["A3"].alignment.wrap_text is True
    assert sheet.column_dimensions["A"].width == pytest.approx(62.0, abs=0.05)
    workbook.close()


def test_validate_workbook_rejects_uncentered_merged_cell(tmp_path: Path) -> None:
    document = _formatted_document()
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    workbook["재무상태표"]["B8"].alignment = Alignment()
    workbook.save(path)
    workbook.close()

    ok, issue, cell = _revalidate(document, path)

    assert ok is False
    assert issue == "merge_alignment_mismatch"
    assert cell == "B8"


def test_validate_workbook_rejects_narrowed_column_width(tmp_path: Path) -> None:
    document = _formatted_document()
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    workbook["재무상태표"].column_dimensions["B"].width = 4
    workbook.save(path)
    workbook.close()

    ok, issue, _cell = _revalidate(document, path)

    assert ok is False
    assert issue == "column_width_mismatch"


def test_validate_workbook_rejects_deleted_column_width(tmp_path: Path) -> None:
    document = _parse_document(
        "<document>"
        "<heading>재무상태표</heading>"
        "<table><tr><td>자산</td><td>12345678901</td></tr></table>"
        "<heading>손익 및 포괄손익계산서</heading>"
        "<table><tr><td>매출</td><td>(10)</td></tr></table>"
        "<heading>자본변동표</heading>"
        "<table><tr><td>자본</td><td>5</td></tr></table>"
        "<heading>현금흐름표</heading>"
        "<table><tr><td>현금</td><td>6</td></tr></table>"
        "</document>"
    )
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    del workbook["재무상태표"].column_dimensions["B"]
    workbook.save(path)
    workbook.close()

    ok, issue, _cell = _revalidate(document, path)

    assert ok is False
    assert issue == "column_width_mismatch"


def test_validate_workbook_rejects_missing_text_wrap(tmp_path: Path) -> None:
    document = _long_text_document()
    exported = ExcelExportService(tmp_path).export(_context(document))
    assert exported.ok is True
    assert exported.data is not None
    path = exported.data.output_path
    workbook = load_workbook(path)
    workbook["현금흐름표"]["A3"].alignment = Alignment()
    workbook.save(path)
    workbook.close()

    ok, issue, cell = _revalidate(document, path)

    assert ok is False
    assert issue == "text_wrap_missing"
    assert cell == "A3"


def test_export_writes_metadata_and_marks_mixed_image_section_partial(
    tmp_path: Path,
) -> None:
    result = ExcelExportService(tmp_path).export(
        _context(_document(include_image=True))
    )

    assert result.ok is True
    assert result.data is not None
    assert result.data.collection_status == "partial"
    assert result.data.validation_status == "passed"
    assert result.data.validated_cell_count > 0
    assert result.data.validated_merge_count > 0
    assert result.data.output_path.exists()
    assert result.data.output_path.name.endswith("_부분수집.xlsx")
    assert result.data.output_path.stat().st_size > 0
    workbook = load_workbook(result.data.output_path)
    assert "A2:B2" in {str(item) for item in workbook["재무상태표"].merged_cells.ranges}
    metadata = {
        str(row[0].value): str(row[1].value)
        for row in workbook["수집정보"].iter_rows(min_col=1, max_col=2)
        if row[0].value is not None and row[1].value is not None
    }
    assert metadata["source_coverage_status"] == "passed"
    assert metadata["source_cell_count"] == metadata["captured_cell_count"]
    assert metadata["source_table_count"] == metadata["captured_table_count"]
    assert metadata["source_image_count"] == metadata["captured_image_count"]
    assert "주석 1:부분수집" in metadata["section_statuses"]
    workbook.close()


def test_export_does_not_write_when_core_statement_is_missing(tmp_path: Path) -> None:
    document = _parse_document(
        "<document><heading>주석 1</heading>"
        "<table><tr><td>내용</td><td>값</td></tr></table></document>"
    )

    result = ExcelExportService(tmp_path).export(_context(document))

    assert result.ok is False
    assert result.error is not None
    assert result.error.code.value == "CORE_STATEMENT_MISSING"
    assert list(tmp_path.glob("*.xlsx")) == []


def test_export_rejects_document_without_source_coverage(tmp_path: Path) -> None:
    document = replace(_document(), source_coverage=None)

    result = ExcelExportService(tmp_path).export(_context(document))

    assert result.ok is False
    assert result.error is not None
    assert result.error.code.value == "VALIDATION_FAILED"
    assert result.error.details["issue"] == "source_coverage_missing"
    assert list(tmp_path.glob("*.xlsx")) == []
