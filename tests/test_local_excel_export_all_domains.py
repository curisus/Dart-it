from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from mcp_types import CallToolResult
from openpyxl import load_workbook

from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.mcp_server import mcp
from dart_crawler.query_limits import EXCEL_POLICY
from dart_crawler.result import (
    ErrorCode,
    JsonObject,
    Result,
    WarningCode,
    WarningInfo,
    error_info,
)
from tests.excel_dispatch_fakes import RecordingExcelServiceFactory
from tests.excel_service_fixture_factory import empty_excel_service_responses
from tests.local_excel_export_test_support import (
    RecordingClock,
    install_local_export_runtime,
)
from tests.local_excel_mcp_result_support import structured_content
from tests.test_excel_validated_arguments import EXCEL_ARGUMENT_CASES


@pytest.mark.anyio
async def test_actual_local_tool_executes_every_domain_exactly_once(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    factory = RecordingExcelServiceFactory(empty_excel_service_responses())
    clock = RecordingClock(datetime(2026, 9, 10, 11, 12, 13, 14, tzinfo=UTC))
    output_root = install_local_export_runtime(
        monkeypatch,
        tmp_path,
        factory,
        clock,
    )

    for case in EXCEL_ARGUMENT_CASES:
        request: JsonObject = {
            "domain": case.domain.value,
            "arguments": case.raw,
        }
        called = await mcp.call_tool("export_query_excel", {"request": request})

        assert isinstance(called, CallToolResult)
        result = Result[ExcelExportResult].model_validate(
            structured_content(called)
        )
        assert result.ok is True
        assert result.warnings == ()
        exported = result.data
        assert exported is not None
        assert exported.filename == f"{case.domain.value}.xlsx"
        assert Path(exported.absolute_path).parent == output_root.resolve()
        workbook = load_workbook(
            exported.absolute_path,
            read_only=False,
            data_only=False,
        )
        try:
            assert workbook.sheetnames == ["data", "metadata", "warnings"]
            metadata = {
                row[0].value: row[1].value
                for row in workbook["metadata"].iter_rows(min_col=1, max_col=2)
            }
            assert metadata["domain"] == case.domain.value
            assert metadata["dataset_id"] == exported.dataset_id
            assert metadata["total_rows"] == exported.total_rows
        finally:
            workbook.close()

    assert len(factory.policies) == len(EXCEL_ARGUMENT_CASES) == 13
    assert factory.policies == [EXCEL_POLICY] * 13
    assert len(factory.services) == 13
    assert [service.calls[0].domain for service in factory.services] == [
        case.domain for case in EXCEL_ARGUMENT_CASES
    ]
    assert all(len(service.calls) == 1 for service in factory.services)
    assert clock.calls == 13
    assert len(list(output_root.glob("*.xlsx"))) == 13
    assert list(output_root.glob("*.lock")) == []
    assert list(output_root.glob(".*.xlsx")) == []


@pytest.mark.anyio
async def test_actual_local_tool_preserves_upstream_failure_semantics(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    upstream_error = error_info(
        ErrorCode.UPSTREAM_RATE_LIMIT,
        "rate limited",
        retryable=True,
        details={"retry_after_seconds": 7},
    )
    upstream_warning = WarningInfo(
        code=WarningCode.PARTIAL_COLLECTION,
        message="partial",
        details={"received": 0},
    )
    responses = replace(
        empty_excel_service_responses(),
        search_companies=Result.failure(
            upstream_error,
            warnings=(upstream_warning,),
            next_action="retry later",
        ),
    )
    factory = RecordingExcelServiceFactory(responses)
    clock = RecordingClock(datetime(2026, 9, 10, tzinfo=UTC))
    output_root = install_local_export_runtime(
        monkeypatch,
        tmp_path,
        factory,
        clock,
    )

    called = await mcp.call_tool(
        "export_query_excel",
        {
            "request": {
                "domain": "search_companies",
                "arguments": {"company_query": "회사"},
            }
        },
    )

    assert isinstance(called, CallToolResult)
    result = Result[ExcelExportResult].model_validate(structured_content(called))
    assert result.ok is False
    assert result.data is None
    assert result.error == upstream_error
    assert result.warnings == (upstream_warning,)
    assert result.next_action == "retry later"
    assert factory.policies == [EXCEL_POLICY]
    assert len(factory.services) == 1
    assert len(factory.services[0].calls) == 1
    assert clock.calls == 0
    assert not output_root.exists()
