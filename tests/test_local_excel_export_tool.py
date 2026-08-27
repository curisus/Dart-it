from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from mcp_types import CallToolResult
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import dart_crawler.mcp_server as mcp_server
from dart_crawler.domain import Company, Market, MatchConfidence
from dart_crawler.excel_query_export_models import ExcelExportResult
from dart_crawler.mcp_server import mcp
from dart_crawler.query_limits import EXCEL_POLICY
from dart_crawler.result import (
    ErrorCode,
    JsonValue,
    Result,
    WarningCode,
    WarningInfo,
)
from tests.excel_dispatch_fakes import RecordingExcelServiceFactory
from tests.excel_service_fixture_factory import empty_excel_service_responses
from tests.local_excel_export_test_support import (
    RecordingClock,
    install_local_export_runtime,
)
from tests.local_excel_mcp_result_support import structured_content


@pytest.mark.anyio
@pytest.mark.parametrize(
    "raw_request",
    [
        None,
        [],
        {},
        {
            "domain": "search_companies",
            "arguments": {"company_query": "회사"},
            "page_size": 1,
        },
        {
            "domain": "search_companies",
            "arguments": {"company_query": 7},
        },
        {"domain": "unknown", "arguments": {}},
        {"domain": "search_companies", "arguments": None},
        {
            "domain": "search_companies",
            "arguments": {"company_query": "회사"},
            "filename": "unsafe.xlsx",
        },
        {
            "domain": "search_companies",
            "arguments": {"company_query": "회사"},
            "output_dir": "unsafe",
        },
    ],
)
async def test_export_query_excel_rejects_non_contract_requests(
    raw_request: JsonValue,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def forbid_settings() -> None:
        raise AssertionError

    monkeypatch.setattr(mcp_server, "load_settings", forbid_settings)
    result = await mcp.call_tool("export_query_excel", {"request": raw_request})

    assert isinstance(result, CallToolResult)
    envelope = structured_content(result)
    assert envelope["ok"] is False
    assert envelope["data"] is None
    assert envelope["warnings"] == []
    assert envelope["next_action"] is None
    error = envelope["error"]
    assert isinstance(error, dict)
    assert error["code"] == ErrorCode.INVALID_INPUT.value
    assert error["retryable"] is False
    assert error["details"] == {"reason": "invalid_request"}


@pytest.mark.anyio
async def test_actual_local_tool_publishes_populated_native_workbook(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source_warning = WarningInfo(
        code=WarningCode.FALLBACK_SOURCE_USED,
        message="fallback",
    )
    company = Company(
        company_name="테스트 회사",
        corp_code="00123456",
        stock_code="123456",
        market=Market.KOSPI,
        ranking=1,
        match_confidence=MatchConfidence.EXACT,
    )
    factory = RecordingExcelServiceFactory(
        replace(
            empty_excel_service_responses(),
            search_companies=Result[tuple[Company, ...]].success(
                (company,),
                warnings=(source_warning,),
            ),
        )
    )
    clock = RecordingClock(datetime(2026, 1, 2, 3, 4, 5, 123456, tzinfo=UTC))
    output_root = install_local_export_runtime(
        monkeypatch, tmp_path, factory, clock
    )

    called = await mcp.call_tool(
        "export_query_excel",
        {
            "request": {
                "domain": "search_companies",
                "arguments": {"company_query": "테스트 회사"},
            }
        },
    )

    assert isinstance(called, CallToolResult)
    result = Result[ExcelExportResult].model_validate(structured_content(called))
    assert result.ok is True
    assert result.warnings == ()
    exported = result.data
    assert exported is not None
    assert list(exported.model_dump()) == [
        "absolute_path",
        "filename",
        "dataset_id",
        "total_rows",
        "sheet_names",
    ]
    assert exported.filename == "search_companies.xlsx"
    assert Path(exported.absolute_path) == (output_root / exported.filename).resolve()
    assert exported.total_rows == 1
    assert exported.sheet_names == ("data", "metadata", "warnings")
    assert factory.policies == [EXCEL_POLICY]
    assert len(factory.services) == 1
    assert len(factory.services[0].calls) == 1
    assert clock.calls == 1

    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == ["data", "metadata", "warnings"]
        assert all(sheet.sheet_state == "visible" for sheet in workbook.worksheets)
        assert workbook["data"].max_row == 2
        headers = [cell.value for cell in workbook["data"][1]]
        company_name_column = headers.index("company_name") + 1
        assert workbook["data"].cell(2, company_name_column).value == "테스트 회사"
        metadata = [
            (row[0].value, row[1].value)
            for row in workbook["metadata"].iter_rows(min_col=1, max_col=2)
        ]
        assert [key for key, _value in metadata] == [
            "key",
            "schema_version",
            "domain",
            "request_fingerprint",
            "source_fingerprint",
            "dataset_id",
            "total_rows",
            "generated_at_utc",
            "provenance",
            "data_sheet_names",
        ]
        assert dict(metadata)["generated_at_utc"] == "2026-01-02T03:04:05.123456Z"
        assert dict(metadata)["dataset_id"] == exported.dataset_id
        warnings = [
            tuple(cell.value for cell in row)
            for row in workbook["warnings"].iter_rows(min_col=1, max_col=4)
        ]
        assert warnings == [
            ("warning_index", "code", "message", "details"),
            (1, "FALLBACK_SOURCE_USED", "fallback", "{}"),
        ]
    finally:
        workbook.close()


@pytest.mark.anyio
async def test_actual_local_tool_publishes_empty_header_only_workbook(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    factory = RecordingExcelServiceFactory(empty_excel_service_responses())
    clock = RecordingClock(datetime(2026, 2, 3, 4, 5, 6, 7, tzinfo=UTC))
    _ = install_local_export_runtime(monkeypatch, tmp_path, factory, clock)

    called = await mcp.call_tool(
        "export_query_excel",
        {
            "request": {
                "domain": "search_companies",
                "arguments": {"company_query": "없는 회사"},
            }
        },
    )

    assert isinstance(called, CallToolResult)
    result = Result[ExcelExportResult].model_validate(structured_content(called))
    exported = result.data
    assert exported is not None
    assert exported.total_rows == 0
    assert exported.sheet_names == ("data", "metadata", "warnings")
    assert clock.calls == 1
    assert len(factory.services[0].calls) == 1
    workbook = load_workbook(exported.absolute_path, read_only=False, data_only=False)
    try:
        assert workbook["data"].max_row == 1
        assert workbook["data"][1][0].value == "company_query"
        assert workbook["warnings"].max_row == 1
        warning_sheet = workbook["warnings"]
        assert isinstance(warning_sheet, Worksheet)
        assert warning_sheet.cell(1, 1).value == "warning_index"
        metadata = {
            row[0].value: row[1].value
            for row in workbook["metadata"].iter_rows(min_col=1, max_col=2)
        }
        assert metadata["total_rows"] == 0
        assert metadata["generated_at_utc"] == "2026-02-03T04:05:06.000007Z"
    finally:
        workbook.close()
